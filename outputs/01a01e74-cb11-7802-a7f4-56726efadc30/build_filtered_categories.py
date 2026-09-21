from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont


OUTPUT_DIR = Path(r"D:\Cursor\creatisignal-app-new\outputs\01a01e74-cb11-7802-a7f4-56726efadc30")
OUTPUT_XLSX = OUTPUT_DIR / "商品品类表_精简版.xlsx"
PREVIEW_PNG = OUTPUT_DIR / "商品品类表_精简版_preview.png"

CATEGORIES = [
    ("Supplements", "保健品/营养补剂"),
    ("Jewelry", "珠宝首饰"),
    ("Intimates", "内衣"),
    ("Technology", "科技/电子产品"),
    ("Health & Wellness", "健康与养生"),
    ("Beauty", "美妆"),
    ("Accessories", "配饰"),
    ("Hair Care", "护发"),
    ("Skincare", "护肤"),
    ("Home Goods", "家居用品"),
    ("Sports & Fitness", "运动健身"),
    ("Outdoors", "户外"),
    ("Pets", "宠物"),
    ("Baby & Kids", "母婴/儿童"),
    ("Personal Care", "个人护理"),
    ("Kitchen & Dining", "厨房餐饮"),
    ("Education", "教育"),
    ("Entertainment", "娱乐"),
    ("Productivity", "效率工具"),
    ("Retail", "零售"),
    ("Travel", "旅行"),
    ("Lifestyle", "生活方式"),
    ("Automotive", "汽车"),
    ("Footwear", "鞋履/鞋类"),
    ("Toys", "玩具"),
    ("Books & Stationery", "图书文具"),
    ("Musical Instruments", "乐器"),
]

REMOVED_ENGLISH = {
    "Apparel",
    "Finance",
    "Food & Beverage",
    "FinTech",
    "Professional Services",
    "Medical",
    "Gaming",
    "Social Media",
}


def build_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "品类表"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    sheet.append(["类型", "英文", "中文"])
    for english, chinese in CATEGORIES:
        sheet.append([None, english, chinese])

    last_row = len(CATEGORIES) + 1
    sheet.merge_cells(start_row=2, start_column=1, end_row=last_row, end_column=1)
    sheet["A2"] = "商品"

    border_side = Side(style="thin", color="D9D9D9")
    standard_border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="222222")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = standard_border

    sheet["A2"].fill = PatternFill("solid", fgColor="D9E2F3")
    sheet["A2"].font = Font(name="Microsoft YaHei", size=11, color="222222")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"].border = standard_border

    for row in range(2, last_row + 1):
        for column in (2, 3):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Microsoft YaHei", size=10, color="222222")
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = standard_border
        sheet.row_dimensions[row].height = 25

    sheet.row_dimensions[1].height = 27
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 27
    sheet.column_dimensions["C"].width = 23
    sheet.auto_filter.ref = f"B1:C{last_row}"
    sheet.sheet_view.zoomScale = 100
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.save(OUTPUT_XLSX)


def validate_workbook() -> dict[str, object]:
    with zipfile.ZipFile(OUTPUT_XLSX, "r") as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise RuntimeError(f"Corrupt ZIP member: {bad_member}")

    workbook = load_workbook(OUTPUT_XLSX, read_only=False, data_only=False)
    if workbook.sheetnames != ["品类表"]:
        raise RuntimeError(f"Unexpected worksheets: {workbook.sheetnames}")
    sheet = workbook["品类表"]
    headers = [sheet.cell(1, column).value for column in range(1, 4)]
    if headers != ["类型", "英文", "中文"]:
        raise RuntimeError(f"Unexpected headers: {headers}")
    if (sheet.max_row, sheet.max_column) != (28, 3):
        raise RuntimeError(f"Unexpected dimensions: {sheet.max_row}x{sheet.max_column}")
    if "A2:A28" not in {str(item) for item in sheet.merged_cells.ranges}:
        raise RuntimeError(f"Missing merged type range: {sheet.merged_cells.ranges}")
    if sheet["A2"].value != "商品":
        raise RuntimeError("Missing 商品 type label")

    english_values = [sheet.cell(row, 2).value for row in range(2, 29)]
    if len(english_values) != 27 or len(set(english_values)) != 27:
        raise RuntimeError("Category count or uniqueness check failed")
    unexpected_removed = sorted(REMOVED_ENGLISH.intersection(english_values))
    if unexpected_removed:
        raise RuntimeError(f"Removed categories still present: {unexpected_removed}")
    if english_values != [item[0] for item in CATEGORIES]:
        raise RuntimeError("Category order changed")

    return {
        "rows_retained": len(english_values),
        "rows_removed": len(REMOVED_ENGLISH),
        "sheet": sheet.title,
        "merged_range": "A2:A28",
        "filter_range": sheet.auto_filter.ref,
        "file_size": OUTPUT_XLSX.stat().st_size,
        "sha256": hashlib.sha256(OUTPUT_XLSX.read_bytes()).hexdigest(),
    }


def render_preview() -> None:
    widths = [190, 300, 260]
    header_height = 44
    row_height = 35
    width = sum(widths)
    height = header_height + len(CATEGORIES) * row_height
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    regular = ImageFont.truetype(str(regular_path), 17) if regular_path.exists() else ImageFont.load_default()
    bold = ImageFont.truetype(str(bold_path), 18) if bold_path.exists() else regular

    x = 0
    for column_width, header in zip(widths, ["类型", "英文", "中文"]):
        draw.rectangle((x, 0, x + column_width, header_height), fill="#FFF2CC", outline="#D9D9D9")
        draw.text((x + 10, 11), header, fill="#222222", font=bold)
        x += column_width

    type_bottom = header_height + len(CATEGORIES) * row_height
    draw.rectangle((0, header_height, widths[0], type_bottom), fill="#D9E2F3", outline="#D9D9D9")
    label_box = draw.textbbox((0, 0), "商品", font=regular)
    label_width = label_box[2] - label_box[0]
    label_height = label_box[3] - label_box[1]
    draw.text(
        ((widths[0] - label_width) / 2, header_height + (type_bottom - header_height - label_height) / 2),
        "商品",
        fill="#222222",
        font=regular,
    )

    for index, (english, chinese) in enumerate(CATEGORIES):
        y = header_height + index * row_height
        x = widths[0]
        for column_width, value in zip(widths[1:], [english, chinese]):
            draw.rectangle((x, y, x + column_width, y + row_height), fill="white", outline="#D9D9D9")
            draw.text((x + 8, y + 7), value, fill="#222222", font=regular)
            x += column_width
    image.save(PREVIEW_PNG)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_workbook()
    validation = validate_workbook()
    render_preview()
    print(
        json.dumps(
            {
                "output": str(OUTPUT_XLSX),
                "preview": str(PREVIEW_PNG),
                "validation": validation,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
