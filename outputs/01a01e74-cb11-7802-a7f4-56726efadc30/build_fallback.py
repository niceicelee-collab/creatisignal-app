from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont


OUTPUT_DIR = Path(r"D:\Cursor\creatisignal-app-new\outputs\01a01e74-cb11-7802-a7f4-56726efadc30")
SOURCE_JS = OUTPUT_DIR / "build_cases.mjs"
OUTPUT_XLSX = OUTPUT_DIR / "CreatiSignal_TikTokShop_Seedance2_Case_Library.xlsx"
PREVIEW_PNG = OUTPUT_DIR / "CreatiSignal_TikTokShop_Seedance2_Case_Library_preview.png"


def decode_js_string(value: str) -> str:
    return json.loads(f'"{value}"')


def extract_entries(source: str) -> tuple[list[dict[str, str]], str]:
    tail_match = re.search(r'const commonTail = "((?:\\.|[^"\\])*)";', source)
    if not tail_match:
        raise RuntimeError("Could not extract common prompt tail")
    common_tail = decode_js_string(tail_match.group(1))

    pattern = re.compile(
        r'\{\s*industry: "((?:\\.|[^"\\])*)",\s*'
        r'category: "((?:\\.|[^"\\])*)",\s*'
        r'keywords: "((?:\\.|[^"\\])*)",\s*'
        r'prompt: "((?:\\.|[^"\\])*)"\s*\}',
        re.DOTALL,
    )
    entries = [
        {
            "industry": decode_js_string(industry),
            "category": decode_js_string(category),
            "keywords": decode_js_string(keywords),
            "prompt": decode_js_string(prompt),
        }
        for industry, category, keywords, prompt in pattern.findall(source)
    ]
    if len(entries) != 32:
        raise RuntimeError(f"Expected 32 entries, extracted {len(entries)}")
    return entries, common_tail


def build_workbook(entries: list[dict[str, str]], common_tail: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "案例库"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"

    headers = ["行业", "品类", "产品关键词", "完整提示词"]
    ws.append(headers)
    for entry in entries:
        ws.append(
            [
                entry["industry"],
                entry["category"],
                entry["keywords"],
                f'{entry["prompt"]}\n\n{common_tail}',
            ]
        )

    navy = "102A43"
    text_color = "243B53"
    light_border = Side(style="thin", color="D9E2EC")
    medium_border = Side(style="medium", color=navy)
    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(name="Microsoft YaHei", size=12, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color=text_color)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium_border, bottom=medium_border)
    ws.row_dimensions[1].height = 30

    industry_colors = {
        "服饰时尚": "F3E8FF",
        "美妆健康": "FCE7F3",
        "家居家装": "DCFCE7",
        "数码家电": "DBEAFE",
        "母婴宠物": "FEF3C7",
        "运动出行": "FFEDD5",
        "兴趣收藏": "EDE9FE",
        "食品饮料": "FEE2E2",
    }

    for row_idx in range(2, ws.max_row + 1):
        industry = ws.cell(row_idx, 1).value
        group_fill = PatternFill("solid", fgColor=industry_colors[industry])
        keyword_fill = PatternFill("solid", fgColor="F8FAFC")
        for col_idx in range(1, 5):
            cell = ws.cell(row_idx, col_idx)
            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="center" if col_idx <= 2 else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=light_border)
        ws.cell(row_idx, 1).font = Font(
            name="Microsoft YaHei", size=10, bold=True, color=navy
        )
        ws.cell(row_idx, 1).fill = group_fill
        ws.cell(row_idx, 2).fill = group_fill
        ws.cell(row_idx, 3).fill = keyword_fill
        ws.row_dimensions[row_idx].height = 122

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 27
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 118

    table_ref = f"A1:D{ws.max_row}"
    table = Table(displayName="CaseLibraryTable", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.auto_filter.ref = table_ref

    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.zoomScale = 75

    wb.save(OUTPUT_XLSX)


def validate_workbook(entries: list[dict[str, str]]) -> dict[str, object]:
    with zipfile.ZipFile(OUTPUT_XLSX, "r") as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise RuntimeError(f"Corrupt ZIP member: {bad_member}")

    wb = load_workbook(OUTPUT_XLSX, read_only=False, data_only=False)
    if wb.sheetnames != ["案例库"]:
        raise RuntimeError(f"Unexpected sheets: {wb.sheetnames}")
    ws = wb["案例库"]
    expected_headers = ["行业", "品类", "产品关键词", "完整提示词"]
    actual_headers = [ws.cell(1, col).value for col in range(1, 5)]
    if actual_headers != expected_headers:
        raise RuntimeError(f"Header mismatch: {actual_headers}")
    if (ws.max_row, ws.max_column) != (33, 4):
        raise RuntimeError(f"Unexpected dimensions: {ws.max_row}x{ws.max_column}")
    if str(ws.freeze_panes) != "C2":
        raise RuntimeError(f"Unexpected freeze pane: {ws.freeze_panes}")
    if "CaseLibraryTable" not in ws.tables:
        raise RuntimeError("Missing Excel table")
    if ws.tables["CaseLibraryTable"].ref != "A1:D33":
        raise RuntimeError(f"Unexpected table range: {ws.tables['CaseLibraryTable'].ref}")

    blank_rows = []
    missing_complete_prompt = []
    for row_idx in range(2, 34):
        values = [ws.cell(row_idx, col).value for col in range(1, 5)]
        if any(value is None or str(value).strip() == "" for value in values):
            blank_rows.append(row_idx)
        prompt = str(values[3])
        if "竖屏9:16" not in prompt or "15秒" not in prompt or "产品在首帧清楚出现" not in prompt:
            missing_complete_prompt.append(row_idx)
    if blank_rows:
        raise RuntimeError(f"Blank required cells in rows: {blank_rows}")
    if missing_complete_prompt:
        raise RuntimeError(f"Incomplete prompts in rows: {missing_complete_prompt}")

    industry_counts = Counter(ws.cell(row_idx, 1).value for row_idx in range(2, 34))
    if len(industry_counts) != 8 or set(industry_counts.values()) != {4}:
        raise RuntimeError(f"Unexpected industry distribution: {industry_counts}")

    digest = hashlib.sha256(OUTPUT_XLSX.read_bytes()).hexdigest()
    return {
        "rows": len(entries),
        "columns": 4,
        "industries": dict(industry_counts),
        "freeze_panes": str(ws.freeze_panes),
        "table_ref": ws.tables["CaseLibraryTable"].ref,
        "file_size": OUTPUT_XLSX.stat().st_size,
        "sha256": digest,
    }


def weighted_wrap(text: str, max_units: int, max_lines: int) -> list[str]:
    lines: list[str] = []
    current = ""
    current_units = 0
    for char in text:
        if char == "\n":
            lines.append(current)
            current = ""
            current_units = 0
            if len(lines) >= max_lines:
                break
            continue
        units = 2 if ord(char) > 127 else 1
        if current and current_units + units > max_units:
            lines.append(current)
            current = char
            current_units = units
            if len(lines) >= max_lines:
                break
        else:
            current += char
            current_units += units
    if len(lines) < max_lines and current:
        lines.append(current)
    if len(lines) == max_lines and sum(len(line) for line in lines) < len(text.replace("\n", "")):
        lines[-1] = lines[-1][:-1] + "…"
    return lines


def render_preview(entries: list[dict[str, str]], common_tail: str) -> None:
    widths = [150, 240, 520, 1080]
    header_height = 56
    row_height = 190
    canvas_width = sum(widths)
    canvas_height = header_height + len(entries) * row_height
    image = Image.new("RGB", (canvas_width, canvas_height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_font_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    font = ImageFont.truetype(str(font_path), 15) if font_path.exists() else ImageFont.load_default()
    bold_font = (
        ImageFont.truetype(str(bold_font_path), 17)
        if bold_font_path.exists()
        else font
    )

    headers = ["行业", "品类", "产品关键词", "完整提示词"]
    x = 0
    for width, header in zip(widths, headers):
        draw.rectangle((x, 0, x + width, header_height), fill="#102A43", outline="#102A43")
        draw.text((x + 10, 15), header, fill="white", font=bold_font)
        x += width

    industry_colors = {
        "服饰时尚": "#F3E8FF",
        "美妆健康": "#FCE7F3",
        "家居家装": "#DCFCE7",
        "数码家电": "#DBEAFE",
        "母婴宠物": "#FEF3C7",
        "运动出行": "#FFEDD5",
        "兴趣收藏": "#EDE9FE",
        "食品饮料": "#FEE2E2",
    }
    wrap_units = [12, 20, 50, 112]
    max_lines = [4, 7, 9, 10]
    for idx, entry in enumerate(entries):
        y = header_height + idx * row_height
        full_prompt = f'{entry["prompt"]}\n\n{common_tail}'
        values = [entry["industry"], entry["category"], entry["keywords"], full_prompt]
        fills = [industry_colors[entry["industry"]], industry_colors[entry["industry"]], "#F8FAFC", "white"]
        x = 0
        for col_idx, (width, value, fill) in enumerate(zip(widths, values, fills)):
            draw.rectangle((x, y, x + width, y + row_height), fill=fill, outline="#D9E2EC")
            lines = weighted_wrap(value, wrap_units[col_idx], max_lines[col_idx])
            draw.multiline_text((x + 8, y + 8), "\n".join(lines), fill="#243B53", font=font, spacing=3)
            x += width

    image.save(PREVIEW_PNG)


def main() -> None:
    source = SOURCE_JS.read_text(encoding="utf-8")
    entries, common_tail = extract_entries(source)
    build_workbook(entries, common_tail)
    validation = validate_workbook(entries)
    render_preview(entries, common_tail)
    print(json.dumps({
        "output": str(OUTPUT_XLSX),
        "preview": str(PREVIEW_PNG),
        "validation": validation,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
