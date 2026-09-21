from __future__ import annotations

from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont


OUTPUT_DIR = Path(r"D:\Cursor\creatisignal-app-new\outputs\01a01e74-cb11-7802-a7f4-56726efadc30")
OUTPUT_FILE = OUTPUT_DIR / "CreatiSignal_Seedance2_红框细分品类全覆盖_GTM案例库.xlsx"
PREVIEW_FILE = OUTPUT_DIR / "CreatiSignal_Seedance2_红框细分品类全覆盖_GTM案例库_preview.png"


def case(industry, subcategory, product, keywords, lock, scene, hook, demo, proof, final, vo1, vo2, vo3, vo4, realism):
    category = f"{subcategory}｜{product}"
    prompt = (
        f"生成一条15秒、竖屏9:16的TikTok Shop原生UGC广告。使用@图片1 @图片2 @图片3作为{product}参考，"
        f"保持{lock}一致。产品在首帧清楚出现。场景：{scene}。\n"
        f"0-3秒：{hook} 英文口播：\"{vo1}\"\n"
        f"3-7秒：{demo} 英文口播：\"{vo2}\"\n"
        f"7-11秒：{proof} 英文口播：\"{vo3}\"\n"
        f"11-15秒：{final} 英文口播：\"{vo4}\"\n"
        f"真实手机拍摄质感，自然光，近景细节清楚；{realism}。手、人物与产品接触关系真实，避免穿模、额外手指和物体瞬移。"
        "不得改变产品Logo、文字、颜色、结构、配件数量或包装，不虚构价格、折扣、评价和未经证实的功效；结尾预留字幕与CTA空间。"
    )
    return [industry, category, keywords, prompt]


ROWS = [
    case(
        "服饰时尚", "男女装", "男士功能Polo",
        "主搜：men's performance polo moisture wicking；长尾：quick dry golf shirt, stretch polo, wrinkle resistant polo, fabric close up",
        "领型、门襟纽扣、面料纹理、版型、颜色和Logo",
        "明亮卧室衣柜前，男性创作者自然试穿",
        "真人把Polo举到胸前，快速切到上身效果，正面Logo与领口可见。",
        "近景捏拉袖口和侧腰面料，再整理领口，展示弹性与垂坠感。",
        "真人抬臂、转身并做轻度挥杆动作，展示活动空间和背部版型。",
        "正侧背三角度快速切换，最后手指产品细节并面向镜头。",
        "This is the polo I keep reaching for.",
        "The fabric feels light and moves with me.",
        "It stays clean-looking from work to the weekend.",
        "An easy upgrade for everyday outfits.",
        "衣物随身体自然产生褶皱，纽扣和Logo位置稳定，不改变人体比例",
    ),
    case(
        "服饰时尚", "内衣", "无缝塑身连体衣",
        "主搜：women's seamless shapewear bodysuit tummy control；长尾：adjustable straps, snap closure, under dress shapewear, front back detail",
        "肩带、领口、裆部扣合、无缝纹理、颜色和裁片结构",
        "更衣镜前，女性穿着基础打底衣进行叠穿演示",
        "真人一手拿产品、一手展示平整无缝面料，产品轮廓完整可见。",
        "镜头切到肩带调节、裆部扣合和面料回弹近景，再叠穿连衣裙。",
        "以自然站姿展示正面与侧面轮廓，不做夸张前后对比。",
        "真人在镜前整理外搭，手持产品包装完成hero展示。",
        "This is the smooth base layer under my outfits.",
        "The seamless fabric feels flexible and easy to layer.",
        "I like how clean it looks under a fitted dress.",
        "A simple wardrobe staple for dress-up days.",
        "服装贴合但不过度压缩身体，保持肩带、扣件和裁片的真实连接",
    ),
    case(
        "服饰时尚", "儿童", "儿童连帽防晒衣",
        "主搜：kids lightweight sun protection hoodie；长尾：UPF hoodie kids, quick dry jacket, zip up sun shirt, fabric and zipper detail",
        "帽型、拉链、口袋、袖口、面料颜色和印花位置",
        "家庭玄关与户外公园，家长视角拍摄，不突出儿童脸部",
        "家长把防晒衣从挂钩取下，衣服正面和帽子在首帧完整出现。",
        "手部拉开拉链、展示口袋与轻薄面料，再自然给孩子穿上。",
        "孩子从背面和侧面慢跑、抬臂，展示帽子与袖口的活动状态。",
        "衣服挂回玄关，家长指向关键细节，画面留出CTA位置。",
        "This is the layer we grab before heading outside.",
        "It feels light, easy to zip, and simple to pack.",
        "The fit gives plenty of room for active play.",
        "Ready for park days and family trips.",
        "儿童动作自然安全，服装尺寸一致，拉链不得与身体或手指穿插",
    ),
    case(
        "服饰时尚", "鞋靴", "轻量运动鞋",
        "主搜：lightweight breathable walking shoes；长尾：mesh athletic sneakers, flexible sole, cushioned walking shoes, side sole detail",
        "鞋型、鞋面网布、鞋带孔、鞋底纹路、颜色和Logo",
        "城市步道与家中换鞋凳，日常运动穿搭",
        "真人把一只鞋靠近镜头，随后快速切到双脚穿着落地。",
        "手指按压鞋面并弯折前掌，近景展示网布和鞋底纹理。",
        "低机位跟拍行走、转向和轻快上台阶，双鞋外观始终一致。",
        "鞋子整齐放在台阶上，真人一脚入镜停在旁边形成hero画面。",
        "These feel light the second I put them on.",
        "The upper flexes without losing its shape.",
        "They move easily through my everyday walk.",
        "A clean pair for errands, travel, and more.",
        "步态和落地受力真实，鞋底与地面明确接触，左右鞋不互换细节",
    ),
    case(
        "服饰时尚", "箱包", "多功能通勤托特包",
        "主搜：women work tote bag laptop compartment；长尾：zipper closure tote, multiple pockets, water bottle pocket, inside compartments",
        "包型、提手、拉链、口袋分区、五金、颜色和Logo",
        "办公桌与通勤玄关，女性创作者收纳随身物品",
        "打开托特包并把电脑、杯子和小物摆在旁边，包体首帧完整可见。",
        "依次把电脑、水杯和钥匙放入对应分区，近景展示拉链闭合。",
        "单手提起装满物品的包，切到肩背行走和放上办公桌。",
        "俯拍包内整齐分区，再合上包展示正面hero角度。",
        "Everything I need actually has a place in this tote.",
        "The compartments keep my workday essentials organized.",
        "It carries easily from the commute to my desk.",
        "This is my grab-and-go work bag.",
        "放入物品的尺寸与包内空间匹配，拉链沿真实轨迹开合，提手承重自然",
    ),
    case(
        "服饰时尚", "配件", "可调节棒球帽",
        "主搜：adjustable baseball cap unisex；长尾：washed cotton dad hat, curved brim cap, metal buckle strap, front side back view",
        "帽冠、帽檐弧度、后部调节扣、刺绣位置、颜色和Logo",
        "玄关镜前与户外街道，轻松日常穿搭",
        "真人把帽子贴近镜头旋转半圈，首帧看到正面刺绣与帽檐。",
        "近景展示后部调节扣并调整一格，随后自然戴上。",
        "镜前展示正面、侧面和背面佩戴效果，轻触帽檐。",
        "真人摘下帽子拿在肩侧，背景留白完成hero展示。",
        "This cap fixes an outfit in seconds.",
        "The adjustable back makes the fit easy.",
        "I like the relaxed shape from every angle.",
        "An effortless finish for everyday looks.",
        "帽子与头部稳定接触，帽檐不变形穿插，调节扣结构前后一致",
    ),
    case(
        "服饰时尚", "珠宝", "极简叠戴项链",
        "主搜：minimalist layered necklace for women；长尾：adjustable chain necklace, delicate pendant, gold tone necklace, clasp detail",
        "链条层数、吊坠形状、扣头、延长链、金属颜色和包装",
        "梳妆台与窗边自然光，女性创作者进行佩戴分享",
        "项链平放在首饰托盘，手指轻提吊坠，让产品在首帧清晰闪现。",
        "微距展示扣头与延长链，双手在颈后完成真实佩戴。",
        "锁骨近景从正面到侧面轻移，展示链条层次和自然反光。",
        "真人把包装放在脸侧，项链仍清楚可见，完成hero画面。",
        "This layered necklace makes a simple top feel finished.",
        "The adjustable chain makes styling really easy.",
        "I love the delicate shine without it feeling too much.",
        "A small detail that works with almost everything.",
        "链条保持连续，不粘连皮肤或凭空断裂，吊坠数量和位置不可变化",
    ),
    case(
        "美妆健康", "美妆个护", "高光泽唇油",
        "主搜：hydrating tinted lip oil glossy non sticky；长尾：large applicator lip oil, clear tube, lip swatch, applicator close up",
        "透明管身、盖子、标签、唇油颜色、刷头尺寸和光泽质地",
        "明亮卧室梳妆台，真实美妆博主对镜口播",
        "真人把唇油放在脸侧，嘴巴和产品都清楚可见。",
        "右手旋开盖子抽出刷头，在左手手背划出一条轻透试色。",
        "右手用刷头轻涂嘴唇中央，近景呈现自然水光与轻微颜色。",
        "左手持管身放在脸侧，嘴唇和手背试色同时可见。",
        "This lip oil gives me instant glossy lips.",
        "The applicator feels plush and the tint stays sheer.",
        "Look at that smooth, juicy-looking shine.",
        "This one is staying in my everyday bag.",
        "左手负责管身和手背，右手负责刷头；刷头与手背、嘴唇接触清楚，质地轻薄通透",
    ),
    case(
        "美妆健康", "保健品", "成人维生素软糖",
        "主搜：adult multivitamin gummies；长尾：women multivitamin gummies, vegan vitamin gummies, berry flavor gummies, bottle gummy close up",
        "瓶型、瓶盖、标签文字、软糖颜色形状、数量感和包装",
        "明亮厨房早餐台，仅由成年人演示",
        "成年人手持瓶身靠近镜头，标签正面在首帧清楚出现。",
        "旋开瓶盖，将两粒软糖倒入掌心，展示真实形状与颜色。",
        "把瓶子放在早餐杯旁，镜头强调便携包装和日常摆放场景。",
        "合上瓶盖并将产品放回台面，手指包装完成hero展示。",
        "These gummies make my morning routine feel simple.",
        "The bottle is easy to keep right by breakfast.",
        "I like the familiar gummy format and everyday convenience.",
        "Always check the label and use as directed.",
        "只展示包装与食用场景，不暗示治疗、减重、治愈或确定健康结果；软糖倒出数量保持连续",
    ),
    case(
        "家居家装", "家居日用", "蒸汽拖把",
        "主搜：steam mop for hardwood and tile floors；长尾：detachable water tank, washable microfiber pads, lightweight floor cleaner, product in use",
        "机身轮廓、水箱、拖布、手柄连接、颜色和Logo",
        "真实厨房与客厅地面，日常清洁场景",
        "镜头对准一小块可见污渍，拖把随即进入画面，产品首帧可见。",
        "手部装入水箱并固定拖布，连接动作清楚，不展示不存在的配件。",
        "低机位跟拍拖把沿地面直线移动，再切到清洁后的同一区域。",
        "拖把直立停靠，水箱和拖布近景可见，画面留出CTA。",
        "This makes my quick floor reset feel much easier.",
        "The setup is simple and the head moves smoothly.",
        "Here is the same spot after one focused pass.",
        "A practical tool for everyday cleanup.",
        "拖把头始终贴地，水汽轻微自然，不夸张表现瞬间消毒或适用于所有地板",
    ),
    case(
        "家居家装", "布艺", "遮光窗帘",
        "主搜：blackout curtains bedroom grommet 2 panels；长尾：room darkening curtains, thermal insulated curtain, linen texture, fabric close up",
        "帘片数量、尺寸比例、打孔环、面料纹理、颜色和包装",
        "卧室窗边，同一机位展示自然光变化",
        "手拉开一侧窗帘露出明亮窗户，产品和面料在首帧可见。",
        "近景摸过布料与打孔环，再沿窗杆顺畅拉合两片窗帘。",
        "固定机位展示拉合前后的室内亮度差异，保留真实边缘漏光。",
        "两片窗帘自然垂落，手持包装在画面一侧完成hero展示。",
        "These curtains change the feel of the room in one pull.",
        "The fabric has a soft texture and hangs neatly.",
        "You can see the real difference in this same window.",
        "A simple bedroom refresh with a practical purpose.",
        "窗帘沿轨道移动并产生自然褶皱，不让房间瞬间全黑，不改变窗户结构",
    ),
    case(
        "家居家装", "厨房用品", "便携无线搅拌杯",
        "主搜：portable rechargeable blender personal size；长尾：USB blender cup, travel smoothie blender, blade detail, leak resistant lid",
        "杯体、杯盖、刀头、按键、充电口、刻度和Logo",
        "明亮厨房台面，制作单杯水果饮品",
        "透明杯内放入切好的水果和液体，产品在首帧完整出现。",
        "手部对齐杯体与底座并锁紧，按下按钮启动真实旋转。",
        "固定近景展示内容物逐渐混合，停止后旋开杯盖倒入玻璃杯。",
        "搅拌杯与成品并排，真人拿起杯体完成hero展示。",
        "This is my shortcut to a quick single-serve blend.",
        "The cup locks into place and starts with one button.",
        "You can watch the texture change as it blends.",
        "Easy to make, carry, and clean up.",
        "刀头只在杯体正确锁定后旋转，液体不穿过杯壁，不夸张处理冰块或硬物能力",
    ),
    case(
        "家居家装", "家具", "可折叠边桌",
        "主搜：small folding side table for living room；长尾：portable end table, foldable tray table, compact bedside table, folded and open view",
        "桌面形状、支腿结构、折叠铰链、颜色、纹理和尺寸比例",
        "小户型客厅沙发旁，空间改造演示",
        "折叠状态的边桌靠墙放置，真人单手拿起，产品首帧清楚可见。",
        "双手展开支腿并锁定铰链，将桌面平稳立在沙发旁。",
        "依次放上杯子、书和小台灯，侧面展示承托与占地大小。",
        "快速收起边桌再展开，最终保留整洁客厅hero画面。",
        "This little table gives the corner an instant purpose.",
        "It opens quickly and fits right beside the sofa.",
        "There is room for the essentials without taking over the space.",
        "A flexible piece for compact rooms.",
        "铰链按真实方向转动并锁定，桌腿始终接触地面，不夸张承重",
    ),
    case(
        "家居家装", "建材", "自粘防溅墙板",
        "主搜：peel and stick backsplash tile kitchen；长尾：self adhesive wall tile, waterproof backsplash panel, subway tile sheet, installation close up",
        "单片尺寸、拼缝、表面纹理、颜色、背胶和包装数量",
        "厨房水槽后方的一小块干净墙面，DIY局部改造",
        "真人把一片墙板贴在待改造墙面旁比对，正面纹理首帧可见。",
        "近景撕开部分背纸，对齐边缘后从一侧缓慢压贴。",
        "手掌与刮板排平表面，镜头展示真实拼缝和局部完成效果。",
        "同一角度展示改造区域，手持剩余墙板与包装完成hero展示。",
        "This peel-and-stick panel makes a small update feel doable.",
        "I line up the edge, peel the backing, and press it down.",
        "The close-up shows the real texture and seams.",
        "A focused DIY refresh without changing the whole room.",
        "只展示局部小面积安装，不让整面墙瞬间变化；板材边缘、拼缝和手部按压关系真实",
    ),
    case(
        "家居家装", "五金", "棘轮螺丝刀套装",
        "主搜：ratcheting screwdriver set with bits；长尾：multi bit screwdriver, magnetic bit holder, compact tool kit, ratchet mechanism close up",
        "手柄、棘轮拨杆、批头数量、收纳盒卡位、颜色和Logo",
        "家庭工作台，组装小家具的真实维修场景",
        "打开收纳盒展示完整批头排列，螺丝刀横放在首帧中央。",
        "手指选择一个批头并插入磁吸接口，近景切换棘轮方向。",
        "螺丝刀与螺钉真实咬合，连续转动完成一处家具紧固。",
        "工具归位到收纳盒，合盖前展示全套hero俯拍。",
        "This compact driver keeps the bits I actually use together.",
        "The bit clicks in and the ratchet direction switches easily.",
        "Here it is tightening a real furniture screw.",
        "A useful kit to keep within reach at home.",
        "批头必须由手装入接口，螺钉沿轴线旋入，不凭空更换批头或增加配件",
    ),
    case(
        "数码家电", "手机与数码", "磁吸手机壳折叠支架",
        "主搜：magnetic phone case with stand；长尾：MagSafe compatible case, foldable kickstand, magnetic ring phone case, front back side view",
        "壳体开孔、磁吸环、折叠支架、颜色、边框和Logo",
        "咖啡桌与办公桌，手机日常使用演示",
        "手机壳背面对镜头，磁吸环与支架在首帧清楚出现。",
        "双手把手机压入壳体，四角依次卡紧，再展开支架。",
        "手机分别横放和竖放，近景展示支架与桌面的真实受力角度。",
        "折回支架并把手机立在键盘旁，展示正背面hero角度。",
        "This case turns into a stand exactly when I need it.",
        "The phone clicks in and the ring folds out smoothly.",
        "It works for both vertical scrolling and horizontal viewing.",
        "One case, less to carry around.",
        "手机型号与壳体开孔匹配，支架沿铰链旋转并接触桌面，不悬空承重",
    ),
    case(
        "数码家电", "电脑办公", "便携显示器",
        "主搜：portable monitor for laptop USB C；长尾：15.6 inch travel monitor, built in stand, second screen for laptop, ports side view",
        "屏幕比例、边框、支架、接口位置、线缆和Logo",
        "家庭办公桌，笔记本双屏工作场景",
        "便携显示器从保护套中取出并立在笔记本旁，产品首帧完整出现。",
        "近景把USB-C线插入对应接口，再展开内置支架。",
        "屏幕亮起显示简洁色块与窗口布局，镜头侧移展示双屏关系。",
        "收拢线缆并展示薄度、接口和桌面全景，预留CTA。",
        "This is how I add a second screen without a permanent setup.",
        "One cable connects the display and keeps the desk simple.",
        "The extra space makes my workflow easier to organize.",
        "Then it packs flat when I am done.",
        "线缆从真实端口连接，屏幕内容稳定且不生成可辨识品牌界面或错误文字",
    ),
    case(
        "数码家电", "家电", "无线手持吸尘器",
        "主搜：cordless handheld vacuum for car and home；长尾：crevice tool vacuum, clear dust cup, mini cordless vacuum, accessories included",
        "主机、透明尘杯、滤芯、吸嘴配件、按键、颜色和Logo",
        "沙发缝隙与汽车座椅，快速清洁场景",
        "少量可见碎屑旁放置手持吸尘器，主机在首帧清楚出现。",
        "手部将缝隙吸嘴卡入主机并按下开关，部件连接过程完整。",
        "吸嘴沿沙发缝隙缓慢移动，透明尘杯内出现对应少量碎屑。",
        "关机后拆下尘杯在垃圾桶上方展示，再恢复产品hero画面。",
        "This little vacuum is made for the spots my big one misses.",
        "The crevice tool clicks on for tight spaces.",
        "You can see exactly what it picked up from this one area.",
        "A handy reset for the sofa or car.",
        "碎屑被吸入路径连续，不让大面积污物瞬间消失；配件只能按真实接口装卸",
    ),
    case(
        "母婴宠物", "母婴用品", "一键折叠婴儿车",
        "主搜：compact travel stroller one hand fold；长尾：lightweight baby stroller, airplane travel stroller, folded size, front side folded view",
        "车架、座舱、遮阳篷、轮组、折叠关节、安全带和颜色",
        "家庭玄关与平整户外步道，仅由成年人操作空车",
        "折叠婴儿车立在玄关，成年人单手提起，首帧看到完整车体。",
        "按真实解锁位置展开车架，关节依次到位并确认锁定。",
        "推行空车转弯、过小坡并踩下刹车，低机位展示轮组。",
        "单手折叠回紧凑状态，放入汽车后备箱完成hero展示。",
        "This travel stroller folds down in one smooth routine.",
        "The frame opens, locks, and is ready to roll.",
        "The wheels turn easily on a normal walking path.",
        "Compact when the trip is over.",
        "不放置婴儿进行功能测试；折叠关节按真实顺序运动，车轮始终接触地面",
    ),
    case(
        "母婴宠物", "宠物用品", "宠物自动饮水机",
        "主搜：cat water fountain stainless steel；长尾：quiet pet fountain, replaceable filter, water level window, fountain in use",
        "水盆、出水口、滤芯、水位窗、电源线、颜色和Logo",
        "明亮客厅宠物角，真实猫咪自然接近",
        "饮水机放在宠物垫上，细小水流运行，产品首帧完整出现。",
        "成年人双手拆开上盖，放入滤芯并加水，再正确复位。",
        "固定镜头记录猫咪靠近并自然饮水，水流保持稳定不过度飞溅。",
        "猫咪离开后切到水位窗与产品正面，画面留出CTA。",
        "This fountain keeps the water station feeling fresh and tidy.",
        "The filter and top lift out for a simple refill routine.",
        "My cat can approach it naturally without a loud splash.",
        "A neat upgrade for the pet corner.",
        "宠物行为不强迫，舌头与水流接触自然；滤芯、上盖和泵体不瞬移",
    ),
    case(
        "运动出行", "运动与户外", "阻力带套装",
        "主搜：resistance loop bands set with carry bag；长尾：fabric resistance bands, different resistance levels, workout demonstration, set flat lay",
        "每条带子的宽度、颜色、阻力标识、缝线、数量和收纳袋",
        "客厅瑜伽垫，真实居家训练场景",
        "整套阻力带按颜色平铺，真人拿起其中一条，首帧看到完整套装。",
        "把阻力带套到膝上并调整到平整位置，近景展示缝线与弹性。",
        "完成慢速侧步和臀桥各一次，带子随动作自然拉伸回弹。",
        "训练结束将全部阻力带卷好放入收纳袋，完成hero俯拍。",
        "This band set makes a quick home workout easy to start.",
        "The fabric sits flat and feels secure around my legs.",
        "You can see the band stretch and return with each rep.",
        "Then the whole set packs into one small bag.",
        "肢体关节与拉伸方向符合人体运动，带子不穿过腿部，不夸张阻力等级",
    ),
    case(
        "运动出行", "汽车与摩托车", "汽车手机支架",
        "主搜：car phone mount dashboard adjustable；长尾：air vent phone holder, 360 degree rotation, one hand operation, installed view",
        "夹臂、底座、旋钮、球头、出风口夹和颜色",
        "停放状态的汽车内饰，仅演示安装与静态操作",
        "支架放在中控台旁，手持展示夹臂与底座，首帧产品清楚。",
        "把底座固定到指定位置并旋紧旋钮，再单手放入手机。",
        "轻推手机测试稳定性，转动球头切换竖屏和横屏角度。",
        "从驾驶视角展示不遮挡道路的安装位置，车辆保持静止。",
        "This mount puts my phone where I can see it without holding it.",
        "The base tightens down and the phone goes in with one hand.",
        "I can switch the viewing angle with a simple turn.",
        "Set it before the drive and keep your hands free.",
        "只在停车状态演示；底座必须真实固定，手机由夹臂承托，不遮挡驾驶视线",
    ),
    case(
        "兴趣收藏", "玩具", "磁力片拼装玩具",
        "主搜：magnetic building tiles set；长尾：3D construction blocks, transparent magnetic tiles, castle building set, all pieces flat lay",
        "磁力片形状、透明颜色、边框、磁点位置、数量和包装",
        "明亮亲子活动桌，成年人双手演示拼装",
        "全部磁力片按形状平铺，双手拿起两片吸合，首帧产品清楚。",
        "从底座开始逐层搭建一个小型立体房屋，关键连接近景可见。",
        "轻转完成的结构展示各面，再拆下一片证明真实磁吸连接。",
        "成品与剩余零件、包装同框，完成彩色hero俯拍。",
        "These magnetic tiles turn a few simple shapes into a new build.",
        "The pieces connect edge to edge as the structure grows.",
        "You can see the real magnetic connection up close.",
        "Build it, change it, and start again.",
        "每片必须由手移动并沿边缘吸合，不凭空生成完成建筑，不增加套装零件数量",
    ),
    case(
        "兴趣收藏", "爱好", "串珠手工包",
        "主搜：bracelet making kit beads for adults；长尾：friendship bracelet kit, bead organizer box, jewelry making supplies, finished bracelet",
        "珠子颜色与形状、分格盒、线材、工具、数量感和包装",
        "安静书桌，手作创作者第一视角",
        "打开分格收纳盒展示珠子与线材，成套产品在首帧清楚出现。",
        "左手固定线材，右手依次穿入三到四颗珠子，动作近景清晰。",
        "继续完成一小段图案并与参考成品并排，展示真实制作进度。",
        "把半成品、工具和收纳盒排成整洁hero俯拍。",
        "This kit makes it easy to start a small creative project.",
        "I can pick a color and thread each bead one at a time.",
        "The pattern builds gradually, which is the satisfying part.",
        "Everything goes back into the organizer when I am done.",
        "珠子只能由手逐颗移动并穿过线材，不瞬间生成完整手链，左右手任务明确",
    ),
    case(
        "兴趣收藏", "图书", "儿童互动翻翻书",
        "主搜：lift the flap board books for toddlers；长尾：interactive board book, sturdy pages, touch and feel book, page spread close up",
        "封面、书脊、开本、页数厚度、插画风格和翻页结构",
        "温暖阅读角，由成年人双手进行产品展示",
        "书本封面正对镜头放在阅读垫上，首帧完整清楚。",
        "双手打开书本，缓慢翻过两页并掀起一个真实翻翻页机关。",
        "微距展示厚纸板边缘、装订与插画细节，避免停留在可读正文。",
        "合上书本并与阅读玩偶同框，封面清楚、背景留出CTA。",
        "This interactive book makes story time feel hands-on.",
        "The sturdy pages are easy to turn and explore.",
        "Each flap adds a small moment of discovery.",
        "A simple pick for the reading corner.",
        "页面必须由手沿书脊真实翻动，翻翻页沿折线开启；不得改写或生成新的可读文字",
    ),
    case(
        "兴趣收藏", "杂志", "家居设计杂志",
        "主搜：interior design magazine print；长尾：home decor magazine, architecture magazine issue, coffee table magazine, cover and spread view",
        "封面版式、书脊、尺寸、页数厚度、纸张质感和原有标题",
        "客厅咖啡桌，慢节奏生活方式开箱",
        "杂志平放在咖啡桌中央，封面在首帧清楚可见。",
        "手部从右下角缓慢翻过三页，展示跨页图片与纸张光泽。",
        "把杂志靠在沙发边桌上，再切到书脊与厚度近景。",
        "杂志与咖啡杯、眼镜组成整洁hero画面，封面保持清楚。",
        "This issue is full of visual ideas for a slow coffee break.",
        "The print, paper, and full-page layouts feel great in hand.",
        "It also looks right at home on the coffee table.",
        "One issue to browse, keep, and come back to.",
        "页面顺序连续且由手翻动，不复制或改写正文，不生成错误可读标题",
    ),
    case(
        "兴趣收藏", "音频", "收藏版黑胶唱片",
        "主搜：vinyl record album limited edition；长尾：colored vinyl record, gatefold album, record sleeve collectible, turntable setup",
        "封套图案、唱片颜色、中心标签、尺寸、曲目侧标和包装",
        "音乐角与唱片机旁，收藏者开箱展示",
        "封套正面靠近镜头，随后抽出一小段唱片，产品首帧清楚。",
        "双手托住唱片边缘完全取出，微距展示彩胶与中心标签。",
        "唱片沿中孔放上唱盘并缓慢旋转，仅使用无歌词环境音乐。",
        "封套立在唱片机旁，旋转唱片与包装同框完成hero画面。",
        "This pressing looks as good on display as it does on the turntable.",
        "The color and sleeve details are the reason I picked it up.",
        "Here it is spinning in my listening corner.",
        "A collectible piece for vinyl fans.",
        "手只能接触唱片边缘与中心区域，唱片沿唱盘轴旋转；不生成受版权保护的歌词",
    ),
    case(
        "兴趣收藏", "收藏品", "盲盒潮玩公仔",
        "主搜：blind box collectible figures；长尾：mystery box vinyl figure, mini figurine series, sealed package, figure detail",
        "外盒图案、公仔造型、面部特征、配色、底座、配件和系列标识",
        "收藏展示架前，第一视角真实开盒",
        "未拆封盲盒放在展示架前，手指盒面系列图，首帧包装清楚。",
        "双手沿盒盖打开包装并取出内袋，再真实撕开封口。",
        "公仔从袋中取出并缓慢旋转，近景展示正侧背与底座。",
        "把公仔放入展示架，包装立在旁边形成hero画面。",
        "Let us open one and see which figure is inside.",
        "The reveal happens one layer at a time.",
        "Here are the sculpt and color details from every side.",
        "This one is going straight onto the display shelf.",
        "开盒顺序连续，公仔必须从真实包装中取出；不得中途改变款式、配色或配件数量",
    ),
]


EXPECTED = {
    "服饰时尚": ["男女装", "内衣", "儿童", "鞋靴", "箱包", "配件", "珠宝"],
    "美妆健康": ["美妆个护", "保健品"],
    "家居家装": ["家居日用", "布艺", "厨房用品", "家具", "建材", "五金"],
    "数码家电": ["手机与数码", "电脑办公", "家电"],
    "母婴宠物": ["母婴用品", "宠物用品"],
    "运动出行": ["运动与户外", "汽车与摩托车"],
    "兴趣收藏": ["玩具", "爱好", "图书", "杂志", "音频", "收藏品"],
}


COLORS = {
    "服饰时尚": "E8DDF8",
    "美妆健康": "FCE0EA",
    "家居家装": "DDF3E3",
    "数码家电": "DCEBFA",
    "母婴宠物": "FFF0C7",
    "运动出行": "FFE4C7",
    "兴趣收藏": "E9E2FA",
}


def validate_rows(rows):
    expected_pairs = [(industry, subcategory) for industry, subs in EXPECTED.items() for subcategory in subs]
    actual_pairs = [(row[0], row[1].split("｜", 1)[0]) for row in rows]
    assert len(rows) == 28, len(rows)
    assert actual_pairs == expected_pairs, (actual_pairs, expected_pairs)
    assert len({row[1] for row in rows}) == 28
    for industry, category, keywords, prompt in rows:
        assert keywords.startswith("主搜：") and "长尾：" in keywords, (industry, category)
        for token in ("15秒", "9:16", "TikTok Shop", "@图片1", "产品在首帧清楚出现", "英文口播"):
            assert token in prompt, (category, token)
        assert 450 <= len(prompt) <= 900, (category, len(prompt))


def build_workbook():
    validate_rows(ROWS)
    wb = Workbook()
    ws = wb.active
    ws.title = "细分品类案例库"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 75
    ws.freeze_panes = "C2"

    headers = ["行业", "品类", "产品关键词", "完整提示词"]
    ws.append(headers)
    for row in ROWS:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="102A43")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="243B53")
    thin = Side(style="thin", color="D9E2EC")
    section = Side(style="medium", color="9FB3C8")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=section)
    ws.row_dimensions[1].height = 28

    previous_industry = None
    for r in range(2, ws.max_row + 1):
        industry = ws.cell(r, 1).value
        fill = PatternFill("solid", fgColor=COLORS[industry])
        for c in range(1, 5):
            cell = ws.cell(r, c)
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        ws.cell(r, 1).fill = fill
        ws.cell(r, 2).fill = fill
        ws.cell(r, 1).font = Font(name="Microsoft YaHei", size=10, bold=True, color="102A43")
        ws.cell(r, 2).font = Font(name="Microsoft YaHei", size=10, bold=True, color="102A43")
        if previous_industry is not None and industry != previous_industry:
            for c in range(1, 5):
                ws.cell(r, c).border = Border(top=section, bottom=thin)
        previous_industry = industry
        ws.row_dimensions[r].height = 150

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 112
    ws.auto_filter.ref = f"A1:D{ws.max_row}"

    table = Table(displayName="RedBoxCoverageTable", ref=f"A1:D{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.print_title_rows = "1:1"

    wb.properties.title = "CreatiSignal Seedance 2.0 红框细分品类全覆盖 GTM 案例库"
    wb.properties.subject = "7个行业、28个红框细分品类的Amazon US关键词与视频生成提示词"
    wb.properties.creator = "CreatiSignal"
    wb.save(OUTPUT_FILE)


def render_preview():
    width = 2200
    header_h = 62
    row_h = 198
    col_widths = [190, 320, 530, width - 190 - 320 - 530]
    height = header_h + row_h * len(ROWS)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    font_paths = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    font_path = next((p for p in font_paths if p.exists()), None)
    header_font = ImageFont.truetype(str(font_path), 25) if font_path else ImageFont.load_default()
    body_font = ImageFont.truetype(str(font_path), 19) if font_path else ImageFont.load_default()
    bold_font = ImageFont.truetype(str(font_path), 20) if font_path else ImageFont.load_default()

    def wrap(text, font, max_width, max_lines):
        lines, current = [], ""
        for ch in str(text).replace("\n", " "):
            candidate = current + ch
            if draw.textlength(candidate, font=font) <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = ch
                if len(lines) >= max_lines:
                    break
        if len(lines) < max_lines and current:
            lines.append(current)
        return lines[:max_lines]

    x = 0
    for header, col_w in zip(["行业", "品类", "产品关键词", "完整提示词"], col_widths):
        draw.rectangle((x, 0, x + col_w, header_h), fill="#102A43")
        draw.text((x + 12, 17), header, font=header_font, fill="white")
        x += col_w

    for idx, row in enumerate(ROWS):
        y = header_h + idx * row_h
        x = 0
        for col, (value, col_w) in enumerate(zip(row, col_widths)):
            fill = "#FFFFFF"
            if col < 2:
                fill = "#" + COLORS[row[0]]
            draw.rectangle((x, y, x + col_w, y + row_h), fill=fill, outline="#D9E2EC", width=1)
            font = bold_font if col < 2 else body_font
            max_lines = 7 if col == 3 else 6
            lines = wrap(value, font, col_w - 24, max_lines)
            for line_idx, line in enumerate(lines):
                draw.text((x + 12, y + 10 + line_idx * 27), line, font=font, fill="#243B53")
            x += col_w
    image.save(PREVIEW_FILE, quality=95)


def verify():
    with ZipFile(OUTPUT_FILE) as zf:
        assert zf.testzip() is None
    wb = load_workbook(OUTPUT_FILE, read_only=False, data_only=False)
    ws = wb["细分品类案例库"]
    assert [ws.cell(1, c).value for c in range(1, 5)] == ["行业", "品类", "产品关键词", "完整提示词"]
    assert ws.max_row == 29 and ws.max_column == 4
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == "A1:D29"
    assert any(table.ref == "A1:D29" for table in ws.tables.values())
    values = list(ws.iter_rows(min_row=2, max_col=4, values_only=True))
    validate_rows(values)
    counts = Counter(row[0] for row in values)
    assert counts == Counter({industry: len(subs) for industry, subs in EXPECTED.items()})
    assert not [cell.coordinate for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    wb.close()
    assert PREVIEW_FILE.exists() and PREVIEW_FILE.stat().st_size > 100_000
    print(f"OUTPUT={OUTPUT_FILE}")
    print(f"PREVIEW={PREVIEW_FILE}")
    print(f"ROWS={len(values)}; INDUSTRIES={len(counts)}; COUNTS={dict(counts)}")
    print(f"BYTES={OUTPUT_FILE.stat().st_size}")


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_workbook()
    render_preview()
    verify()
