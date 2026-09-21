from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from PIL import Image, ImageDraw, ImageFont

import build_representative_products as base

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "CreatiSignal_28品类_代表商品与TikTok爆款检索_20260902.xlsx"
OUTPUT = ROOT / "CreatiSignal_28品类_TikTok搜A_Amazon搜B_20260903.xlsx"
QA = ROOT / "ab_keywords_qa"
MAPPINGS = []


def item(category, b_type, queries, match, boundary, replacements=()):
    assert len(queries) == 3
    MAPPINGS.append(dict(category=category, b_type=b_type, queries=queries, match=match,
                         boundary=boundary, replacements=replacements))


item("男女装", "高腰弹力阔腿牛仔裤",
     ["women high waisted stretch wide leg jeans", "women wide leg stretch denim pants", "high waist wide leg jeans with pockets"],
     "选非HALARA品牌/款式；保留高腰、阔腿、弹力牛仔面料和可展示口袋，颜色可不同。",
     "复用镜前试穿、腰头轻拉和侧面走动；不把A的洗水色、走线或显瘦承诺套给B。")

item("内衣", "可调肩带无缝短裤式塑身连体衣",
     ["seamless shapewear bodysuit shorts adjustable straps", "women mid thigh shapewear bodysuit", "seamless shorts body shaper under dress"],
     "选非SHAPERX品牌/款式；保留连体、可调肩带与短裤式裤腿，不选普通吊带或丁字裤式。",
     "复用面料近景与连衣裙叠穿；版型、肩带和裁片按B展示，不迁移身材改变效果。")

item("儿童", "双向拉链印花连体睡衣",
     ["baby bamboo viscose pajamas two way zipper", "baby zipper sleeper fold over feet cuffs", "printed baby romper double zipper pajamas"],
     "选非Little Sleepies品牌；保留连体、双向拉链和可翻袖脚口结构，印花与颜色可以不同。",
     "复用平铺、拉链和袖脚口展示；不复制Sunshine太阳图案，不虚构B不存在的翻折结构。",
     [("A sunny print, a two-way zipper, and little details worth a closer look.", "A closer look at the fabric, zipper, and little everyday details.")])

item("鞋靴", "硬挺后跟免手穿休闲步行鞋",
     ["hands free slip in walking shoes", "hands free walking sneakers reinforced heel", "easy step in slip on walking shoes"],
     "选非Skechers品牌；必须在商品说明中确认可免手穿/踏入，普通软后跟一脚蹬不直接替代。",
     "复用低机位穿鞋和走动；B没有免手穿结构就只展示已穿好后走动，不强行套用A的功能。")

item("箱包", "绗缝软质电脑通勤托特包",
     ["quilted puffer tote bag laptop compartment", "padded work tote bag zipper pockets", "quilted laptop tote bag with bottle pocket"],
     "选非BAGSMART品牌；保留软质绗缝、双提手、拉链和电脑/水瓶收纳区，容量按B实际规格。",
     "复用pack with me和肩背展示；分区、装入物品大小与数量按B调整。")

item("配件", "大号弹簧鲨鱼夹",
     ["large hair claw clips thick hair", "large spring claw hair clip half up", "tortoise shell hair claw clips large"],
     "选非KITSCH品牌；保留夹齿和捏压弹簧开合，颜色和轮廓可不同；单视频只用一个B款。",
     "复用发夹开合、半扎发和后侧展示；不复制A的纹理与夹齿数量。")

item("珠宝", "金色双圈耳环",
     ["gold double hoop earrings women", "chunky double hoop earrings", "gold tone double hoop earrings jewelry"],
     "选非PAVOI品牌；保留双圈造型和耳边佩戴场景，尺寸、金属成分与扣件以B为准。",
     "复用托盘展示、金属反光与侧脸佩戴；不照搬A的纯度、克重和防过敏属性。")

item("美妆个护", "带刷头的轻透有色唇油",
     ["tinted lip oil glossy large applicator", "sheer tinted lip oil wand", "glossy lip oil clear tube applicator"],
     "选非e.l.f.品牌；优先透明管身、刷头涂抹、轻透有色唇油，不用哑光唇釉或固体唇膏替代。",
     "复用开盖、手背试色和轻涂；管身、刷头、色号与妆效只跟随B商品图。")

item("保健品", "瓶装成人苹果醋营养软糖",
     ["apple cider vinegar gummies adults", "adult apple cider vinegar gummy supplement", "apple cider vinegar gummy vitamins bottle"],
     "选非Goli品牌；保持成人营养补充剂、瓶装软糖剂型，瓶色、配方和单粒外观可不同。",
     "只复用开瓶和剂型近景；不得复制A的红瓶、配方、剂量或功效。",
     [("红瓶", "B商品瓶身")])

item("家居日用", "双面洗碗清洁海绵",
     ["non scratch dish sponge dual sided scrubber", "double sided kitchen cleaning sponges", "dishwashing sponge non scratch scrub pad"],
     "选非Scrub Daddy品牌的洗碗海绵；长方形、椭圆形均可，需有可实际擦洗的接触面。",
     "复用擦洗—冲水—沥水结构；不把A的笑脸开孔、圆形外观或温感材质迁移到B。",
     [("笑脸清楚可见", "B海绵外形清楚可见"), ("笑脸海绵", "B清洁海绵")])

item("布艺", "双片打孔式遮光窗帘",
     ["blackout curtains grommet 2 panels lined", "bedroom blackout window curtains pair", "thermal blackout curtains grommet top"],
     "选非NICETOWN品牌；优先双片、顶部金属环和遮光背衬，长度需匹配拍摄窗户。",
     "复用面料特写和固定机位开合；只呈现B真实光线变化，不复制A的遮光比例。")

item("厨房用品", "带集料盒的手压蔬菜切丁器",
     ["vegetable chopper press dicer with container", "manual vegetable dicer hinged lid container", "onion chopper grid blade food container"],
     "选非Fullstar品牌；必须为铰链压盖+刀网+下方集料盒结构，不选拉绳绞碎机或电动料理机。",
     "复用放食材—按压—倒出切丁；刀网形状、配色与切丁大小按B实际结构。")

item("家具", "无扶手无轮宽座盘腿办公椅",
     ["criss cross desk chair no wheels fabric", "armless wide seat cross legged chair", "fabric swivel criss cross office chair"],
     "选非PUKAMI品牌；保留宽座、布艺、无扶手、无轮和稳固底座，核验B允许的使用姿态。",
     "复用座垫近景和desk setup；不照搬A的底座、尺寸与人体工学承诺。")

item("建材", "自粘地铁砖纹防溅墙贴",
     ["peel and stick subway tile backsplash", "self adhesive kitchen backsplash subway tiles", "thick peel stick backsplash wall tiles"],
     "选非Art3d品牌；保留自粘背胶、地铁砖纹、单片铺贴形式，不选需要水泥的真瓷砖。",
     "复用揭背纸—对齐—压平；边缘、厚度、砖缝和适用墙面按B确定。")

item("五金", "充电式小型电动螺丝刀套装",
     ["cordless electric screwdriver rechargeable kit", "compact electric screwdriver bits case", "USB rechargeable electric screwdriver set"],
     "选非HOTO品牌；优先一手握持的小型电动螺丝刀、可换批头和收纳盒，不选冲击钻。",
     "复用开箱和单颗螺丝装配；不照搬A的环形灯、扭矩、批头数量或操作按钮。")

item("手机与数码", "带折叠支架的磁吸充电宝",
     ["magnetic wireless power bank foldable stand", "magnetic portable charger built in kickstand", "wireless magnetic battery pack stand"],
     "选非Anker品牌；保留磁吸面和可展开支架，并确认手机/壳兼容；显示屏不是必要条件。",
     "复用贴合—展开支架—桌面支撑；电量、功率、屏幕和续航不得沿用A的数据。",
     [("屏幕数值后期处理", "如B有屏幕，准确数值后期处理；没有屏幕则不添加")])

item("电脑办公", "带旋钮的75%机械键盘",
     ["75 percent mechanical keyboard knob wireless", "75% hot swappable keyboard rotary knob", "compact mechanical keyboard knob desk setup"],
     "选非AULA/EPOMAKER品牌；保留75%布局与旋钮，配色和轴体可不同，核验B真实键位。",
     "复用桌搭、敲击和旋钮近景；B键位、轴体与声音不得套用A的版本。")

item("家电", "带透明吸头的便携布艺清洗机",
     ["portable carpet upholstery cleaner spot cleaner", "upholstery cleaner machine hose clear nozzle", "compact spot cleaner couch carpet machine"],
     "选非BISSELL品牌；需有水箱、软管、喷吸清洗和透明吸头，不能用普通干式吸尘器替代。",
     "复用局部吸洗和清洁轨迹；机身、管线、用水和实测效果按B，不能保留A的绿机外形。")

item("母婴用品", "透明盖台面式奶瓶清洗机",
     ["baby bottle washer sterilizer dryer machine", "automatic countertop baby bottle washer", "baby bottle washing machine clear lid"],
     "选非Momcozy品牌；必须具备清洗功能和装载支架，优先透明盖，不选只有消毒烘干的设备。",
     "复用奶瓶装载—合盖—启动；支架位置、装载数量和操作步骤必须按B说明。")

item("宠物用品", "带饮水盘的循环式宠物饮水机",
     ["automatic cat water fountain drinking tray", "cat water fountain circulating water bowl", "pet drinking fountain detachable tank"],
     "选非PETLIBRO品牌；保留水箱、出水口和饮水盘即可，B可有线或电池供电。",
     "复用宠物靠近、低幅水流和舔水；按B出水模式拍摄，不擅自加入A的感应、APP或RFID功能。",
     [("电池模式只在接近触发的合理时间内出水", "按B实际供电和出水模式展示，不额外添加感应功能")])

item("运动与户外", "附阻力带的双踏板液压迷你踏步机",
     ["mini stepper exercise machine resistance bands", "hydraulic mini stair stepper home workout", "compact stepper two pedals resistance bands"],
     "选非Sunny Health & Fitness品牌；保留双踏板液压机构和阻力带，不选跑步机或平衡板。",
     "复用慢速交替踏步和空间展示；B的踏板轨迹、行程与连接点独立核验。")

item("汽车与摩托车", "带气管的无线轮胎充气泵",
     ["cordless tire inflator portable air compressor", "rechargeable tire pump car motorcycle", "portable digital tire inflator hose"],
     "选非Fanttik品牌；保留充电式主机、软管和轮胎气门连接方式，确认适用汽车/摩托车。",
     "复用接管—启动—整机展示；不能沿用A的胎压数值、充气时长和按键布局。")

item("玩具", "彩色透明几何磁力片套装",
     ["magnetic building tiles translucent squares triangles", "magnetic tiles building blocks starter set", "magnetic construction tiles house building"],
     "选非MAGNA-TILES品牌；保留透明正方形/三角形薄片与边缘磁吸，不选磁力棒或普通积木。",
     "复用简单小屋搭建；数量、尺寸和配色按B，不沿用A的32片数量或包装。")

item("爱好", "企鹅新手钩织材料包",
     ["penguin crochet kit for beginners", "beginner penguin crochet kit yarn hook", "penguin amigurumi crochet starter kit"],
     "选非The Woobles品牌；优先包含钩针、毛线与企鹅样品图的材料包，不选只有成品的毛绒玩具。",
     "复用材料开箱—一针过程—完成样品；B没有预起针服务时，由拍摄者提前准备织片，不暗示为随包赠送。",
     [("预起针小片", "拍摄者提前准备的一小片织片")])

item("图书", "个人成长主题英文纸质精装书",
     ["personal growth self help book hardcover", "self improvement book hardcover English", "mindset personal development hardcover book"],
     "选择不同书名/作者的个人成长纸质书；不选The Let Them Theory，先核验B的真实主题与版本。",
     "复用举封面、翻页、书签和阅读氛围；书名、作者、正文与观点必须换成B，不沿用A书中观点。",
     [("Today’s reading pick: The Let Them Theory. A quiet moment, a bookmark, and a new chapter.", "A quiet reading moment, a bookmark, and a new chapter.")])

item("杂志", "时尚类纸质单期杂志",
     ["fashion magazine single issue print", "style fashion magazine current issue", "fashion photography magazine print edition"],
     "选择不同于VOGUE的时尚刊物；购买/取图前固定B的刊名、国家版、期号和封面。",
     "复用封面展示、翻页和桌面氛围；不保留A的刊名、2025年9月号、封面人物或内页。",
     [("不换封面人物、国家版或期号", "B的封面人物、国家版和期号在视频内保持一致")])

item("音频", "其他歌手或专辑的彩色流行音乐黑胶LP",
     ["colored vinyl LP pop album", "pop music album colored vinyl record", "marbled vinyl LP album record"],
     "选择不同歌手或不同专辑的实体黑胶，不选Taylor Swift Midnights；固定B盘色、封套和压片版本。",
     "复用取盘、盘面展示和唱机旋转；封套、盘色、歌手与曲目全部换成B，不搬用A音乐。",
     [("同一张蓝色唱片", "B商品的同一张唱片"),
      ("不使用Taylor Swift歌曲、歌词、声音或仿唱", "不使用A或B专辑中未获授权的音乐、歌词、歌手声音或仿唱")])

item("收藏品", "不同原创角色的搪胶毛绒挂件盲盒",
     ["vinyl plush pendant blind box", "collectible plush bag charm blind box", "designer vinyl face plush keychain mystery box"],
     "选不同于THE MONSTERS/Labubu的正版角色或系列；保留盲盒、毛绒挂件形态，不选A仿品或换色盗版。",
     "复用开盒—细节—包挂结构；角色、牙齿、毛色、吊牌和包装完全跟随B，不能把B变成Labubu。",
     [("A closer look at this Big into Energy character, from the box to the bag.", "A closer look at this little character, from the box to the bag.")])


def updated_prompt(old, mapping):
    # Retain useful shot structures, but remove the A-specific product-lock paragraph.
    body = "场景与主体：" + old.split("场景与主体：", 1)[1]
    for a, b in mapping["replacements"]:
        assert a in body, (mapping["category"], a)
        body = body.replace(a, b)
    prefix = (
        f"生成15秒、9:16竖屏TikTok原生商品展示视频，展示Amazon选定的B商品：{mapping['b_type']}。"
        "用已上传的@图片1 @图片2 @图片3锁定B的品牌、外观、颜色、标签、结构和真实配件。"
        "如提供已授权@视频1，它仅是A商品的创意参考；借鉴钩子、节奏、构图及适用于B的动作，不复制A商品、包装、人物身份、声音、水印或评价。"
        "B不具备的功能或结构不展示。\n"
    )
    return prefix + body


GUIDE = [
    ("版本与范围", "2026-09-03｜覆盖7个行业、28个细分品类，原顺序与六列表头不变。此次从‘两边搜同款’改为‘TikTok搜A视频，Amazon搜同类B商品’。原版文件保留。"),
    ("A与B分别是什么", "主表‘具体商品’和‘TikTok爆款关键词’指A：用来寻找参考视频。‘Amazon商品关键词’用于寻找B：与A属于同一细分品类、核心用途和演示动作相近，但选择不同品牌或不同商品。‘完整提示词’用于生成B，而不是重新生成A。"),
    ("本次B的确定程度", "本表提供B的类型、3条无A品牌的英文搜索词和选款条件，尚未替你锁定B的品牌、链接或ASIN。搜索结果也可能包含A，需要按每行条件选择不同的B。词语相似不代表每个搜索结果都适合复刻。"),
    ("TikTok怎么搜", "沿用A的4条独立英文查询词，一次搜一条。先找同一A商品的视频，再查看钩子、机位、手部动作和商品展示是否清晰；搜索词不是已核验的爆款榜。"),
    ("Amazon怎么搜", "一次复制主表Amazon单元格中的一行英文关键词，‘B方向’及‘选款条件’是人工筛选说明，不要整格复制进搜索框。先看功能和形态，再从结果中排除A，确认B的品牌、型号及商品页；不要依赖未经验证的排除搜索语法。"),
    ("怎么判断A和B相似", "同时满足：同一细分品类；核心用途相近；演示动作相近；结构能支持原分镜。颜色和品牌可以不同。手压切丁器不能换成拉绳绞碎机；喷吸布艺清洗机不能换成干式吸尘器；只有消毒功能的设备不能套用奶瓶清洗视频。"),
    ("具体例子", "TikTok A：e.l.f.唇油 → Amazon B：tinted lip oil glossy large applicator，选其他品牌唇油，复用试色/上唇结构。TikTok A：Scrub Daddy → Amazon B：non scratch dish sponge dual sided scrubber，选其他品牌海绵，复用擦洗过程，不保留笑脸造型。"),
    ("参考素材如何上传", "@视频1只放有权使用的A参考视频；@图片1/2/3全部使用B同一SKU的正面、侧面和关键细节图，不混入A的商品图。引用标签替换为界面中的真实名称；没有的图片/视频引用删除。B一经选定，生成过程中不换品牌、配色或款式。"),
    ("提示词已同步调整", "A仅提供创意结构，B图片决定商品身份。移除了对A品牌、包装颜色、图案、书名、专辑和盲盒角色的锁定；沿用能适配B的短分镜。B缺少某功能时删掉该镜头，不通过提示词强行生成。真实字幕、标签或封面细字必要时后期合成。"),
    ("10–15秒使用方式", "默认15秒四段结构。缩为10秒时保留一个主要动作，采用0–2秒开场、2–6秒演示、6–10秒收尾，删去一段细节。先选B实物/图片，再调整动作与口播；不要同时复刻A的完整长视频。"),
    ("爆款筛选与广告入口", "在同品类、同市场和相近发布时间内比较候选，记录链接、账号、日期、播放/互动和购买意图。高播放不等于高转化。Top Ads官方检索说明（沿用旧版来源，本轮未重新核验）：https://ads.tiktok.com/help/article/how-to-use-the-top-ads-dashboard?lang=en"),
    ("权益与品牌边界", "参考公开视频不等于拥有原片、人物、声音、音乐或品牌素材的商用授权。只借鉴结构，不伪装品牌合作或真实用户见证。B选独立商品或正版不同角色，不通过‘换品牌’制作A的仿品。GTM公开展示前核验相关权益与披露要求。"),
    ("功效和模拟演示", "A的功效、数据、容量、速度、除菌率、评论及销量不能迁移到B。保健品不生成医疗/减重承诺；儿童母婴按B说明使用；清洁/遮光/键盘声音等AI镜头不是实测证据。B的参数和实际表现需独立确认。"),
    ("核验与来源边界", "此次为关键词与提示词逻辑修订，未新增实时选品或Amazon库存核验。‘A-B匹配说明’的链接仅是2026-09-02旧版用于确认A名称/结构的来源，未重新检查时效，也不是B的链接或TikTok爆款视频证据。"),
    ("最终检查", "下载B图片前核对品牌、型号、颜色、规格/色号和结构；杂志固定刊名与期号，图书固定书名与版本，黑胶固定专辑/压片版本，盲盒固定角色/系列。出片逐帧检查B一致性，必要时用实拍商品镜头结合AI场景，并标注AI生成/模拟演示。"),
]


def fit_changed_rows(ws):
    widths = [ws.column_dimensions[ws.cell(1, c).column_letter].width for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        lines = max(len(base.wrapped(ws.cell(r, c).value or "", base.MEASURE_FONT, widths[c - 1] * 7 - 22))
                    for c in range(1, ws.max_column + 1))
        height = max(ws.row_dimensions[r].height or 44, math.ceil((lines * 22 + 22) * .75))
        assert height <= 409, (ws.title, r, height)
        ws.row_dimensions[r].height = height


def build():
    original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    original = load_workbook(SOURCE)
    wb = load_workbook(SOURCE)
    ws = wb["28品类案例库"]
    assert [m["category"] for m in MAPPINGS] == base.EXPECTED
    for row, m in enumerate(MAPPINGS, 2):
        assert ws.cell(row, 2).value == m["category"]
        ws.cell(row, 6).value = (
            "B方向：" + m["b_type"] + "\n"
            + "\n".join(m["queries"]) + "\n选款：" + m["match"]
        )
        ws.cell(row, 5).value = updated_prompt(ws.cell(row, 5).value, m)
    ws["C1"].comment = Comment("本列为TikTok参考商品A；Amazon列用于另找同类B商品，提示词生成B。", "ice.li")
    ws["D1"].comment = Comment("沿用A商品的TikTok视频检索词，不是B的搜索词或已核验爆款清单。", "ice.li")
    ws["E1"].comment = Comment("@视频1=A的视频结构参考；@图片1/2/3=B的同一SKU图片。", "ice.li")
    ws["F1"].comment = Comment("仅复制各行英文查询词。B方向和选款说明用于筛选，搜索后排除A品牌/商品。", "ice.li")
    fit_changed_rows(ws)

    guide = wb["使用说明"]
    assert guide.max_row == len(GUIDE) + 1
    for r, values in enumerate(GUIDE, 2):
        guide.cell(r, 1).value, guide.cell(r, 2).value = values
    fit_changed_rows(guide)

    matching = wb["选品理由与来源"]
    matching.title = "A-B匹配说明"
    labels = ["序号", "品类", "A参考商品（TikTok）", "B选款条件（Amazon）", "可复用结构与迁移边界", "A商品信息来源（旧版核验）"]
    for c, label in enumerate(labels, 1):
        matching.cell(1, c).value = label
    for c, table_column in enumerate(matching.tables["ProductSources"].tableColumns):
        table_column.name = labels[c]
    for row, m in enumerate(MAPPINGS, 2):
        matching.cell(row, 3).value = ws.cell(row, 3).value
        matching.cell(row, 4).value = m["b_type"] + "。" + m["match"]
        matching.cell(row, 5).value = m["boundary"]
    fit_changed_rows(matching)
    for sheet in wb:
        sheet.oddFooter.center.text = "CreatiSignal | A video -> B product | 2026-09-03 | &P / &N"
    wb.properties.title = "CreatiSignal 28品类：TikTok搜A，Amazon搜B"
    wb.properties.subject = "跨商品创意参考：A的视频结构与B的真实商品图片分离"
    wb.save(OUTPUT)

    # Read the saved workbook, compare immutable columns, and validate every modified row.
    checked = load_workbook(OUTPUT)
    result_sheet = checked["28品类案例库"]
    assert [c.value for c in result_sheet[1]] == base.HEADERS
    assert (result_sheet.max_row, result_sheet.max_column) == (29, 6)
    assert result_sheet.freeze_panes == "D2" and result_sheet.auto_filter.ref == "A1:F29"
    brand_tokens = ["HALARA", "SHAPERX", "Little Sleepies", "Skechers", "BAGSMART", "KITSCH", "PAVOI", "e.l.f.", "Goli", "Scrub Daddy", "NICETOWN", "Fullstar", "PUKAMI", "Art3d", "HOTO", "Anker", "AULA", "EPOMAKER", "BISSELL", "Momcozy", "PETLIBRO", "Sunny Health", "Fanttik", "MAGNA-TILES", "The Woobles", "The Let Them Theory", "VOGUE", "Taylor Swift", "Midnights", "Moonstone", "POP MART", "Big into Energy", "Labubu"]
    def contains_brand(text):
        return any(re.search(r"(?<![A-Za-z0-9])" + re.escape(t) + r"(?![A-Za-z0-9])", text, re.I) for t in brand_tokens)
    for r, m in enumerate(MAPPINGS, 2):
        for c in range(1, 5):
            assert result_sheet.cell(r, c).value == original["28品类案例库"].cell(r, c).value
        for c in range(1, 7):
            assert result_sheet.cell(r, c).has_style
            assert result_sheet.cell(r, c)._style == original["28品类案例库"].cell(r, c)._style
        prompt = result_sheet.cell(r, 5).value
        assert all(t in prompt for t in ["B商品", "@图片1", "@视频1", "A商品", "15秒", "9:16", "11–15秒"])
        assert not contains_brand(prompt), (r, "A brand in B prompt")
        assert not contains_brand(" ".join(m["queries"])), (r, "A brand in query")
        assert result_sheet.cell(r, 5).value != original["28品类案例库"].cell(r, 5).value
        assert result_sheet.cell(r, 6).value != original["28品类案例库"].cell(r, 6).value
    for s in checked:
        for row in s:
            for c in row:
                assert c.data_type not in ("f", "e")
    with ZipFile(OUTPUT) as z:
        assert z.testzip() is None
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == original_hash

    QA.mkdir(exist_ok=True)
    base.QA = QA
    previews = []
    for s, step in [(result_sheet, 3), (checked["使用说明"], 7), (checked["A-B匹配说明"], 7)]:
        previews.extend(base.render_sheet(s, step))
    width, height = 640, 520
    index = Image.new("RGB", (width * 3, height * math.ceil(len(previews) / 3)), "#D9E2EC")
    draw = ImageDraw.Draw(index)
    font = ImageFont.truetype(str(base.FONT_PATH), 18)
    for i, p in enumerate(previews):
        im = Image.open(p)
        im.thumbnail((width - 18, height - 36))
        x, y = i % 3 * width + 9, i // 3 * height + 28
        index.paste(im, (x, y))
        draw.text((x, y - 24), p.stem, font=font, fill="#102A43")
    index.save(QA / "all_sheets.png")
    report = {"output": str(OUTPUT), "rows": 28, "independent_amazon_queries": 84,
              "headers_and_A_columns": "UNCHANGED", "A_brands_in_B_prompts_or_search_queries": 0,
              "original_file": "UNCHANGED", "sheets": checked.sheetnames, "rendered_pages": len(previews),
              "max_row_height": max(result_sheet.row_dimensions[r].height for r in range(2, 30)),
              "checks": "PASS: content, exact order, original styles, filters, ZIP, all cell-fit checks"}
    (QA / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
