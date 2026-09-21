from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Cursor\creatisignal-app-new")
OUT_DIR = ROOT / "deliverables" / "ai-friendly-inspiration"
OUT_PATH = OUT_DIR / "灵感发现_AI友好型短视频素材筛选_PRD_V1.0.docx"
PROTOTYPE_PATH = OUT_DIR / "灵感发现_AI友好素材筛选_产品原型.png"

SKILL_ROOT = Path(
    r"C:\Users\ice.li\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.826.12353\skills\documents"
)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_TEXT = "8B8F98"
LINE = "D7DEE8"
HEADER_FILL = "E8EEF5"
ROW_FILL = "F8FAFC"
CALLOUT_FILL = "F4F6F9"
LIME = "A8FF1A"
LIME_SOFT = "F2FFD9"
GREEN = "147D46"
CAUTION = "7A5A00"
RISK = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "17181C"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 90, "bottom": 90, "start": 120, "end": 120}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 name: str = FONT_LATIN) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size: float, color: str, bold: bool = False) -> None:
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = tc_borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            tc_borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:color"), color)


def set_cell_margins(cell, margins: dict[str, int] = CELL_MARGINS) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != PAGE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {PAGE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_num_definition(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 1
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 11, BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    set_style_font(title, 27, BLACK, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, 13.5, MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.keep_with_next = True

    h1 = styles["Heading 1"]
    set_style_font(h1, 16, BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    set_style_font(h2, 13, BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    set_style_font(h3, 12, DARK_BLUE, bold=True)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.keep_with_next = True

    if "Table Text" not in styles:
        table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = styles["Table Text"]
    set_style_font(table_text, 9.2, BLACK)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.15

    if "Table Header" not in styles:
        table_header = styles.add_style("Table Header", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_header = styles["Table Header"]
    set_style_font(table_header, 9.2, INK, bold=True)
    table_header.paragraph_format.space_before = Pt(0)
    table_header.paragraph_format.space_after = Pt(0)
    table_header.paragraph_format.line_spacing = 1.10

    if "Caption Text" not in styles:
        caption = styles.add_style("Caption Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption Text"]
    set_style_font(caption, 9, MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, 9, MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("CreatiSignal  |  灵感发现 PRD")
    set_run_font(run, 9, MUTED, bold=True)
    footer_p = section.footer.paragraphs[0]
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    add_page_field(footer_p)


def add_body(doc: Document, text: str, bold: bool = False, color: str = BLACK,
             italic: bool = False, after: float | None = None) -> None:
    p = doc.add_paragraph(style="Normal")
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, 11, color, bold=bold, italic=italic)


def add_labeled(doc: Document, label: str, text: str, label_color: str = INK,
                after: float = 5) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, 11, label_color, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, 11, BLACK)


def add_bullet(doc: Document, text: str, bullet_num_id: int, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="Normal")
    apply_numbering(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 11, INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, 11, BLACK)
    else:
        run = p.add_run(text)
        set_run_font(run, 11, BLACK)


def add_callout(doc: Document, label: str, text: str, fill: str = CALLOUT_FILL,
                accent: str = BLUE) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "14")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, 10.5, accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, 10.5, BLACK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int],
              zebra: bool = True, keep_rows: bool = True) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    mark_repeat_header(hdr)
    keep_row_together(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        set_run_font(run, 9.2, INK, bold=True)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        if keep_rows:
            keep_row_together(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            if zebra and row_idx % 2 == 1:
                set_cell_shading(cell, ROW_FILL)
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, 9.2, BLACK)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(6)


def add_prototype(doc: Document) -> None:
    if not PROTOTYPE_PATH.exists():
        raise FileNotFoundError(PROTOTYPE_PATH)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    shape = run.add_picture(str(PROTOTYPE_PATH), width=Inches(6.5))
    shape._inline.docPr.set("descr", "灵感发现页面的 AI 友好型短视频素材筛选高保真产品原型")
    caption = doc.add_paragraph(style="Caption Text")
    caption.add_run("图 1  开发完成态原型：AI友好型（≤15s）+ 近7天 + 美国 + 投放中，按播放量排序")


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    configure_header_footer(doc)
    bullet_num_id = add_num_definition(doc, "bullet")

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    props = doc.core_properties
    props.title = "灵感发现｜AI友好型短视频素材筛选 PRD"
    props.subject = "CreatiSignal 灵感发现筛选、排序、搜索、数据与验收需求"
    props.author = "CreatiSignal Product"
    props.keywords = "灵感发现, AI友好素材, 15秒, 时效性, 素材筛选, PRD"
    props.comments = "V1.0 engineering-ready product requirements"

    # Opening masthead (memo_masthead pattern, without decorative bottom border).
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(5)
    kr = kicker.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_run_font(kr, 10, GREEN, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("灵感发现｜AI友好型短视频素材筛选")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("以 15 秒以内与近 7/14 天为核心，完成可筛选、可排序、可追溯的素材发现能力")

    metadata = [
        ("版本", "V1.0"),
        ("日期", "2026-08-31"),
        ("状态", "可进入技术方案评审"),
        ("适用页面", "/discover/inspiration（灵感发现）"),
        ("评审角色", "产品、设计、数据采集、后端、前端、QA"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        lr = p.add_run(f"{label}：")
        set_run_font(lr, 10.5, BLACK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, 10.5, BLACK)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        "核心决策",
        "P0 将“AI友好型素材”定义为时长已知且 ≤15.000 秒的视频素材；这是可计算的筛选标签，不等同于对 AI 复刻成功率或稳定批量交付能力的保证。",
        fill=LIME_SOFT,
        accent=GREEN,
    )
    add_body(
        doc,
        "本需求在现有“灵感发现”页面内增强，不新增一级导航。默认优先帮助用户看到近 7 天的短视频；当近 7 天无结果时，不自动混入更老素材，而是明确提示切换至近 14 天。",
    )
    add_bullet(doc, "交付范围：筛选、排序、关键词搜索、结果卡片信息、服务端查询、数据新鲜度与埋点。", bullet_num_id)
    add_bullet(doc, "默认视图：AI友好型开启、近7天、沿用当前平台默认项；国家与投放状态默认“全部”。", bullet_num_id)
    add_bullet(doc, "关键依赖：CTR、广告消耗预估与投放状态必须有真实来源和更新时间；不可用时展示“-”并在排序中置底，禁止伪造。", bullet_num_id)

    doc.add_page_break()

    doc.add_heading("背景与问题", level=1)
    add_body(
        doc,
        "AI 视频生成在 10–15 秒区间更容易控制连续性、人物/商品一致性与镜头稳定性。客户当前主要生成 10–15 秒短视频，或复刻同等长度的优秀素材。现有素材池若不能同时控制时长与发布时间，用户仍需逐条打开判断，难以形成稳定的“发现—筛选—复刻”工作流。",
    )
    add_labeled(doc, "用户痛点", "素材数量大，但“短、近期、可投、表现好”的交集无法一次筛出；同一素材的时长、发布时间、投放状态、互动指标与商业指标也缺少统一呈现。")
    add_labeled(doc, "业务机会", "把 AI 生产约束前置到灵感发现，让用户先找到更适合短视频生成/复刻的候选素材，再进入一键复刻，降低无效浏览和抽卡成本。")

    doc.add_heading("现状基线", level=2)
    add_bullet(doc, "截图基线：页面已有品类、媒体平台、国家地区、推荐/互动/点击率/最新、表现筛选与关键词入口，结果以视频瀑布流展示。", bullet_num_id)
    add_bullet(doc, "仓库基线：src/components/materials/materials-content.tsx 当前使用本地 mock，页面状态只有品类与搜索，卡片主数据仅 likes/comments/shares，现有筛选控件尚未连接统一查询。", bullet_num_id)
    add_bullet(doc, "邻近可复用能力：src/lib/replicate/source-filters.ts 已有 platform/country/status/likes/saves/spend/startDate 等模拟字段和过滤排序语义，但不能直接视为生产数据合同。", bullet_num_id)
    add_callout(doc, "工程判断", "本需求不是单纯补 UI 控件。要达到可用状态，必须同步完成标准化数据模型、采集/更新策略、服务端过滤排序与字段覆盖率治理。")

    doc.add_heading("目标与边界", level=1)
    doc.add_heading("产品目标", level=2)
    for item in [
        "用户可一键筛出时长 ≤15 秒的视频素材，并能看到精确时长。",
        "用户可优先查看近 7 天、近 14 天素材，时间边界可验证且不静默扩大。",
        "用户可按品类、投放国家、媒体平台、投放状态筛选，并进行关键词搜索。",
        "用户可按播放、点赞、收藏、评论、点击率、广告消耗预估进行稳定排序。",
        "所有指标都带数据时间语义；缺失或过期不会被误当成 0 或高质量结果。",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("非目标（P0 不做）", level=2)
    for item in [
        "不构建基于镜头复杂度、人物动作、商品一致性风险的“AI 复刻成功率评分”。",
        "不承诺任一素材能被特定模型稳定复刻；≤15 秒只代表时长友好，不代表生产结果友好。",
        "不做视频内容语义检索、OCR/ASR 关键词召回；P0 搜索结构化文本字段。",
        "不改变素材详情页和一键复刻主流程，不新增一级页面或新的素材资产体系。",
        "不在数据源缺失时推算并展示未经验证的 CTR、消耗或投放状态。",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("核心定义与业务规则", level=1)
    add_table(
        doc,
        ["概念", "P0 定义", "展示/计算规则"],
        [
            ["AI友好型", "media_type=video 且 0<duration_ms≤15000", "15.000s 包含；15.001s 排除；时长未知默认排除"],
            ["近7天/近14天", "published_at ≥ now-7d/14d", "UTC 计算、边界包含；必须使用平台发布时间"],
            ["国家", "广告实际投放市场 country_code", "不以创作者所在地替代；未知单列"],
            ["媒体平台", "TikTok / Facebook（后续可扩展）", "P0 单选；跨平台指标口径不同，UI 给出提示"],
            ["投放状态", "投放中 / 已停投 / 状态未知", "由 normalized_status 映射；保留 raw status 与更新时间"],
            ["广告消耗预估", "估算区间及其 USD 中位值", "展示区间和置信度；排序按 midpoint_usd；缺失置底"],
            ["指标快照", "同一 metrics_at 下的指标集合", "播放/点赞/收藏/评论/CTR 不混用不同更新时间"],
        ],
        [1700, 3360, 4300],
    )

    doc.add_heading("范围与优先级", level=1)
    add_table(
        doc,
        ["优先级", "包含", "明确排除/后置"],
        [
            ["P0", "≤15s、7/14/30天、品类/国家/平台/状态、关键词、8种排序、卡片字段、URL 状态、服务端查询、埋点", "AI成功率评分、语义搜索、多选批量收藏"],
            ["P1", "自定义时长/日期、多选筛选、保存筛选条件、OCR/ASR 搜索、字段覆盖率提示", "自动生成与投放闭环"],
            ["P2", "AI友好度 V2：镜头数、运动复杂度、人物-商品交互、SKU 风险、历史复刻通过率", "模型供应商能力承诺"],
        ],
        [1100, 4510, 3750],
    )

    doc.add_page_break()
    doc.add_heading("产品原型与信息架构", level=1)
    add_prototype(doc)
    add_callout(
        doc,
        "原型状态说明",
        "效果图故意展示“AI友好型 ≤15s + 近7天 + TikTok + 美国 + 投放中 + 播放量倒序”的组合，用于暴露完整控件与卡片字段；首次进入页面时国家和投放状态仍默认为“全部”。",
    )
    doc.add_heading("界面结构", level=2)
    for item in [
        "第一层：沿用品类快捷入口与搜索框，保持当前浏览心智。",
        "第二层：AI友好型、时间、平台、国家、投放状态、更多筛选，全部使用现有胶囊/下拉视觉。",
        "第三层：结果数量、已选条件 chips、清除全部、排序菜单；筛选变更后立即可见且可撤销。",
        "素材卡片：封面内展示 AI友好、平台、投放状态、时长与发布时间；卡片底部展示 6 类指标和更新时间入口。",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("筛选控件规格", level=1)
    add_table(
        doc,
        ["控件", "类型/默认值", "选项", "关键行为"],
        [
            ["AI友好型", "开关型筛选；默认开启", "≤15s", "开启发送 duration_lte_ms=15000；关闭不限制时长"],
            ["时间", "单选；默认近7天", "近7天、近14天、近30天、自定义", "无结果不自动扩大；展示切换建议"],
            ["品类", "单选；默认全部", "沿用现有品类树", "一级快捷 pill；更多内展示完整分类"],
            ["媒体平台", "单选；沿用当前默认 TikTok", "全部、TikTok、Facebook", "切平台时保留其他条件"],
            ["国家", "可搜索单选；默认全部", "国家/地区标准列表", "含“未知”；语义是投放市场"],
            ["投放状态", "单选；默认全部", "全部、投放中、已停投、状态未知", "已停投可包含 pause/end，详情保留原始原因"],
            ["关键词", "文本输入；默认空", "标题、品牌、账号、平台文案 caption", "300ms 防抖；Enter 立即查询；多个词 AND"],
            ["排序", "单选；默认推荐", "推荐、最新、播放、点赞、收藏、评论、CTR、消耗预估", "数值默认倒序；同值按发布时间再 material_id"],
        ],
        [1500, 1970, 2550, 3340],
        keep_rows=False,
    )

    doc.add_heading("组合与状态规则", level=2)
    for item in [
        "不同维度之间为 AND；P0 每个维度仅允许一个值，因此无需定义维度内 OR。",
        "每次筛选、搜索、排序都更新 URL query；刷新、复制链接、浏览器前进/后退均可恢复。",
        "筛选变更后回到列表顶部并重置 cursor；保留用户主动选择的排序。",
        "清除全部恢复产品默认：AI友好型开启、近7天、平台沿用当前默认、品类/国家/状态为全部、排序推荐。",
        "结果请求期间保留上一屏并显示轻量 loading；新结果返回后整体替换，避免卡片闪烁。",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("详细功能需求", level=1)
    doc.add_heading("功能项", level=2)
    requirements = [
        ("FR-01  AI友好型筛选", "仅返回 duration_ms 已知、>0 且 ≤15000 的视频。卡片统一显示 mm:ss；后台保存毫秒，显示按秒向下/四舍五入规则需固定为“总秒数四舍五入到整数显示”，边界判断始终用原始毫秒。", "15.000s 可见；15.001s 与未知时长不可见。"),
        ("FR-02  时效筛选", "支持近7/14/30天和自定义区间。过滤字段只能是 published_at；first_seen_at 用于数据审计，不能伪装成发布时间。", "在精确边界时间发布的素材包含；无发布时间素材在日期筛选下排除。"),
        ("FR-03  品类/国家/平台", "沿用当前页面入口与视觉。国家使用 ISO country_code，展示本地化名称；platform 使用稳定枚举。", "切换任一项后其余条件保持，URL 可还原同一结果。"),
        ("FR-04  投放状态", "UI 使用投放中/已停投/状态未知；服务端保留 normalized_status、platform_status、status_reason、status_updated_at。", "状态映射可追溯；更新时间超过阈值时卡片可显示“状态待更新”。"),
        ("FR-05  关键词搜索", "对 title、brand_name、account_name、caption 做 trim、Unicode 规范化与不区分大小写匹配；多个非空词为 AND。", "空格输入等同空查询；清空后恢复原筛选结果。"),
        ("FR-06  指标排序", "播放/点赞/收藏/评论/CTR/消耗按数值倒序；缺失始终置底，不得按 0 混排；稳定 tie-breaker 为 published_at DESC、material_id ASC。", "相同请求连续返回顺序一致；翻页无重复/漏项。"),
        ("FR-07  卡片信息", "封面区显示 AI友好、平台、状态、时长、发布时间；指标区显示播放、点赞、收藏、评论、CTR、预估消耗。", "缺失字段显示“-”；hover/点击入口与现有详情/一键复刻保持一致。"),
        ("FR-08  数据时间", "每条素材返回 published_at、metrics_at、status_updated_at；详情或 tooltip 展示数据更新时间。", "任何指标都能判断数据采集时间；过期数据不伪装为实时。"),
        ("FR-09  列表状态", "提供首屏 skeleton、加载更多、空结果、请求失败与重试；空结果保留已选条件。", "错误不清空条件；重试使用同一 query。"),
        ("FR-10  深链与埋点", "筛选条件写入 URL，并记录筛选、排序、搜索、卡片点击、详情点击和一键复刻事件。", "刷新恢复条件；事件参数含完整 query 与 material_id。"),
    ]
    for title_text, rule, acceptance in requirements:
        doc.add_heading(title_text, level=3)
        add_labeled(doc, "规则", rule)
        add_labeled(doc, "验收", acceptance, label_color=GREEN)

    doc.add_page_break()
    doc.add_heading("排序口径", level=1)
    add_table(
        doc,
        ["UI 名称", "sort_by", "数据字段/计算", "缺失与并列处理"],
        [
            ["推荐", "recommendation", "沿用现有 recommendation_score DESC", "缺失按 0；再按发布时间、ID"],
            ["最新发布", "published_at", "published_at DESC", "无发布时间置底"],
            ["播放量", "play_count", "最新 metrics snapshot 的 play_count", "NULL 置底；不把 NULL 当 0"],
            ["点赞数", "like_count", "最新 snapshot 的 like_count", "NULL 置底"],
            ["收藏数", "save_count", "最新 snapshot 的 save_count", "NULL 置底"],
            ["评论数", "comment_count", "最新 snapshot 的 comment_count", "NULL 置底"],
            ["点击率", "ctr", "同一窗口 clicks/impressions；0–1 小数存储", "NULL 或分母为 0 置底"],
            ["广告消耗预估", "estimated_spend_mid_usd", "估算区间 low/high 的 USD 中位值", "NULL 置底；卡片展示区间与 confidence"],
        ],
        [1450, 1900, 3650, 2360],
    )
    add_callout(doc, "跨平台提示", "TikTok 播放与 Facebook 视频播放/展示的原始口径可能不同。P0 允许统一字段展示，但当用户选择“全部平台”并按播放量或 CTR 排序时，界面需提示“跨平台口径仅供趋势参考”。", fill="FFF8E8", accent=CAUTION)

    doc.add_heading("卡片与结果呈现", level=1)
    add_table(
        doc,
        ["区域", "必须字段", "显示规则"],
        [
            ["封面左上", "AI友好", "仅当当前规则命中；文案固定，不显示推测分数"],
            ["封面右上", "平台 + 投放状态", "平台图标/名称；状态带颜色点，未知为灰色"],
            ["封面底部", "时长 + 发布时间", "00:12；“3天前”，hover 展示绝对时间"],
            ["指标第一行", "播放、点赞、收藏", "K/M 紧凑格式；tooltip 展示完整数值"],
            ["指标第二行", "评论、CTR、预估消耗", "CTR 百分比两位；消耗显示区间或 midpoint"],
            ["交互", "详情/一键复刻", "沿用现有路径；埋点携带来源筛选上下文"],
        ],
        [1450, 2750, 5160],
    )

    doc.add_heading("数据模型", level=1)
    add_body(doc, "建议以“素材主表 + 指标快照 + 投放状态快照”建模。下表是前后端需要共同确认的最小合同；具体存储可由技术方案决定。")
    data_rows = [
        ["material_id", "string", "内部稳定 ID", "必填；分页 tie-breaker"],
        ["platform", "enum", "tiktok / facebook", "必填"],
        ["platform_material_id", "string", "平台素材/广告 ID", "与 platform 唯一"],
        ["content_url", "string", "原始内容链接", "权限/合法性校验"],
        ["media_type", "enum", "video / image", "AI友好型只取 video"],
        ["duration_ms", "integer", "原始视频时长毫秒", "筛选必需；未知为 NULL"],
        ["title", "string", "素材标题", "搜索字段"],
        ["brand_name", "string", "品牌名称", "搜索字段"],
        ["account_name", "string", "发布账号", "搜索字段"],
        ["caption", "text", "平台原始文案", "P0 搜索；缺失可空"],
        ["category_id", "string", "统一品类 ID", "与现有 taxonomy 对齐"],
        ["country_code", "string", "ISO 投放市场", "未知为 NULL"],
        ["published_at", "timestamp", "平台发布时间 UTC", "时效筛选唯一口径"],
        ["first_seen_at", "timestamp", "系统首次抓取时间", "仅审计/回补"],
        ["metrics_at", "timestamp", "指标快照时间 UTC", "与全部指标同批"],
        ["play_count", "bigint", "平台播放/视频浏览", "可空"],
        ["like_count", "bigint", "点赞数", "可空"],
        ["save_count", "bigint", "收藏数", "可空"],
        ["comment_count", "bigint", "评论数", "可空"],
        ["ctr", "decimal", "0–1 的点击率", "分母/时间窗写入 metadata"],
        ["estimated_spend_low_usd", "decimal", "预估消耗下界 USD", "与 high 成对"],
        ["estimated_spend_high_usd", "decimal", "预估消耗上界 USD", "与 low 成对"],
        ["estimated_spend_mid_usd", "decimal", "排序用中位值", "服务端计算"],
        ["spend_confidence", "enum", "high / medium / low", "卡片/tooltip 展示"],
        ["normalized_status", "enum", "active / stopped / unknown", "UI 过滤字段"],
        ["platform_status", "string", "平台原始状态", "不可丢弃"],
        ["status_reason", "string", "停投/异常原因", "可空"],
        ["status_updated_at", "timestamp", "状态更新时间 UTC", "新鲜度判断"],
        ["cover_url", "string", "封面资源", "必填"],
        ["video_url", "string", "播放资源", "按权限返回"],
        ["recommendation_score", "decimal", "现有推荐分", "定义保持现状"],
    ]
    add_table(doc, ["字段", "类型", "含义", "约束/用途"], data_rows, [2400, 1400, 3200, 2360], keep_rows=False)

    doc.add_heading("查询接口建议", level=1)
    add_callout(doc, "接口边界", "建议新增服务端列表接口 GET /api/discover/inspiration/materials。前端不得在已下载的局部 mock 列表上做“看似可用”的全量排序；筛选、排序与 cursor 分页必须在同一服务端查询完成。")
    api_rows = [
        ["q", "string", "关键词", "空值忽略"],
        ["category_id", "string", "品类", "缺省全部"],
        ["country_code", "string", "投放国家", "缺省全部"],
        ["platform", "enum", "平台", "缺省沿用当前默认"],
        ["duration_lte_ms", "integer", "最大时长", "AI友好为 15000"],
        ["published_within_days", "7|14|30", "相对日期窗口", "与自定义区间互斥"],
        ["published_from/to", "ISO timestamp", "自定义区间", "边界包含"],
        ["delivery_status", "enum", "active/stopped/unknown", "缺省全部"],
        ["sort_by", "enum", "排序字段", "缺省 recommendation"],
        ["sort_order", "asc|desc", "方向", "数值默认 desc"],
        ["cursor", "string", "稳定分页游标", "包含 sort value + id"],
        ["limit", "integer", "页大小", "建议 30，最大 60"],
    ]
    add_table(doc, ["参数", "类型", "用途", "规则"], api_rows, [2250, 1700, 3000, 2410], keep_rows=False)
    add_labeled(doc, "响应最小字段", "items、next_cursor、total_estimate、applied_filters、data_updated_at；每个 item 返回卡片所需字段与对应更新时间。")
    add_labeled(doc, "稳定排序", "所有 sort_by 必须追加 published_at DESC、material_id ASC；NULLS LAST，确保 cursor 翻页不重复、不漏项。")

    doc.add_heading("数据采集与新鲜度", level=1)
    add_table(
        doc,
        ["对象", "目标 SLA", "前台处理"],
        [
            ["新素材发现", "采集周期 ≤6 小时", "显示平台发布时间；采集延迟不伪装成发布时间"],
            ["投放中指标", "metrics_at p95 ≤24 小时", "超过24小时提示数据时间；超过72小时可标记陈旧"],
            ["投放状态", "status_updated_at p95 ≤24 小时", "过期显示“状态待更新”；仍保留 raw status"],
            ["已停投素材", "每日更新或事件回调", "停投时间可追溯"],
            ["日期覆盖率", "默认列表 published_at 覆盖率 100%", "日期筛选下未知发布时间素材直接排除"],
        ],
        [1850, 2500, 5010],
    )
    add_callout(doc, "禁止静默兜底", "当 CTR、消耗预估或投放状态的数据覆盖率不足时，排序选项应置灰并显示“当前仅覆盖 X% 素材”，或允许使用但明确提示覆盖率；禁止用随机数、0 或其他指标替代。", fill="FFF4F2", accent=RISK)

    doc.add_heading("异常、空状态与可访问性", level=1)
    state_rows = [
        ["首屏加载", "卡片 skeleton；筛选区可操作或显示统一 loading", "无布局跳变"],
        ["筛选刷新", "保留旧结果并显示顶部进度；成功后替换", "防止白屏"],
        ["无结果", "展示当前条件与“切换近14天/清除部分条件”", "不自动扩大时间"],
        ["请求失败", "保留条件，显示重试；记录 request_id", "重试同一 query"],
        ["指标缺失", "显示“-”与缺失原因 tooltip", "排序 NULLS LAST"],
        ["键盘/焦点", "Tab 可达全部控件；Enter/Space 操作；Esc 关闭菜单", "焦点环对比度达标"],
        ["读屏", "筛选控件有可读 label；状态不只靠颜色", "卡片指标有 aria-label"],
    ]
    add_table(doc, ["场景", "产品行为", "验收重点"], state_rows, [1650, 4900, 2810])

    doc.add_heading("埋点与衡量", level=1)
    analytics_rows = [
        ["inspiration_filter_change", "filter_name、old_value、new_value、full_query", "筛选使用率、无结果率"],
        ["inspiration_sort_change", "sort_by、sort_order、full_query", "高价值排序偏好"],
        ["inspiration_search", "query_length、term_count、result_count", "不采集敏感全文（按策略）"],
        ["inspiration_card_click", "material_id、rank、full_query", "筛选后详情点击率"],
        ["inspiration_replicate_click", "material_id、rank、full_query", "筛选后复刻启动率"],
        ["inspiration_empty_result", "full_query", "筛选可用性与数据缺口"],
    ]
    add_table(doc, ["事件", "关键参数", "用途"], analytics_rows, [2750, 3950, 2660])
    add_labeled(doc, "上线后观察", "先收集两周基线，再确定 AI友好筛选使用率、卡片点击率、一键复刻启动率与无结果率的目标；不在无基线时虚设业务指标。")

    doc.add_page_break()
    doc.add_heading("验收测试矩阵", level=1)
    test_rows = [
        ["AC-01", "duration_ms=15000，AI友好开启", "素材包含", "FR-01"],
        ["AC-02", "duration_ms=15001 或 NULL", "素材排除", "FR-01"],
        ["AC-03", "published_at=now-7d 精确边界", "近7天包含", "FR-02"],
        ["AC-04", "近7天 0 结果、近14天有结果", "显示切换建议，不自动混入", "FR-02/09"],
        ["AC-05", "≤15s+美国+TikTok+投放中", "仅返回四条件交集", "FR-03/04"],
        ["AC-06", "按播放量排序且存在 NULL", "非空倒序，NULL 最后", "FR-06"],
        ["AC-07", "消耗 low/high 不同币种", "先统一 USD midpoint，再排序", "FR-06"],
        ["AC-08", "同一排序值的多条素材分页", "按时间/ID 稳定，无重复漏项", "FR-06"],
        ["AC-09", "关键词命中标题/品牌/账号/caption", "全部字段可召回；多词 AND", "FR-05"],
        ["AC-10", "刷新带 query 的 URL", "恢复筛选、排序和搜索", "FR-10"],
        ["AC-11", "状态更新时间 >24h", "显示数据时间或状态待更新", "FR-04/08"],
        ["AC-12", "CTR/消耗字段缺失", "卡片显示“-”，排序置底", "FR-06/07"],
        ["AC-13", "API 失败后点击重试", "条件不变，重复同一请求", "FR-09"],
        ["AC-14", "键盘操作全部筛选与排序", "可聚焦、可选择、可关闭", "无障碍"],
    ]
    add_table(doc, ["编号", "前置/输入", "预期结果", "关联需求"], test_rows, [1050, 3700, 3340, 1270], keep_rows=False)

    doc.add_heading("非功能要求", level=1)
    for item in [
        "性能：缓存命中时列表查询 p95 ≤1.5s；关键词查询 p95 ≤2s；首屏默认 30 条，cursor 加载。",
        "一致性：卡片、详情、排序使用同一指标快照，不允许前端二次计算导致口径漂移。",
        "可观测性：记录 query hash、数据源、耗时、命中数量、字段覆盖率、request_id 与错误码。",
        "安全与合规：外部平台内容、封面与数据的采集、缓存、展示、跳转遵守授权、平台条款与区域政策。",
        "兼容性：桌面宽屏保持现有 5–7 列布局；较窄窗口控件换行但不遮挡，菜单不越界。",
        "可访问性：键盘可操作、可见焦点、状态文本化、颜色对比与读屏标签满足产品现有标准。",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("工程拆分建议", level=1)
    add_table(
        doc,
        ["工作流", "建议产出", "验收门槛"],
        [
            ["数据/采集", "字段覆盖率审计、平台映射、发布时间/状态/指标快照、消耗估算来源", "能回答每个字段“从哪里来、何时更新、缺失怎么办”"],
            ["后端", "统一列表 API、过滤排序、cursor、缓存、数据时间与覆盖率", "AC-01～13 的接口用例通过"],
            ["前端", "FilterBar、ActiveFilterChips、SortMenu、MaterialCard、URL 状态、空/错/加载", "与原型一致；键盘与响应式通过"],
            ["QA", "边界、组合、排序稳定性、NULL、跨平台、深链、陈旧数据", "验收矩阵全量覆盖"],
            ["灰度", "feature flag、字段覆盖率监控、事件看板", "数据未达标时可回退旧列表"],
        ],
        [1450, 4500, 3410],
        keep_rows=False,
    )
    add_labeled(doc, "建议前端组件", "InspirationFilterBar、ActiveFilterChips、InspirationSortMenu、InspirationMaterialCard、useInspirationQueryState；保留 /discover/inspiration 路由。")
    add_labeled(doc, "建议索引方向", "published_at + duration_ms 为核心，叠加 platform/country/category/status；搜索单独使用全文索引。具体索引需根据数据量与查询计划确定。")

    doc.add_heading("依赖、风险与待决策项", level=1)
    decision_rows = [
        ["D-01", "CTR 数据来源与时间窗", "没有可验证来源则隐藏/置灰排序", "数据负责人"],
        ["D-02", "广告消耗预估方法、币种转换、置信度", "决定是否能上线该排序", "数据/算法"],
        ["D-03", "国家语义是否稳定代表投放市场", "避免把创作者国家当投放国家", "产品/数据"],
        ["D-04", "投放状态映射与更新频率", "避免“投放中”已过期", "数据/后端"],
        ["D-05", "现有“推荐”分的来源与保持方式", "防止新 API 改变原排序体验", "产品/后端"],
        ["D-06", "页面截图版本与仓库当前 mock 的合并基线", "技术方案需确认最终组件入口", "前端负责人"],
        ["D-07", "平台内容展示与缓存授权", "影响封面、播放与长期留存", "法务/平台合作"],
    ]
    add_table(doc, ["编号", "待确认", "不确认的影响", "Owner"], decision_rows, [950, 3300, 3700, 1410], keep_rows=False)
    add_callout(doc, "技术方案进入条件", "D-01～D-04 至少要有明确的数据来源、覆盖率和降级策略；否则只能开发 UI 骨架，不能宣称筛选/排序能力完成。", fill="FFF8E8", accent=CAUTION)

    doc.add_heading("上线门槛", level=1)
    for item in [
        "功能：P0 需求与 AC-01～14 全部通过；无 P0 blocker。",
        "数据：默认结果 100% 有 duration_ms 与 published_at；CTR/消耗/状态显示覆盖率与降级行为符合约定。",
        "性能：核心查询达到 p95 目标；连续翻页无重复/漏项。",
        "体验：效果图中的筛选层级、已选条件、排序菜单与卡片字段在目标桌面分辨率下无重叠/截断。",
        "灰度：feature flag、错误监控、字段覆盖率看板和回退方案就绪。",
    ]:
        add_bullet(doc, item, bullet_num_id)

    doc.add_heading("附录：版本说明", level=1)
    add_labeled(doc, "V1.0", "根据客户需求、现有灵感发现截图与当前仓库实现梳理；新增 AI友好型、时效、投放状态、完整排序口径、数据合同、API 建议和验收矩阵。")
    add_labeled(doc, "设计说明", "原型图为产品评审参考，视觉保持现有 CreatiSignal 语言；具体像素值由前端基于现有 design tokens 复用，不在本 PRD 另建视觉系统。")

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
