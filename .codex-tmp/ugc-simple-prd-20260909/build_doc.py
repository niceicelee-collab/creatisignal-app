from pathlib import Path
import re, json, shutil, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent.parent / 'deliverables' / 'ugc-simple-prd-20260909'
OUT.mkdir(parents=True, exist_ok=True)
IMG = ROOT / 'images'
IMG.mkdir(exist_ok=True)
GEN = Path('C:/Users/ice.li/.codex/generated_images/01a07c8b-8080-7143-9436-d01cab6ec7f3')
mapping = {
    'feed.png': GEN / 'exec-a6251423-b42d-4eb4-8114-64c97c16246d.png',
    'progress.png': GEN / 'exec-5d979730-1d61-4815-8121-32740753e647.png',
    'actions.png': GEN / 'exec-1a00d457-f721-4ad7-937d-a1a1d64e3735.png',
    'favorites.png': GEN / 'exec-4c49c1b7-59c8-416e-8c66-d93552bfb56d.png',
    'flow.png': Path('C:/Users/ice.li/AppData/Local/Temp/codex-clipboard-d8b1752d-6bf3-496b-901c-20fb8bff2bfd.png'),
}
for name, src in mapping.items():
    shutil.copy2(src, IMG / name)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.left_margin = sec.right_margin = Inches(.60)
sec.top_margin, sec.bottom_margin = Inches(.52), Inches(.52)
sec.footer_distance = Inches(.20)
W = 7.07

def setfont(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rp = style.element.get_or_add_rPr()
    rp.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    rp.rFonts.set(qn('w:ascii'), 'Calibri')
    rp.rFonts.set(qn('w:hAnsi'), 'Calibri')
    lang = OxmlElement('w:lang'); lang.set(qn('w:eastAsia'), 'zh-CN'); rp.append(lang)

setfont(doc.styles['Normal'], 10.5)
pf = doc.styles['Normal'].paragraph_format
pf.line_spacing = Pt(15.5)
pf.space_after = Pt(5)
pf.widow_control = True
for name, size, leading in [('Title',25,34),('Heading 1',17,24),('Heading 2',14,21),('Caption',9,12)]:
    setfont(doc.styles[name], size, name!='Caption')
    f=doc.styles[name].paragraph_format
    f.line_spacing=Pt(leading)
    f.space_before=Pt(8 if name.startswith('Heading') else 0)
    f.space_after=Pt(6)
    f.keep_with_next = name!='Caption'
for st in doc.styles:
    for b in list(st.element.iter(qn('w:pBdr'))): b.getparent().remove(b)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('UGC爆款需求文档  ·  V2.0  ·  '); r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def inline(p, text):
    for i,s in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        if s:
            r=p.add_run(s); r.bold=bool(i%2)

def table(lines):
    rows=[[c.strip() for c in l.strip('|').split('|')] for l in lines]
    rows=[row for row in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in row)]
    is_accept=rows[0][0]=='编号'
    widths=([.48,1.49,5.10] if is_accept else [1.33,5.04,.70]) if len(rows[0])==3 else [1.37,5.70]
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    for col,wi in zip(t.columns,widths):col.width=Inches(wi)
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(round(W*1440)));tw.set(qn('w:type'),'dxa')
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
        borders.append(b)
    pr.append(borders)
    for ri,values in enumerate(rows):
        row=t.add_row();trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,(c,value) in enumerate(zip(row.cells,values)):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDE4FA' if ri==0 else 'FFFFFF');cp.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top',65),('bottom',65),('left',85),('right',85)]:
                v=OxmlElement('w:'+edge);v.set(qn('w:w'),str(val));v.set(qn('w:type'),'dxa');mar.append(v)
            cp.append(mar)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(13)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (is_accept and ci==0) or (not is_accept and len(values)==3 and ci==2) else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(value);r.font.size=Pt(9.5);r.bold=ri==0
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)

source=(ROOT/'content.md').read_text(encoding='utf-8')
image_count=0
for pi,page in enumerate(source.split('\n---PAGE---\n')):
    lines=page.strip().splitlines();i=0;first=True
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if s.startswith('|'):
            ls=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                ls.append(lines[i].strip());i+=1
            table(ls);continue
        if s.startswith('# '):p=doc.add_paragraph(s[2:],'Title')
        elif s.startswith('## '):p=doc.add_paragraph(s[3:],'Heading 1')
        elif s.startswith('### '):p=doc.add_paragraph(s[4:],'Heading 2')
        elif s.startswith('!['):
            m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',s)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(IMG/m[2]),width=Inches(W))
            shape._inline.docPr.set('descr',m[1]);shape._inline.docPr.set('title',m[1]);image_count+=1
        elif s.startswith('> '):
            p=doc.add_paragraph(s[2:],'Caption')
        else:
            p=doc.add_paragraph();inline(p,s)
        if first and pi:p.paragraph_format.page_break_before=True
        first=False;i+=1

doc.core_properties.title='UGC爆款需求文档'
doc.core_properties.subject='简版需求补充 第7点功能原型与说明及我的收藏'
doc.core_properties.author='CreatiSignal'
doc.core_properties.keywords='UGC爆款, 沉浸式浏览, 我的收藏, 简版需求'
doc.core_properties.comments=''
path=OUT/'UGC爆款需求文档_简版补充稿_V2.0.docx'
doc.save(path)
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    xml=z.read('word/document.xml').decode('utf-8')
    assert '我的收藏' in xml and '商品库' in xml
    assert len(z.namelist())>10
manifest={'output':str(path),'size':path.stat().st_size,'images':image_count,'tables':len(doc.tables),'source_chars':len(source),'source_flow':str(mapping['flow.png']),'prototype_files':{k:str(v) for k,v in mapping.items() if k!='flow.png'}}
(ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest,ensure_ascii=False))
