from __future__ import annotations

from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"D:\Cursor\creatisignal-app-new")
OUT_DIR = ROOT / "deliverables" / "gmv-max-diagnosis-prd"
OUT_PATH = OUT_DIR / "CreatiSignal_GMV_Max素材诊断闭环_PRD_重构版_V2.0.docx"
HELPER_PATH = ROOT / ".codex-tmp" / "ai-friendly-inspiration-prd" / "build_prd.py"

SOURCE_IMAGES = {
    "overview": Path(r"C:\Users\ice.li\AppData\Local\Temp\codex-clipboard-a1d48a3d-d7e1-4ebd-9d48-aaec9a158b09.png"),
    "diagnosis": Path(r"C:\Users\ice.li\AppData\Local\Temp\codex-clipboard-f10c97ba-887c-4905-8238-6bfbbee58343.png"),
    "delivery": Path(r"C:\Users\ice.li\AppData\Local\Temp\codex-clipboard-e2a9ad97-6064-463b-ad03-9538c18bbd69.png"),
    "callback": Path(r"C:\Users\ice.li\AppData\Local\Temp\codex-clipboard-110955f6-1e17-4282-990d-46436c02d8cd.png"),
}


def load_helpers():
    spec = spec_from_file_location("doc_helpers", HELPER_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load helpers from {HELPER_PATH}")
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


h = load_helpers()


def add_numbered(doc: Document, text: str, num_id: int, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="Normal")
    h.apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        h.set_run_font(r1, 11, h.INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        h.set_run_font(r2, 11, h.BLACK)
    else:
        r = p.add_run(text)
        h.set_run_font(r, 11, h.BLACK)


def add_formula(doc: Document, name: str, formula: str, note: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F9")
    p_pr.append(shd)
    r1 = p.add_run(f"{name}  ")
    h.set_run_font(r1, 10.5, h.DARK_BLUE, bold=True)
    r2 = p.add_run(formula)
    h.set_run_font(r2, 10.5, h.BLACK, bold=True, name="Consolas")
    r3 = p.add_run(f"\n{note}")
    h.set_run_font(r3, 9.5, h.MUTED)


def add_source_image(doc: Document, key: str, caption: str, alt: str, width: float = 6.5) -> None:
    path = SOURCE_IMAGES[key]
    if not path.exists():
        h.add_callout(doc, "原型图缺失", f"未找到原始截图：{path}", fill="FFF4E5", accent="A15C00")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    cp = doc.add_paragraph(style="Caption Text")
    cp.add_run(caption)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    h.configure_page(doc)
    h.configure_styles(doc)
    h.configure_header_footer(doc)
    bullet_id = h.add_num_definition(doc, "bullet")
    number_id = h.add_num_definition(doc, "decimal")

    # Replace inherited header label.
    header = doc.sections[0].header.paragraphs[0]
    header.clear()
    hr = header.add_run("CreatiSignal  |  GMV Max 素材诊断闭环 PRD")
    h.set_run_font(hr, 9, h.MUTED, bold=True)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    props = doc.core_properties
    props.title = "CreatiSignal GMV Max 素材诊断闭环 PRD（重构版）"
    props.subject = "从经营异常定位、素材诊断、动作执行、投放草稿到结果回流的工程可落地需求"
    props.author = "CreatiSignal Product"
    props.keywords = "GMV Max, 素材诊断, 决策闭环, 投放中心, 结果回流, PRD"
    props.comments = "V2.0 restructured engineering-ready PRD based on supplied draft and current repository baseline"

    # Cover / memo masthead.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(5)
    kr = kicker.add_run("PRODUCT REQUIREMENTS DOCUMENT · RESTRUCTURED")
    h.set_run_font(kr, 10, h.GREEN, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("CreatiSignal GMV Max 素材诊断闭环")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("以可解释诊断为入口，以可审计执行与结果回流为终点")

    h.add_labeled(doc, "版本", "V2.0（基于领导 AI 初稿重构）")
    h.add_labeled(doc, "日期", "2026-08-31")
    h.add_labeled(doc, "目标读者", "产品、数据、算法、前端、后端、测试、TikTok Shop GMV Max 投放 AO")
    h.add_labeled(doc, "文档状态", "可用于技术方案评审；平台接口、归因口径与阈值仍需业务/数据 Owner 确认")
    h.add_labeled(doc, "产品范围", "经营总览 → 素材诊断 → 创意生成/动作配置 → 投放中心 → 任务与结果回流")

    h.add_callout(
        doc,
        "一句话结论",
        "原稿覆盖面很全，但把业务规则、原型交互、Demo 数据和工程实现混在了一起。重构后采用“一套诊断对象、一条闭环主流程、六个用户结果、三类动作分支”，并把口径、优先级、状态机和验收条件单独定义。",
        fill=h.LIME_SOFT,
        accent=h.GREEN,
    )

    doc.add_heading("阅读导航", level=1)
    nav_rows = [
        ["先看", "第 1–4 章", "背景、原稿问题、目标边界、业务闭环"],
        ["产品/设计", "第 5–7 章", "诊断模型、功能清单、逐功能需求"],
        ["数据/算法", "第 5、8 章", "指标公式、样本门槛、Benchmark、数据实体"],
        ["工程/测试", "第 8–11 章", "状态机、接口契约、优先级、验收标准"],
        ["评审决策", "第 12 章", "仍需确认的业务口径与平台能力"],
    ]
    h.add_table(doc, ["角色", "章节", "重点"], nav_rows, [1500, 1800, 6060])

    page_break(doc)
    doc.add_heading("1. 重构结论与关键改动", level=1)
    h.add_callout(
        doc,
        "推荐结构",
        "诊断结果不是六套互不相干的产品流程。系统先完成数据资格校验与归因保护，再输出一个诊断快照；结果只决定主动作是放量、保持、观察、迭代、换新或关停。生成、配置、发布、观察和回流均复用统一任务与审计能力。",
        fill="EAF4FF",
        accent=h.BLUE,
    )
    changes = [
        ["原稿表达", "主要风险", "重构后的处理"],
        ["六类结果各走独立步骤", "重复状态与交互，后端难维护，规则容易冲突", "一条主流程 + 生成/监控/关停三类动作分支；六类仅是用户结果"],
        ["商品、素材、账户混作诊断主体", "同一素材跨商品/账户时结论失真", "主诊断键固定为店铺×账户×商品×素材×投放上下文×时间窗；商品列表只做聚合"],
        ["ROI 与 ROAS 混用", "公式看似正确但指标含义不一致", "产品字段保留 target_roi 兼容平台命名，指标字典明确其口径为 Revenue/Spend；展示层可写“目标 ROI（平台口径）”"],
        ["阈值写死在 PRD/Demo 中", "无法调参、回溯与解释历史结论", "所有阈值进入版本化规则集 rule_version；V2.0 只给 MVP 默认建议"],
        ["样本成熟条件过于宽松", "短时间高消耗或少量订单被误判", "数据完整 + Active≥24h +（Orders≥5 或 Spend≥2×Target CPO）才进入强动作判断"],
        ["Benchmark 不足时直接扩品类", "可比性下降却仍给同等置信度", "分层回退并输出 benchmark_level、sample_size、confidence；低置信度禁止关停/强放量"],
        ["预计 GMV 影响未定义", "排序结果不可复现，容易与真实 GMV 混淆", "P0 使用可解释的风险/机会分；预测金额后置 P2，必须标“预估”并记录 model_version"],
        ["创建成功等同投放成功", "有广告 ID 但未审核/未消耗也会显示成功", "拆分创建状态与投放生命周期；保留 raw platform_status 及原因"],
        ["关停等同删除素材", "不可恢复且可能影响其他商品", "仅对当前投放关系创建可恢复的停用/移出建议，二次确认并记录影响范围"],
        ["大表同时描述目标、UI、Demo 和数据", "工程师难找到真正约束，测试无法逐条验收", "按业务规则、功能需求、数据模型、状态机、验收矩阵分别编排"],
    ]
    # First row above is deliberately removed from content and represented by real header.
    h.add_table(doc, ["原稿问题", "风险", "修订原则"], changes[1:], [2600, 2900, 3860], keep_rows=False)

    doc.add_heading("1.1 本版保留的核心价值", level=2)
    for text in [
        "AO 在经营总览中知道今天优先处理哪个商品、原因是什么、预计影响方向如何。",
        "从商品经营问题进入素材诊断后，平台、店铺、账户、国家、商品与时间窗保持一致。",
        "每个诊断快照只给一个主结果与一个主动作，证据、规则版本和数据时间可追溯。",
        "任何修改投放、目标、预算或停用关系的动作均由 AO 确认，不进行静默自动执行。",
        "只有投放后的观测窗口与样本门槛同时满足，才能验收并回写下一轮诊断。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("1.2 当前仓库实现基线", level=2)
    h.add_body(doc, "当前 Insights 页面主要基于本地 mock 数据和固定 ROI 阈值：账户状态为可放量/学习中/需改写/已停用，实验页按 ROI 2.0、1.3、0.8 等固定阈值判定；尚未形成产品×素材×投放上下文的诊断快照、规则版本、Benchmark 置信度、投放草稿和回流实体。")
    h.add_callout(doc, "工程约束", "前端不得继续直接计算正式诊断结论。P0 应先定义服务端/领域层的指标与规则输出，前端只渲染结果、证据与允许动作；Demo 数值不得成为生产默认值。", fill="FFF4E5", accent="A15C00")

    page_break(doc)
    doc.add_heading("2. 背景、问题与目标", level=1)
    doc.add_heading("2.1 业务背景", level=2)
    h.add_body(doc, "TikTok Shop GMV Max 将商品、素材池、预算与目标回报共同参与投放。运营人员可以看到大量经营与素材指标，但当前页面能力彼此割裂：发现经营异常后，仍需人工跨页面寻找问题素材、判断是否属于素材原因、安排创意处理，再单独记录投放结果。这个过程耗时、不可复盘，也难以形成稳定规则。")

    doc.add_heading("2.2 目标用户与核心任务", level=2)
    roles = [
        ["AO / 投放负责人", "决定优先处理对象，确认预算、目标和停用动作", "今日处理队列、诊断证据、投放草稿、结果回流"],
        ["素材负责人", "理解需要改变的变量并获得可执行 brief", "迭代变量、参考素材、生成结果、版本关系"],
        ["数据/算法", "维护口径、Benchmark 与规则版本", "指标快照、样本置信度、模型/规则版本、误判反馈"],
        ["管理者", "查看闭环效率与动作收益", "任务进度、验证率、胜出率、阻塞原因"],
    ]
    h.add_table(doc, ["用户", "主要职责", "产品支持"], roles, [1900, 3300, 4160])

    doc.add_heading("2.3 产品目标", level=2)
    for text in [
        "30 秒内回答：今天先处理哪个商品、为什么、下一步做什么。",
        "把经营信号定位到可解释的 Product × Creative 诊断快照，并区分素材问题与非素材问题。",
        "把诊断结论转换为可执行任务：生成/复刻、监控复查、投放配置或可恢复关停。",
        "所有外部写操作均经过人工确认、幂等执行、状态跟踪与审计。",
        "用统一观测窗口比较原素材与新素材/新配置，并将结果回写下一轮决策。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("2.4 本期非目标", level=2)
    for text in [
        "不做素材级独立竞价；预算与目标仍属于商品计划/GMV Max 配置。",
        "不自动修改预算、Target ROI、素材关系或关停计划；系统只生成建议与待确认草稿。",
        "不把价格、折扣、库存、毛利、授权、账户或店铺状态错误归因成素材问题。",
        "不承诺仅凭 CTR/CVR/ROI 能给出因果结论；产品输出为规则性建议和证据，不声称统计因果。",
        "P0 不建设黑盒 GMV 预测模型、多级审批、批量自动化与跨组织权限系统。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("2.5 成功指标", level=2)
    success = [
        ["效率", "从经营总览进入可执行动作的中位时长", "上线前埋点建立基线；P0 目标较基线下降 ≥30%"],
        ["解释性", "诊断快照具备完整证据、规则版本、数据时间的比例", "≥99%"],
        ["闭环率", "已执行任务中完成观测并回写结果的比例", "P1 上线 4 周后 ≥70%"],
        ["安全性", "未经确认发生的外部写操作", "0"],
        ["稳定性", "重复请求造成重复草稿/重复发布", "0；幂等校验覆盖 100% 写操作"],
        ["可追溯", "动作记录包含 operator/before/after/request_id/platform_response", "100%"],
    ]
    h.add_table(doc, ["维度", "指标", "建议目标"], success, [1500, 4600, 3260])

    page_break(doc)
    doc.add_heading("3. 核心对象、术语与口径", level=1)
    doc.add_heading("3.1 诊断对象与聚合层级", level=2)
    add_formula(doc, "诊断主键", "shop_id × account_id × product_id × creative_id × delivery_context × window", "campaign/adgroup 等平台层级按实际接入能力放入 delivery_context。相同 creative_id 在不同商品、账户或计划下必须分别诊断。")
    h.add_body(doc, "经营总览以商品为展示粒度，但商品行只是多个诊断快照的聚合。点击进入详情后，必须落到明确的素材-商品-投放上下文，不能用素材全局均值代替。")
    object_rows = [
        ["Product", "商品/SKU", "承载目标 ROI、AOV、库存/Offer 等商品条件"],
        ["Creative", "视频/图片素材", "保留平台素材 ID、指纹、来源、版本与可投状态"],
        ["Delivery context", "账户/计划/广告组/国家", "决定指标口径、素材关系与写操作影响范围"],
        ["Diagnosis snapshot", "不可变诊断快照", "冻结指标、Benchmark、规则版本、结果、证据与允许动作"],
        ["Action intent", "待执行动作意图", "承接放量、保持、观察、迭代、换新、关停中的唯一主动作"],
        ["Delivery draft", "待发布配置", "保存目标、预算、素材清单、观察与止损规则，不等同平台已创建"],
        ["Validation result", "观测结果", "比较执行前后/原素材与变体，形成 winner 与下一动作"],
    ]
    h.add_table(doc, ["对象", "定位", "关键约束"], object_rows, [1900, 2000, 5460], keep_rows=False)

    doc.add_heading("3.2 指标术语", level=2)
    metric_rows = [
        ["Spend", "归因窗口内消耗", "货币统一后聚合"],
        ["Gross revenue", "归因窗口内 GMV/成交金额", "明确是否含退款、取消与自然成交"],
        ["CTR", "Clicks / Impressions", "Impressions=0 时为 null，不显示 0%"],
        ["CVR", "Orders / Clicks", "Clicks=0 时为 null"],
        ["AOV", "Gross revenue / Orders", "Orders=0 时为 null"],
        ["ROI（平台口径）", "Gross revenue / Spend", "数学上等价 ROAS 类口径；全产品统一叫法"],
        ["Target CPO", "Expected AOV / Target ROI", "仅在 Target ROI>0 且 Expected AOV 可用时计算"],
        ["Index", "Actual / Benchmark 或 Actual / Target", "分子分母必须同时间窗、国家、币种与归因口径"],
    ]
    h.add_table(doc, ["指标", "计算", "空值/口径规则"], metric_rows, [2000, 3100, 4260])
    h.add_callout(doc, "必须确认", "原稿把 ROI、ROAS、Target ROI 混用。若平台字段名固定为 Target ROI，可在接口层保留原名，但数据字典必须写明其分子分母，前端不能在不同页面切换语义。", fill="FFF4E5", accent="A15C00")

    page_break(doc)
    doc.add_heading("4. 业务整体逻辑（闭环流程）", level=1)
    h.add_body(doc, "所有业务分支共享以下主干。任何一步产生的数据快照、动作意图与状态变更都必须可追溯。")
    steps = [
        ("选择上下文", "AO 选择平台、店铺、账户、国家与时间范围；系统返回 data_as_of 与时区。"),
        ("经营聚合与排队", "按商品聚合 GMV/Spend/ROI、素材供给与异常信号，输出今日处理队列。"),
        ("数据资格校验", "检查数据完整性、延迟、归因窗口、样本门槛与 Benchmark 置信度；不合格直接进入待观察。"),
        ("归因保护", "当素材漏斗指标正常而 ROI 不达标，检查 AOV、Offer、库存、毛利/目标、账户或平台状态；命中后禁止素材动作。"),
        ("生成诊断快照", "冻结 Product × Creative × Context 的实际值、基准、Index、趋势、规则版本与唯一主结果。"),
        ("创建唯一主动作", "根据结果生成放量、保持、观察、迭代、换新或关停的 Action Intent；次级入口仅用于查看证据。"),
        ("执行分支", "生成类进入方向/变量选择与素材生成；监控类创建复查任务；关停类进入影响确认。"),
        ("保存投放草稿", "生成或复用素材后配置 Target ROI、日预算、观察时长、胜出订单线与止损 ROI，保存草稿。"),
        ("人工确认与平台执行", "AO 确认后调用平台；创建成功与实际投放状态分离，并持续同步审核/投放状态。"),
        ("观测、验收与回流", "同时满足观察时间和订单样本后，比较原素材与变体/新配置；结果回写 Benchmark、任务和下一轮诊断。"),
    ]
    for title_text, desc in steps:
        add_numbered(doc, f"{title_text}：{desc}", number_id, bold_prefix=f"{title_text}：")

    doc.add_heading("4.1 三类动作分支", level=2)
    branch_rows = [
        ["生成/配置", "可放量、需迭代、需换新", "选择参考/变量 → 生成/复用素材 → 选择结果 → 配置并保存草稿"],
        ["监控/复查", "稳定投放、待观察", "创建复查规则 → 累积样本/等待数据 → 自动或人工重诊断"],
        ["风险处置", "建议关停、非素材保护", "展示影响范围 → 二次确认/跳转商品或投放诊断 → 记录可恢复操作"],
    ]
    h.add_table(doc, ["分支", "适用结果", "共用路径"], branch_rows, [1800, 2500, 5060])

    doc.add_heading("4.2 一条诊断只能有一个主结果和一个主动作", level=2)
    h.add_body(doc, "同一 diagnosis_snapshot_id 只允许一个 active action_intent。AO 变更结论时，必须关闭旧意图并创建新版本，不可在同一快照上同时“放量”和“关停”。查看商品问题、查看证据、返回上游不算主动作。")

    page_break(doc)
    doc.add_heading("5. 数据计算与诊断决策模型", level=1)
    doc.add_heading("5.1 基础指标", level=2)
    add_formula(doc, "点击率", "ctr = clicks / impressions", "impressions 为 0 或数据缺失时返回 null；不得以 0 参与 Index。")
    add_formula(doc, "转化率", "cvr = orders / clicks", "clicks 为 0 时返回 null；订单口径必须与 GMV 归因一致。")
    add_formula(doc, "客单价", "aov = gross_revenue / orders", "orders 为 0 时返回 null；禁止展示为 $0 以免误导。")
    add_formula(doc, "平台口径 ROI", "roi = gross_revenue / spend", "spend 为 0 时返回 null；若平台提供权威值，保留 raw value 与本地重算差异。")
    add_formula(doc, "目标单均成本", "target_cpo = expected_aov / target_roi", "仅当 target_roi>0；expected_aov 优先取同商品近期稳定样本，不使用当前少样本素材自身 AOV。")

    doc.add_heading("5.2 标准化 Index 与趋势", level=2)
    add_formula(doc, "ROI 达成指数", "roi_index = actual_roi / target_roi", "目标与实际必须使用同一口径。")
    add_formula(doc, "CTR 指数", "ctr_index = creative_ctr / benchmark_ctr", "Benchmark 为 null 或置信度低时，不做强结论。")
    add_formula(doc, "CVR 指数", "cvr_index = creative_cvr / benchmark_cvr", "Benchmark 为 null 或置信度低时，不做强结论。")
    add_formula(doc, "短期趋势", "trend_3d_vs_prev7d = metric_last_3d / metric_prev_7d - 1", "两个窗口必须连续、不重叠；任一窗口样本不足则返回 null。")
    add_formula(doc, "波动率", "volatility = stdev(daily_metric) / mean(daily_metric)", "均值为 0 或有效天数不足 3 天时返回 null。")

    doc.add_heading("5.3 Benchmark 分层与置信度", level=2)
    bench_rows = [
        ["L1", "同商品 + 同国家 + 同账户 + 近7天成熟素材", "sample≥5", "High"],
        ["L2", "同商品 + 同国家 + 同账户 + 近14天成熟素材", "sample≥5", "Medium"],
        ["L3", "同商品 + 同国家 + 同店铺其他账户 + 近14天", "sample≥5", "Medium-Low"],
        ["L4", "同类目 + 同国家 + 同店铺 + 近14天", "sample≥10", "Low；只用于方向提示"],
        ["None", "以上均不足", "—", "No benchmark；只能待观察"],
    ]
    h.add_table(doc, ["级别", "样本池", "最低样本", "置信度/动作限制"], bench_rows, [900, 4900, 1400, 2160])
    h.add_body(doc, "Benchmark 默认取样本中位数，并记录 benchmark_level、sample_size、window、filter、calculated_at。L4 或 None 不允许直接触发强放量/关停；前端必须显示“参考范围扩大”。")

    doc.add_heading("5.4 样本成熟门槛", level=2)
    add_formula(doc, "MVP 成熟条件", "is_mature = data_complete AND active_hours ≥ 24 AND (orders ≥ 5 OR spend ≥ 2 × target_cpo)", "与原稿相比，避免仅凭短时间高消耗或极少订单形成强结论。orders、24h、2× 均进入规则配置。")
    h.add_callout(doc, "待观察优先", "数据延迟、归因未闭合、Benchmark 不可用、样本未成熟时，统一输出“待观察”，并给出缺口与预计重诊断时间；不得生成关停、放量或替换素材动作。", fill="EAF4FF", accent=h.BLUE)

    doc.add_heading("5.5 非素材问题保护", level=2)
    h.add_body(doc, "当素材漏斗相对正常，但 ROI 未达标时，不把问题归因到素材。MVP 默认保护条件：ctr_index≥0.90 且 cvr_index≥0.90 且 roi_index<1.00；再结合 AOV、Offer、库存、毛利/目标、授权、账户余额、审核和投放状态生成原因码。")
    h.add_callout(doc, "结果语义", "“非素材问题”不是第七个素材诊断状态，而是归因保护分支。命中后锁定生成/换新/关停素材按钮，主动作改为“查看商品与投放诊断”。", fill="F1F5FF", accent="4169A1")

    doc.add_heading("5.6 六类用户结果与默认规则", level=2)
    result_rows = [
        ["可放量", "SCALE", "成熟；roi_index≥1.15；orders≥max(10, benchmark_orders_median)；ROI 3日趋势≥-10%；Benchmark≥Medium", "创建放量草稿；可选生成衍生变体"],
        ["稳定投放", "HOLD", "成熟；1.00≤roi_index<1.15；ctr_index、cvr_index≥0.90；核心指标波动<15%", "保持配置并创建复查任务"],
        ["待观察", "OBSERVE", "数据不完整、样本未成熟、趋势/基准不可计算，或无其他强结论", "展示缺口并排定重诊断"],
        ["需迭代", "ITERATE", "成熟；roi_index≥0.90，且 CTR/CVR 某一环弱于0.80，或短期趋势≤-20%", "一次只改变一个主变量并生成对照版本"],
        ["需换新", "REFRESH", "成熟；0.70≤roi_index<0.90 且 ctr_index、cvr_index均<0.80；或同方向连续2轮未改善", "选择新创意方向并生成整套新素材"],
        ["建议关停", "STOP", "成熟；roi_index<0.70；spend≥2×target_cpo；连续2个诊断窗口恶化；Benchmark≥Medium", "二次确认后对当前投放关系做可恢复停用/移出"],
    ]
    h.add_table(doc, ["展示结果", "代码", "MVP 默认判定（可配置）", "唯一主动作"], result_rows, [1300, 900, 4900, 2260], keep_rows=False)

    doc.add_heading("5.7 规则执行优先级", level=2)
    h.add_body(doc, "原稿给出的“观察→关停→换新→放量→迭代→稳定”没有说明是展示排序、执行顺序还是冲突优先级。V2.0 定义为安全优先的规则短路：")
    precedence = [
        "数据不合格或样本不成熟 → 待观察。",
        "命中非素材保护 → 跳转商品/投放诊断，锁定素材动作。",
        "满足且连续验证建议关停 → 建议关停。",
        "满足换新 → 需换新。",
        "满足迭代 → 需迭代。",
        "满足放量 → 可放量。",
        "其余成熟样本 → 稳定投放。",
    ]
    for item in precedence:
        add_numbered(doc, item, number_id)
    h.add_body(doc, "每次计算返回 matched_rule_id、rule_version、reason_codes 与 suppressed_rules，便于解释为何某条规则没有成为主结果。")

    page_break(doc)
    doc.add_heading("6. 功能清单与优先级", level=1)
    h.add_body(doc, "优先级按依赖与风险定义：P0 构成可演示且可保存的最小闭环；P1 接通平台执行与真实回流；P2 提升决策质量；P3 扩展组织效率。")
    feature_rows = [
        ["F01", "全局上下文", "P0", "平台/店铺/账户/国家/时间窗统一，输出 data_as_of", "数据接入"],
        ["F02", "经营总览与商品处理队列", "P0", "KPI、趋势、事件、商品聚合与可解释排序", "F01、指标服务"],
        ["F03", "诊断快照与证据抽屉", "P0", "唯一结论、证据、上下文、允许动作", "F01、F04"],
        ["F04", "规则引擎与 Benchmark", "P0", "资格校验、保护分支、六类结果、规则版本", "数据仓库/计算服务"],
        ["F05", "动作意图与统一任务", "P0", "一快照一主动作、状态与审计", "F03、任务存储"],
        ["F06", "方向/变量选择与生成承接", "P0/P1", "P0 跳转并接收结果；P1 多轮生成、跨轮选择", "现有生成能力"],
        ["F07", "投放配置与待发布草稿", "P0", "Target ROI、预算、素材、观察/止损规则，保存草稿", "F05、草稿存储"],
        ["F08", "发布确认与投放状态跟踪", "P1", "幂等发布、创建/审核/投放状态分离、失败重试", "TikTok API、F07"],
        ["F09", "稳定/观察复查任务", "P1", "累积样本、数据恢复后重诊断", "调度/规则服务"],
        ["F10", "任务记录与结果回流", "P1", "时间线、原素材/变体对照、winner、回写", "F08、回流数据"],
        ["F11", "投放中调整建议", "P2", "放量、保持、补素材、调目标、止损建议", "F10、毛利/库存数据"],
        ["F12", "GMV 影响与优先级模型", "P2", "风险/机会估计、置信区间、模型版本", "稳定历史样本"],
        ["F13", "权限、审批、批量与通知", "P3", "角色权限、多级审批、批量任务、消息协作", "组织/权限系统"],
    ]
    h.add_table(doc, ["ID", "功能", "优先级", "交付边界", "依赖"], feature_rows, [700, 2200, 900, 3950, 1610], keep_rows=False)

    page_break(doc)
    doc.add_heading("7. 功能详细需求", level=1)

    doc.add_heading("7.1 F01 全局上下文（P0）", level=2)
    h.add_labeled(doc, "业务目的", "保证经营指标、Benchmark、诊断、草稿和回流使用同一数据范围。")
    for text in [
        "默认：GMV Max、当前店铺、全部有权限账户、当前国家、近7天；日期使用账户时区。",
        "平台单选、店铺单选、账户多选、国家单选、日期预设/自定义；切换任一维度后取消旧请求并整体重算。",
        "所有响应返回 context_id、timezone、currency、data_as_of、attribution_window 与数据延迟状态。",
        "从商品行进入诊断、从诊断保存草稿、从任务查看回流时必须携带原 context_id；若数据范围变化，创建新诊断快照。",
        "加载失败显示错误原因与“重新加载”，不得用旧数据伪装新上下文。",
    ]:
        h.add_bullet(doc, text, bullet_id)
    h.add_labeled(doc, "验收", "跨标签、抽屉、草稿和任务返回后上下文不丢失；所有货币指标统一币种并标明换算时间。")

    doc.add_heading("7.2 F02 经营总览与商品处理队列（P0）", level=2)
    h.add_labeled(doc, "页面结构", "KPI 卡 → 趋势与事件轴 → 今日处理摘要 → 商品问题队列。")
    for text in [
        "KPI：Spend、CTR、CVR、Orders、Gross revenue、ROI；点击卡片可切换趋势，至少保留一个指标。",
        "趋势按日显示 7/14/30 天数据；事件 Chip 展示预算、目标、素材、授权、状态与备注变更，点击打开日志。",
        "商品行展示商品/SKU、GMV/ROI/Target ROI、成熟素材数、Benchmark 摘要、主诊断、证据与唯一 CTA。",
        "P0 排序使用可解释优先级：安全风险 > GMV 规模 > ROI 缺口 > 素材供给缺口 > 数据新鲜度；每个分项输出 contribution。",
        "筛选：全部、增长机会、素材问题、商品/投放问题、待观察；空态应解释无匹配商品而非显示 0 数据。",
    ]:
        h.add_bullet(doc, text, bullet_id)
    add_formula(doc, "P0 优先级分", "priority_score = risk_weight + normalized_gmv × w1 + roi_gap × w2 + supply_gap × w3 - staleness_penalty", "权重进入 rule_version；只用于排序，不作为诊断结论。预计 GMV 金额模型后置 P2。")

    doc.add_heading("7.3 F03 诊断快照与证据抽屉（P0）", level=2)
    h.add_labeled(doc, "业务目的", "回答“数据说明了什么、系统为什么这样判断、下一步是什么”。")
    for text in [
        "抽屉默认显示最新快照，不在打开时实时覆盖；用户可手动“重新诊断”生成新 snapshot。",
        "头部：商品、素材、国家/账户、时间窗、data_as_of、规则版本、四阶段进度。",
        "结论区：结果 Badge、一句话结论、主动作；不得出现两个同等级主按钮。",
        "证据区：实际值、Benchmark、Index、趋势、样本量与置信度；空值解释原因。",
        "受影响素材默认列 3 条并可展开；跨商品使用时必须显示影响范围。",
        "命中非素材保护时，素材生成/关停入口禁用并说明保护原因。",
        "关闭抽屉后保留页面上下文；未提交的本地表单需二次确认。",
    ]:
        h.add_bullet(doc, text, bullet_id)
    h.add_labeled(doc, "计算", "读取 F04 生成的 diagnosis_snapshot，不在前端重算结果。")

    doc.add_heading("7.4 F04 规则引擎与 Benchmark（P0）", level=2)
    for text in [
        "输入：不可变 metric_snapshot、product_target、benchmark_snapshot、data_quality、platform status。",
        "输出：eligibility、attribution_scope、display_result、primary_action、reason_codes、evidence、rule_version、suppressed_rules。",
        "所有阈值可配置、版本化且可灰度；历史快照永远引用当时版本，不随新规则重算。",
        "规则计算失败时返回 UNKNOWN/待观察与 error_code，不允许前端兜底判定。",
        "相同幂等键 context+creative+window+rule_version 重复请求返回同一快照。",
    ]:
        h.add_bullet(doc, text, bullet_id)
    h.add_labeled(doc, "验收", "构造边界值、空值、低样本、低置信度与冲突规则测试；优先级短路结果可解释且可复现。")

    doc.add_heading("7.5 F05 动作意图与统一任务（P0）", level=2)
    h.add_body(doc, "用户接受诊断后创建 action_intent 与 task。一个快照只允许一个 active intent；改判或放弃通过状态变更完成，不物理删除。")
    intent_rows = [
        ["SCALE", "创建放量方案", "current creative、预算建议、Target ROI、证据"],
        ["HOLD", "保持并复查", "复查时间、指标门槛、异常条件"],
        ["OBSERVE", "继续观察", "缺失样本、预计满足时间、重诊断时间"],
        ["ITERATE", "选择一个迭代变量", "变量轴、锁定项、参考素材"],
        ["REFRESH", "选择新方向", "方向标签、保留约束、参考素材"],
        ["STOP", "确认停用/移出", "影响商品/账户、可恢复说明、二次确认"],
    ]
    h.add_table(doc, ["结果", "主动作", "必须持久化"], intent_rows, [1500, 2600, 5260])

    doc.add_heading("7.6 F06 方向/变量选择与生成承接（P0/P1）", level=2)
    h.add_labeled(doc, "P0", "将 action_intent、商品上下文和已选策略带入现有生成能力；接收生成任务 ID、结果素材 ID 和版本关系。")
    variable_rows = [
        ["换 Hook", "CTR", "前3秒文案、首帧、节奏整体变更", "正文/卖点/Offer/CTA"],
        ["换出镜达人", "CTR", "人设/人物，脚本结构不变", "脚本结构/卖点/Offer"],
        ["换场景", "CTR", "拍摄环境与氛围", "人物/台词/Offer"],
        ["换卖点与证明", "CVR", "主卖点及支撑证据形式", "Hook/人物/Offer"],
        ["换 Offer 与 CTA 呈现", "CVR", "利益表达与行动指令的呈现", "Hook/正文/价格本身"],
    ]
    h.add_table(doc, ["变量", "目标指标", "允许改变", "默认锁定"], variable_rows, [1800, 1200, 3300, 3060])
    for text in [
        "需迭代默认一次只改变一个主变量；多选时必须提示无法归因到单一变量，并将实验类型标为 multivariate。",
        "可放量允许不选变量，直接复用当前获胜素材创建草稿；生成衍生变体是可选次动作。",
        "需换新至少选择一个新方向，不沿用已连续失败的方向标签。",
        "P1 支持多轮生成、跨轮选择、每个方案最多 5 条入选素材；所有结果必须可完整播放。",
        "生成失败保留已完成轮次与任务上下文，可重试但不重复生成成功项。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("7.7 F07 投放配置与待发布草稿（P0）", level=2)
    for text in [
        "配置页与生成结果独立，固定展示来源商品、原素材、诊断快照与入选新素材。",
        "字段：Target ROI、每日预算、素材清单、首轮观察时长、胜出订单线、止损 ROI。",
        "Target ROI 默认继承商品目标；每日预算展示诊断建议与可编辑范围，不接受素材自身出价。",
        "校验：Target ROI>0、预算>0、至少1条可投素材、观察时长>0、胜出订单线>0、0<止损ROI<Target ROI。",
        "保存仅写入 delivery_draft，状态为 DRAFT/PENDING_CONFIRM；不直接调用平台接口。",
        "重复保存使用幂等键 diagnosis_snapshot_id+intent_version，若已有草稿则覆盖同草稿版本或提示。",
    ]:
        h.add_bullet(doc, text, bullet_id)
    h.add_labeled(doc, "素材覆盖健康", "coverage_rate = 达到最低可投素材数的在投商品数 / 在投商品总数；同时展示 avg_eligible_creatives_per_product。最低素材数由配置决定，原稿中的“6 条”不能写死。")

    doc.add_heading("7.8 F08 发布确认与投放状态跟踪（P1）", level=2)
    h.add_body(doc, "发布前展示变更摘要和 Before/After。点击确认后创建 execution，不删除草稿；平台写入需幂等、可重试、可审计。")
    state_rows = [
        ["创建阶段", "DRAFT → CREATING → CREATE_SUCCEEDED / CREATE_FAILED", "返回平台 ID 只代表创建阶段成功"],
        ["投放阶段", "PENDING_REVIEW / SCHEDULED / ACTIVE / LIMITED_DELIVERY / PAUSED / NOT_DELIVERING / ENDED / DELETED / RE_REVIEWING", "应用归一化状态；同时保留平台原始状态和原因"],
        ["异常", "UNKNOWN / SYNC_DELAYED", "保留最后成功状态与 data_as_of，不伪造 ACTIVE"],
    ]
    h.add_table(doc, ["状态域", "建议状态", "说明"], state_rows, [1600, 5200, 2560], keep_rows=False)
    for text in [
        "必须保存 normalized_status、platform_status、status_reason、status_message、updated_at。",
        "CREATE_SUCCEEDED 不得在 UI 显示为“投放中”；只有平台生命周期为 ACTIVE 才显示投放中。",
        "失败时保留草稿、request_id、platform_response 与可重试入口；同一幂等键不得创建重复计划。",
        "平台原始状态枚举、发布接口权限和字段映射需接入前用官方文档/沙箱重新确认。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("7.9 F09 稳定/观察复查任务（P1）", level=2)
    for text in [
        "HOLD 创建固定复查时间和触发门槛；默认 72h 仅为建议值，可配置。",
        "OBSERVE 记录未满足条件（数据延迟、active_hours、orders、spend、benchmark sample），展示预计满足时间。",
        "数据条件满足后创建新 snapshot，不覆盖旧结论；任务时间线关联前后快照。",
        "连续同步失败转为 BLOCKED 并提示重试；不得丢弃任务。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("7.10 F10 任务记录与结果回流（P1）", level=2)
    h.add_labeled(doc, "任务列表", "展示生成中、待发布、审核/投放中、观察中、已验证、失败与有阻塞；支持商品、状态、时间搜索。")
    h.add_labeled(doc, "任务详情", "按诊断 → 方向/生成 → 草稿 → 发布 → 状态 → 观测 → 回流显示完整时间线，人工修改记录操作者与时间。")
    add_formula(doc, "验收门槛", "validation_ready = observation_hours ≥ configured_hours AND orders ≥ win_orders AND data_complete", "两项样本条件同时满足；仅时间到达不能验收。")
    for text in [
        "对照优先使用同商品、同上下文、重叠观测窗口的原素材与变体；不满足可比性时标“不可直接比较”。",
        "winner 默认优先看 ROI 达成，再比较 Orders/Spend 稳定性；CTR/CVR 用于解释，不单独决定商业胜出。",
        "回流写入 validation_result，并创建下一轮诊断输入；历史结果不可修改，只能追加更正记录。",
        "若创意同时影响多个商品，按各商品上下文分别验证，不用全局平均代替。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    doc.add_heading("7.11 F11–F13 后续能力（P2/P3）", level=2)
    future_rows = [
        ["F11 调整建议", "P2", "放量、保持、补素材、调目标、止损；所有建议展示触发条件、证据、Before/After 和风险护栏，仍需 AO 确认"],
        ["F12 影响模型", "P2", "预测金额必须标 estimated，输出区间、confidence、model_version；模型不可用时回退可解释排序分"],
        ["F13 组织效率", "P3", "多人审批、批量动作、自定义规则、跨店铺对比、通知与协作；不阻塞 P0/P1 闭环"],
    ]
    h.add_table(doc, ["功能", "优先级", "边界"], future_rows, [2000, 1200, 6160])

    page_break(doc)
    doc.add_heading("8. 数据实体与接口契约", level=1)
    doc.add_heading("8.1 核心数据实体", level=2)
    entity_rows = [
        ["ContextSnapshot", "context_id, platform, shop_id, account_ids, country, timezone, currency, date_range, data_as_of", "冻结查询上下文"],
        ["MetricSnapshot", "spend, impressions, clicks, orders, gross_revenue, ctr, cvr, aov, roi, active_hours, data_quality", "只读指标快照"],
        ["BenchmarkSnapshot", "level, filters, window, sample_size, median metrics, confidence, calculated_at", "可解释基准"],
        ["DiagnosisSnapshot", "id, product_id, creative_id, context_id, metrics, benchmark, indexes, result, reason_codes, rule_version", "不可变诊断结论"],
        ["ActionIntent", "id, diagnosis_id, action_type, strategy, locked_fields, status, version", "唯一主动作"],
        ["GenerationBatch", "intent_id, round, source_assets, selected_strategy, outputs, status", "生成轮次与素材谱系"],
        ["DeliveryDraft", "intent_id, target_roi, daily_budget, creative_ids, observe_hours, win_orders, stop_roi, status", "待发布配置"],
        ["DeliveryExecution", "draft_id, idempotency_key, platform_ids, create_status, delivery_status, reason, raw_status", "平台执行与状态"],
        ["ValidationResult", "execution_id, window, baseline, variant, winner, confidence, decision", "观测验收"],
        ["AuditLog", "operator, action, before, after, request_id, platform_response, created_at", "操作审计"],
    ]
    h.add_table(doc, ["实体", "关键字段", "用途"], entity_rows, [2000, 4800, 2560], keep_rows=False)

    doc.add_heading("8.2 建议接口边界", level=2)
    api_rows = [
        ["GET /insights/overview", "context", "KPI、趋势、事件、商品队列、data_as_of"],
        ["POST /diagnoses", "context+product+creative+rule_version", "创建/返回幂等诊断快照"],
        ["POST /action-intents", "diagnosis_id+action", "创建唯一主动作"],
        ["POST /generation-batches", "intent_id+strategy", "承接生成任务；返回 batch_id"],
        ["POST /delivery-drafts", "intent_id+config", "保存/更新草稿版本"],
        ["POST /delivery-executions", "draft_id+idempotency_key", "确认发布并返回 execution_id"],
        ["GET /delivery-executions/{id}", "—", "创建状态、投放状态、原始状态与原因"],
        ["POST /diagnoses/{id}/recheck", "trigger", "基于新窗口创建新快照"],
        ["GET /tasks/{id}", "—", "任务时间线、阻塞与结果"],
        ["POST /validations", "execution_id+window", "计算并持久化验收结果"],
    ]
    h.add_table(doc, ["接口", "关键输入", "输出/约束"], api_rows, [3200, 2800, 3360], keep_rows=False)
    h.add_callout(doc, "接口原则", "正式 URL、TikTok 字段和平台枚举由技术方案确定。本 PRD 约束的是领域边界、幂等、状态分离和审计字段，而不是要求照搬示例路由。", fill="F4F6F9", accent=h.BLUE)

    doc.add_heading("8.3 错误与降级", level=2)
    error_rows = [
        ["数据延迟/缺失", "诊断=待观察；展示延迟时间与重试，不沿用旧数据冒充当前值"],
        ["Benchmark 不足", "降低 confidence；强动作禁用；显示回退层级"],
        ["规则服务失败", "返回 UNKNOWN/待观察；保留 request_id；不在前端判定"],
        ["生成部分失败", "保留成功项与轮次；失败项可幂等重试"],
        ["保存草稿失败", "保留表单与输入；允许重试；不显示已保存"],
        ["发布失败", "草稿保留；execution=CREATE_FAILED；显示平台原因与重试"],
        ["状态同步异常", "显示最后成功状态+SYNC_DELAYED；禁止误报 ACTIVE"],
        ["回流不可比", "Validation=INCONCLUSIVE；说明窗口/样本/上下文差异"],
    ]
    h.add_table(doc, ["场景", "产品行为"], error_rows, [2600, 6760])

    page_break(doc)
    doc.add_heading("9. 状态机", level=1)
    doc.add_heading("9.1 诊断与动作", level=2)
    status_rows = [
        ["Diagnosis", "COMPUTING → READY / OBSERVE / NON_MATERIAL_PROTECTED / FAILED", "READY 内含六类展示结果；OBSERVE 可定时重诊断"],
        ["Action intent", "DRAFT → CONFIRMED → IN_PROGRESS → COMPLETED / CANCELLED / BLOCKED", "同一诊断仅一个 active intent"],
        ["Generation", "QUEUED → GENERATING → PARTIAL / SUCCEEDED / FAILED / CANCELLED", "多轮追加，不覆盖成功历史"],
        ["Delivery draft", "DRAFT → PENDING_CONFIRM → CONFIRMED → EXECUTING → ARCHIVED", "草稿与平台执行分离"],
        ["Validation", "WAITING_DATA → READY → EVALUATING → WON / LOST / INCONCLUSIVE / FAILED", "满足时间+订单双门槛后 READY"],
    ]
    h.add_table(doc, ["对象", "状态", "关键规则"], status_rows, [1800, 4400, 3160], keep_rows=False)

    doc.add_heading("9.2 关停的可恢复性", level=2)
    h.add_body(doc, "建议关停默认作用于当前 delivery relation，不删除 creative 实体。确认弹窗必须列出受影响商品、账户、计划、当前消耗与素材供给 Before/After；成功后记录 REMOVED/PAUSED 与恢复入口。跨商品共享素材时，只修改用户确认的关系。")

    doc.add_heading("9.3 审计要求", level=2)
    for text in [
        "所有人工确认动作记录 operator、role、timestamp、reason。",
        "所有写操作记录 before、after、request_id、idempotency_key、platform_response。",
        "诊断与规则变更记录 rule_version；模型预测记录 model_version。",
        "失败重试追加 attempt，不覆盖第一次失败证据。",
        "删除/撤销使用软删除或状态变更，不清除历史。",
    ]:
        h.add_bullet(doc, text, bullet_id)

    page_break(doc)
    doc.add_heading("10. 研发分期与依赖", level=1)
    phase_rows = [
        ["P0：可解释最小闭环", "F01–F05、F06 基础承接、F07", "统一口径与上下文；商品队列；诊断快照；规则与 Benchmark；动作任务；跳转生成并接收结果；保存待发布草稿", "不调用平台发布；用明确 mock/沙箱状态演示到草稿"],
        ["P1：真实执行与回流", "F06 完整、F08–F10、审计/重试", "多轮结果；人工确认发布；创建/审核/投放状态；复查任务；结果回流与比较", "平台权限、回流延迟、归因一致性"],
        ["P2：决策增强", "F11–F12", "投放调整建议、毛利/库存护栏、GMV 影响模型、跨商品供给", "模型验证、置信区间与安全阈值"],
        ["P3：规模化协作", "F13", "审批、批量、自定义规则、跨店铺、通知", "组织与权限体系"],
    ]
    h.add_table(doc, ["阶段", "范围", "完成定义", "主要风险"], phase_rows, [1700, 1900, 3760, 2000], keep_rows=False)

    doc.add_heading("10.1 P0 技术拆解建议", level=2)
    for text in [
        "先冻结指标字典、诊断主键、rule_version 和 context_id，再实现页面。",
        "将当前前端固定 ROI 阈值迁出 UI，建立单一诊断服务/领域函数与边界测试。",
        "建立 DiagnosisSnapshot、ActionIntent、DeliveryDraft、Task、AuditLog 最小表结构。",
        "完成商品聚合队列与诊断抽屉，再接现有素材生成入口；避免先做六套详情页。",
        "P0 以“保存待发布草稿 + 可恢复任务记录”为终点；平台发布不得用假成功替代。",
    ]:
        add_numbered(doc, text, number_id)

    doc.add_heading("10.2 关键依赖", level=2)
    dep_rows = [
        ["TikTok Ads/Shop 数据", "指标、商品-素材关系、账户/投放状态", "确认权限、归因窗口、更新频率、币种与退款口径"],
        ["Benchmark 服务", "样本选择、中位数、置信度", "支持分层回退和快照化"],
        ["创意生成", "策略输入、批次、结果与版本关系", "支持幂等任务与完整播放结果"],
        ["投放草稿/发布", "目标、预算、素材清单、平台写入", "确认平台可写字段及账户权限"],
        ["任务调度", "复查、重诊断、状态同步", "支持延迟任务、重试与阻塞状态"],
        ["审计日志", "人工/平台动作", "不可变、可检索、脱敏"],
    ]
    h.add_table(doc, ["依赖", "用途", "上线前确认"], dep_rows, [2100, 3000, 4260])

    page_break(doc)
    doc.add_heading("11. 验收标准", level=1)
    acceptance_rows = [
        ["AC01", "上下文", "切换店铺/账户/国家/日期后，KPI、队列、诊断、Benchmark 同步刷新；跨页返回保持 context"],
        ["AC02", "诊断粒度", "同一 creative 在不同 product/account 下产生独立快照，不使用全局结论"],
        ["AC03", "空值", "分母为0/数据缺失时指标为 null 并解释，不能显示为0或参与 Index"],
        ["AC04", "样本门槛", "未同时满足数据完整、24h 与订单/消耗门槛时只可待观察"],
        ["AC05", "Benchmark", "样本不足按 L1→L4 回退并显示 level/sample/confidence；低置信度禁用强动作"],
        ["AC06", "保护分支", "CTR/CVR 正常且 ROI 不达标时锁定素材动作并跳转商品/投放诊断"],
        ["AC07", "唯一结果", "每个快照只有一个展示结果、一个主按钮和一个 active action intent"],
        ["AC08", "规则可追溯", "结果展示 reason_codes、rule_version、data_as_of；历史快照不随新规则改变"],
        ["AC09", "生成", "需迭代只改一个主变量；多轮成功结果不被后续失败覆盖；最多5条入选"],
        ["AC10", "草稿校验", "至少1条可投素材，0<stop_roi<target_roi，其余必填值有效；保存不触发发布"],
        ["AC11", "发布安全", "发布前二次确认；重复请求不重复创建；失败保留草稿和 request_id"],
        ["AC12", "状态分离", "CREATE_SUCCEEDED 不显示投放中；只有 ACTIVE 显示投放中，原始状态/原因可查"],
        ["AC13", "关停可恢复", "仅修改确认范围内的投放关系，展示 Before/After，并提供恢复入口"],
        ["AC14", "回流", "观察时间与订单线同时满足才验收；不可比数据输出 INCONCLUSIVE"],
        ["AC15", "审计", "预算、目标、素材、关停、发布均记录 operator/before/after/request_id/platform_response"],
        ["AC16", "异常", "数据/规则/平台失败均有明确错误、重试和阻塞状态，不用旧值或 mock 值伪装成功"],
    ]
    h.add_table(doc, ["编号", "范围", "可测试验收条件"], acceptance_rows, [900, 1800, 6660], keep_rows=False)

    doc.add_heading("11.1 关键边界测试", level=2)
    edge_rows = [
        ["roi_index=0.70/0.90/1.00/1.15", "验证包含/不包含边界与规则优先级"],
        ["active_hours=23.99/24，orders=4/5", "验证样本门槛，不提前强判"],
        ["spend=2×target_cpo 前后", "验证门槛与浮点精度"],
        ["benchmark sample=4/5/9/10", "验证回退级别和置信限制"],
        ["CTR/CVR 正常、AOV/库存异常", "验证非素材保护"],
        ["同一素材跨两个商品", "验证关停/验证不串上下文"],
        ["发布超时后重试", "验证幂等与未知状态查询"],
        ["创建成功但审核中/不投放", "验证状态不误报 ACTIVE"],
        ["仅观察时间到、订单不足", "验证不能验收"],
    ]
    h.add_table(doc, ["测试数据", "验证点"], edge_rows, [3600, 5760])

    page_break(doc)
    doc.add_heading("12. 待确认事项与风险", level=1)
    decision_rows = [
        ["D01", "平台所谓 Target ROI 的精确口径、退款/取消处理与自然成交归因", "影响 ROI、AOV、Target CPO 及所有 Index", "数据+投放"],
        ["D02", "商品-素材-账户/计划可获得的最细关系与历史窗口", "决定诊断主键和 Benchmark 可比性", "数据+后端"],
        ["D03", "P0 默认阈值与连续窗口定义", "决定误判率；必须用历史数据回放验证", "业务+算法"],
        ["D04", "Expected AOV 取值：稳定期中位数、商品长期值或平台值", "影响 target_cpo 与成熟门槛", "数据+业务"],
        ["D05", "平台发布可写字段、审核状态和原始状态枚举", "影响 F08 接口与状态映射", "后端+投放"],
        ["D06", "“移出素材池/暂停/删除”在平台的实际可恢复动作", "影响关停安全与跨商品影响", "投放+后端"],
        ["D07", "生成模块可否接收变量锁定、批次、轮次与素材谱系", "影响 F06 P0/P1 范围", "生成团队"],
        ["D08", "Winner 的统计置信标准与最小订单数", "影响回流的可信度", "数据+算法"],
        ["D09", "素材覆盖健康的最低可投素材数", "原稿的6条不可直接写死", "投放业务"],
        ["D10", "GMV 影响预测是否具备训练样本和可解释性", "不足则保持 P2，不在 P0 展示金额", "算法+产品"],
    ]
    h.add_table(doc, ["编号", "待确认", "不确认的影响", "Owner"], decision_rows, [850, 4300, 3000, 1210], keep_rows=False)

    h.add_callout(doc, "上线门槛", "D01–D06 未确认前，只能完成 P0 的可解释诊断与草稿闭环；不得把模拟阈值、假平台状态或演示发布当作生产可用能力。", fill="FFECEC", accent=h.RISK)

    page_break(doc)
    doc.add_heading("附录 A：原型参考（非最终交互规范）", level=1)
    h.add_body(doc, "以下图片来自原始 AI 文档，用于说明信息架构和交互意图。最终页面应以本 PRD 的业务规则、状态机、空值和安全约束为准，不继承截图中的 Mock 数字或未定义状态。")
    add_source_image(doc, "overview", "图 A-1  原稿：经营总览与商品处理队列", "GMV Max 经营总览原型，包含 KPI、趋势和商品问题队列")
    add_source_image(doc, "diagnosis", "图 A-2  原稿：商品经营诊断抽屉", "商品经营诊断抽屉原型，包含诊断结论、证据和受影响素材")
    add_source_image(doc, "delivery", "图 A-3  原稿：投放草稿确认", "GMV Max 投放草稿详情原型，包含目标、预算、素材和观察止损规则")
    add_source_image(doc, "callback", "图 A-4  原稿：任务记录与原素材/变体回流", "原素材与变体回流对照表原型")

    doc.add_heading("附录 B：本版与原稿的范围映射", level=1)
    mapping_rows = [
        ["原稿 1–2：定位/菜单", "第 2–4 章", "保留目标与四标签，重构为统一闭环"],
        ["原稿 3：经营诊断", "第 5、7.1–7.4", "补齐诊断粒度、空值、Benchmark 与规则优先级"],
        ["原稿 4：六类执行流", "第 4、5、7.5–7.7", "六流程合并为主流程与三类动作分支"],
        ["原稿 5：投放中心", "第 7.7–7.8、9 章", "拆分草稿、创建与投放生命周期"],
        ["原稿 6：任务回流", "第 7.9–7.10", "明确时间+订单双门槛与不可比结果"],
        ["原稿 8：验收", "第 11 章", "改成可测试边界与异常行为"],
        ["原稿 9：优先级", "第 6、10 章", "按依赖与风险重排 P0–P3"],
    ]
    h.add_table(doc, ["原稿范围", "本版章节", "关键变化"], mapping_rows, [2500, 2200, 4660])

    doc.add_heading("文档结束", level=1)
    h.add_body(doc, "建议下一步：由产品、数据与投放业务先完成 D01–D06 决策，再由工程基于 P0 实体与接口边界输出技术方案和拆分排期。", bold=True, color=h.INK)

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
