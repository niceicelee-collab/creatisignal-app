from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Cursor\creatisignal-app-new")
OUTPUT = ROOT / "docs" / "prd" / "产品库需求文档_v1.0_2026-08-04.docx"
CONTENT_DXA = 9360

BLUE = "2E74B5"
DARK_BLUE = "1F4E79"
TEXT = "1F2937"
MUTED = "667085"
LIGHT = "F2F4F7"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "ECFDF3"
GREEN = "027A48"
LIGHT_AMBER = "FFFAEB"
AMBER = "B54708"
WHITE = "FFFFFF"
BORDER = "D0D5DD"


def set_east_asia_font(run, east_asia: str = "Microsoft YaHei"):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ["sz", "val", "color", "space"]:
            if key in kwargs[edge]:
                element.set(qn("w:" + key), str(kwargs[edge][key]))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int], indent: int = 120):
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[min(idx, len(widths) - 1)]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header_and_style(table, header_fill: str = LIGHT):
    if not table.rows:
        return
    repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, header_fill)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(TEXT)
    border = {"val": "single", "sz": "4", "color": BORDER, "space": "0"}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_east_asia_font(run)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 页")
    set_east_asia_font(tail)


def hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    link.append(new_run)
    paragraph._p.append(link)


def add_paragraph(doc, text: str = "", style: str | None = None, *, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_east_asia_font(r1)
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia_font(r2)
    else:
        r = p.add_run(text)
        set_east_asia_font(r)
    return p


def add_bullets(doc, items: Iterable[str], level: int = 0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.50 + 0.25 * level)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.10
        r = p.add_run(item)
        set_east_asia_font(r)


def add_numbered(doc, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.50)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.10
        r = p.add_run(item)
        set_east_asia_font(r)


def add_callout(doc, title: str, body: str, fill: str = LIGHT_BLUE, accent: str = BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, left={"val": "single", "sz": "20", "color": accent, "space": "0"})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    set_east_asia_font(r)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    set_east_asia_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int], header_fill: str = LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = value
    set_table_geometry(table, widths)
    set_repeat_table_header_and_style(table, header_fill)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    set_east_asia_font(r)
                    r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_page_break()


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in (("Title", 26, DARK_BLUE), ("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(13.5)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    doc.core_properties.title = "产品库需求文档"
    doc.core_properties.subject = "商品库与商品智能识别工程评审"
    doc.core_properties.author = "CreatiSignal Product Team"
    doc.core_properties.keywords = "产品库, 商品识别, 图片解析, Qwen3.7-Plus, 工程评审"
    doc.core_properties.comments = "2026-08-04 v1.0"

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "CreatiSignal  ·  Product Requirements"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in hp.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        set_east_asia_font(r)
    fp = section.footer.paragraphs[0]
    fp.add_run("内部评审材料  ·  v1.0  ·  ")
    for r in fp.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        set_east_asia_font(r)
    add_page_field(fp)

    # Cover / memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("产品需求文档 · ENGINEERING REVIEW")
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    set_east_asia_font(r)

    title = doc.add_paragraph(style="Title")
    title.add_run("产品库与商品智能识别").bold = True
    for r in title.runs:
        set_east_asia_font(r)
    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("统一沉淀商品信息，降低重复输入，让爆款成片、高保真复刻与增强复刻共享同一数据底座")
    for r in sub.runs:
        set_east_asia_font(r)

    meta = add_table(
        doc,
        ["文档信息", "内容"],
        [
            ["版本 / 状态", "v1.0 / 工程评审稿"],
            ["计划范围", "2026 年 8 月 MVP；W1 启动"],
            ["面向客户", "电商客户、小微组织与小型内容团队"],
            ["核心角色", "运营、设计师、独立卖家"],
            ["评审对象", "产品、前端、后端、算法、测试、设计、安全与数据"],
            ["负责人 / 日期", "待补充 / 2026-08-04"],
        ],
        [1900, 7460],
        LIGHT_BLUE,
    )

    add_callout(
        doc,
        "评审结论建议",
        "MVP 先打通“链接或图片导入 → AI 结构化识别 → 用户确认 → 入库 → 下游复用”。图片识别主模型建议采用 Qwen3.7-Plus（非思考模式 + JSON Schema），Gemini 作为海外/英文商品与质量对照备选，字节视觉模型作为国内链路 A/B 候选。",
        LIGHT_GREEN,
        GREEN,
    )

    add_heading(doc, "一页结论", 1)
    add_bullets(doc, [
        "产品价值：把重复输入和重复上传变成一次确认、多处复用，缩短从商品素材到视频生成的链路。",
        "MVP 核心：链接识别、图片识别、人工确认、商品库列表/详情、历史商品一键复用、下游快照引用。",
        "8 月相邻交付：数字人库扩充至 30–50 个，优先覆盖欧美、东南亚、南美；支持用户上传自有数字人。",
        "成功衡量：上线后观察产品导入成功率、确认后入库率、复用率、首条视频生成时长、WAU 与周视频增量。",
    ])
    add_callout(doc, "业务目标（来自需求清单）", "预计带来 WAU +500、周视频生成量 +5,000。该数值作为业务目标，不作为单周技术验收门槛。", LIGHT_AMBER, AMBER)

    page_break(doc)

    # 1-3 Context
    add_heading(doc, "1. 背景与问题", 1)
    add_paragraph(doc, "当前用户在不同创作链路中反复填写商品名称、描述、卖点、受众、场景与价格信息，并重复上传同一组商品图片。商品信息没有稳定的团队级载体，导致创作流程过重、上下文易丢失、历史项目难复用；即使用户仍有积分，也可能因为准备成本过高而放弃继续生成。")
    add_paragraph(doc, "本期将“商品库”建设为 SaaS 底座：从商品链接或图片中提取可审核的结构化信息，用户确认后入库，并以版本快照方式提供给爆款成片、高保真复刻、增强复刻等功能。")

    add_heading(doc, "2. 目标、指标与非目标", 1)
    add_heading(doc, "2.1 产品目标", 2)
    add_numbered(doc, [
        "一次导入：支持链接或 1–8 张图片创建商品草稿。",
        "一次确认：AI 结果必须可见、可编辑、可追溯；未经确认不得作为正式商品进入下游。",
        "多处复用：同一商品可在多个创作任务中一键选择，不再重复填写与上传。",
        "稳定引用：下游任务引用不可变快照，避免商品后续编辑影响历史项目。",
    ])
    add_heading(doc, "2.2 核心指标", 2)
    add_table(doc, ["指标", "口径", "MVP 观察目标"], [
        ["导入成功率", "进入识别任务后产出可审核结果的任务数 / 发起任务数", "≥ 95%（支持渠道内）"],
        ["确认后入库率", "完成确认并保存的商品数 / 产出识别结果的商品数", "≥ 80%"],
        ["字段人工修改率", "被用户修改的字段数 / AI 有值字段数", "用于评估模型，不设硬性门槛"],
        ["30 天商品复用率", "至少被第 2 个创作任务引用的商品 / 活跃商品", "≥ 30%"],
        ["首条视频准备时长", "进入创作到完成商品信息确认的中位时长", "较现状下降 ≥ 40%"],
        ["业务影响", "WAU 与周视频生成量增量", "WAU +500；周视频 +5,000"],
    ], [2200, 4560, 2600])
    add_heading(doc, "2.3 非目标", 2)
    add_bullets(doc, [
        "不建设完整 PIM/ERP，不承担库存、订单、采购与实时价格同步。",
        "不在本期支持全网任意网站的登录态抓取、验证码绕过或反爬规避。",
        "不在本期训练专属识别模型；先通过通用视觉模型、Schema 与评测集验证价值。",
        "数字人上传仅做素材入库与合规校验，不包含声音克隆、人物训练和自动授权。",
    ])

    add_heading(doc, "3. 用户与核心场景", 1)
    add_table(doc, ["角色", "主要任务", "当前痛点", "期望结果"], [
        ["运营", "批量准备商品与视频任务", "信息反复填写、团队资料分散", "商品可搜索、可复用、可批量进入创作"],
        ["设计师", "基于图片提炼视觉卖点与场景", "需要手工理解商品，容易遗漏", "先获得结构化草稿，再人工校正"],
        ["独立卖家", "快速从商品页生成营销视频", "时间有限、非专业文案能力不足", "粘贴链接或上传图片即可完成商品建档"],
    ], [1400, 2500, 2860, 2600])

    page_break(doc)

    # 4 Scope
    add_heading(doc, "4. 范围与优先级", 1)
    add_table(doc, ["优先级", "能力", "本期定义", "验收产物"], [
        ["P0", "商品链接识别", "提取商品名、描述、卖点、目标人群、使用场景、原价、现价与优惠信息", "可编辑识别结果 + 原始来源"],
        ["P0", "商品图片识别", "从 1–8 张图片识别描述、卖点、目标人群与使用场景；可补充可见文字/品牌/品类", "结构化结果 + 字段证据"],
        ["P0", "确认与修改", "逐字段查看、编辑、清空、恢复 AI 值，并明确保存状态", "确认后生成正式商品"],
        ["P0", "商品库与复用", "名称、缩略图、搜索、筛选、详情、编辑、归档、历史商品一键复用", "商品列表、详情与选择器"],
        ["P0", "下游数据底座", "向爆款成片、高保真复刻、增强复刻输出不可变商品快照", "统一选择器与快照 ID"],
        ["P1", "数字人库扩充", "预置 30–50 个数字人，优先覆盖欧美、东南亚、南美", "可筛选数字人列表"],
        ["P1", "自有数字人", "用户上传并形成团队自有数字人库，包含权利确认与审核状态", "上传、审核、选择"],
    ], [900, 1800, 4440, 2220])

    add_heading(doc, "4.1 端到端流程", 2)
    add_callout(doc, "主流程", "选择来源（链接 / 图片） → 创建识别任务 → 解析与视觉识别 → 生成字段级草稿 → 用户确认校正 → 保存商品 → 下游选择商品 → 固化商品快照 → 开始创作", LIGHT_BLUE, BLUE)
    add_heading(doc, "4.2 范围追踪", 2)
    add_table(doc, ["需求清单项", "PRD 对应章节", "责任域"], [
        ["商品链接识别", "5.1、9、11", "后端 / 算法 / 前端"],
        ["商品图片识别", "5.2、8、9", "算法 / 后端 / 前端"],
        ["人工确认与修改", "5.3、6", "前端 / 后端"],
        ["商品库与历史复用", "5.4、7、10", "前端 / 后端"],
        ["数字人库扩充与自定义上传", "5.6、7.4、12", "内容运营 / 后端 / 前端 / 安全"],
        ["下游通用商品底座", "5.5、10", "各创作业务 / 后端"],
    ], [4100, 2460, 2800])

    add_heading(doc, "5. 功能需求", 1)
    add_heading(doc, "5.1 FR-LINK：从商品链接创建", 2)
    add_table(doc, ["ID", "要求", "验收说明"], [
        ["FR-LINK-01", "用户粘贴 HTTP/HTTPS 商品链接并发起识别。", "前端校验格式；后端再次校验并防止 SSRF。"],
        ["FR-LINK-02", "优先解析页面结构化数据、Open Graph、可见文本、图片与价格信息。", "保留来源 URL、采集时间与原始快照摘要。"],
        ["FR-LINK-03", "输出名称、描述、卖点、受众、场景、原价、现价、币种、优惠信息。", "无法确定的字段为 null，不得编造。"],
        ["FR-LINK-04", "遇到登录、验证码、反爬或不支持站点时提供图片上传和手工录入。", "不得无限重试；错误信息可操作。"],
        ["FR-LINK-05", "同团队内按规范化 URL 与来源哈希提示可能重复商品。", "允许继续创建或跳转已有商品，不强制合并。"],
    ], [1420, 4900, 3040])

    page_break(doc)

    add_heading(doc, "5.2 FR-IMAGE：从商品图片创建", 2)
    add_table(doc, ["ID", "要求", "验收说明"], [
        ["FR-IMAGE-01", "支持一次上传 1–8 张 JPG/PNG/WebP；单张建议 ≤ 10 MB。", "数量与大小为 MVP 假设，工程评审确认后固化。"],
        ["FR-IMAGE-02", "识别并生成商品描述、核心卖点、目标人群、使用场景。", "字段必须符合 JSON Schema，显示置信度与证据。"],
        ["FR-IMAGE-03", "允许补充识别商品名、品牌、品类与图片可见文字。", "只提取可见信息；硬规格不确定时返回 null。"],
        ["FR-IMAGE-04", "多图需要跨图去重和合并，主图由用户确认。", "不得把包装配件误当成独立核心商品。"],
        ["FR-IMAGE-05", "同一图片包含多个商品时，提示用户选择主商品或拆分创建。", "禁止模型静默选择。"],
        ["FR-IMAGE-06", "上传后异步处理，支持取消、失败重试与改为手工录入。", "刷新页面后仍可恢复任务状态。"],
    ], [1420, 4900, 3040])

    add_heading(doc, "5.3 FR-REVIEW：确认和校正", 2)
    add_bullets(doc, [
        "识别结果页按“基础信息、营销信息、价格优惠、来源证据”分组；必填字段缺失时明确提示。",
        "每个 AI 字段支持编辑、清空、恢复 AI 原值；展示低置信度提示，不以总分替代字段级判断。",
        "保存草稿不要求完整；点击“确认并入库”时校验商品名、主图、描述，以及至少 1 条卖点。",
        "保存正式商品时记录确认人、确认时间、模型提供方、模型版本、Prompt 版本和源文件哈希。",
        "模型重新识别不得覆盖用户编辑；新结果以对比方式呈现，由用户选择是否采纳。",
    ])

    add_heading(doc, "5.4 FR-LIB：商品库列表、详情与复用", 2)
    add_table(doc, ["ID", "要求", "验收说明"], [
        ["FR-LIB-01", "列表显示主图缩略图、商品名、品类、状态、更新时间与最近使用时间。", "默认按最近使用/更新时间排序。"],
        ["FR-LIB-02", "支持名称/品牌/品类搜索，以及状态、来源、更新时间筛选。", "搜索结果可直接进入详情或选择。"],
        ["FR-LIB-03", "商品支持编辑、复制、归档、恢复；不做物理删除。", "被历史项目引用的商品仍可追溯。"],
        ["FR-LIB-04", "创作入口提供统一商品选择器，支持最近使用和一键复用。", "选择后带入图片与已确认字段。"],
        ["FR-LIB-05", "团队级数据隔离，只有同团队授权成员可见。", "接口必须验证 teamId，不能信任客户端参数。"],
    ], [1420, 4900, 3040])

    add_heading(doc, "5.5 FR-INTEGRATION：下游复用", 2)
    add_bullets(doc, [
        "爆款成片、高保真复刻、增强复刻统一接收 productId + productSnapshotId，不再复制一套商品字段。",
        "开始创作时生成不可变快照；后续编辑商品不影响已存在的项目和生成任务。",
        "下游允许在项目内覆盖个别字段，但项目覆盖不自动回写商品库，避免误修改团队底座。",
        "记录 product_reused 事件，关联业务入口、项目 ID、快照版本和发起用户。",
    ])

    add_heading(doc, "5.6 FR-AVATAR：数字人库（P1 相邻工作流）", 2)
    add_bullets(doc, [
        "预置 30–50 个数字人，标签至少包含地区、性别呈现、年龄段、语言、服装风格与授权范围。",
        "优先覆盖欧美、东南亚、南美；数量按内容资源与授权进度逐步上线。",
        "用户上传自有数字人时必须确认肖像/素材使用权，进入 processing → review → active / rejected 状态。",
        "数字人与商品分离建模，通过创作项目引用；不把人物信息混入商品识别结果。",
    ])

    page_break(doc)

    # UX and state
    add_heading(doc, "6. 信息架构与交互状态", 1)
    add_heading(doc, "6.1 页面与入口", 2)
    add_table(doc, ["页面 / 组件", "关键内容", "主要动作"], [
        ["商品库列表", "商品卡片/表格、搜索筛选、状态、最近使用", "新建、查看、选择、复制、归档"],
        ["新建商品", "链接导入、图片上传、手工创建三种方式", "提交识别、保存草稿"],
        ["识别进度", "排队、解析、识别、合并、结构化、失败原因", "取消、重试、手工录入"],
        ["结果确认", "字段分组、置信度、证据、AI 值与编辑值", "修改、恢复、确认入库"],
        ["商品详情", "最新版本、历史版本、引用记录、来源", "编辑、复制、归档、用于创作"],
        ["统一商品选择器", "最近使用、搜索、缩略图、关键卖点", "选择并生成项目快照"],
    ], [2000, 4260, 3100])

    add_heading(doc, "6.2 商品与任务状态", 2)
    add_callout(doc, "商品状态", "draft → recognizing → review_required → active → archived；识别失败进入 manual_required，可重试后回到 recognizing。", LIGHT_BLUE, BLUE)
    add_callout(doc, "任务状态", "queued → fetching / uploading → parsing → recognizing → merging → succeeded；任一步骤可进入 failed 或 cancelled，并保留 errorCode。", LIGHT_BLUE, BLUE)

    add_heading(doc, "6.3 必须覆盖的界面状态", 2)
    add_bullets(doc, [
        "空状态：说明商品库价值，并提供“粘贴链接”“上传图片”“手工创建”三种入口。",
        "加载状态：分阶段展示进度，不用不可解释的无限转圈；耗时超过阈值显示可离开提示。",
        "部分成功：已识别字段可继续审核，失败图片可单独重试。",
        "冲突状态：检测到疑似重复商品时展示对比，不静默覆盖。",
        "权限状态：无编辑权限时只读；无查看权限时返回统一无权限结果，不泄露是否存在。",
    ])

    add_heading(doc, "7. 数据模型", 1)
    add_heading(doc, "7.1 Product（可变业务实体）", 2)
    add_table(doc, ["字段组", "字段建议", "说明"], [
        ["标识与归属", "id, teamId, sku, status", "teamId 为强隔离边界；sku 可为空"],
        ["基础信息", "name, brand, category, locale", "name 为正式入库必填"],
        ["营销信息", "description, sellingPoints[], audience, scenes[]", "卖点建议 1–5 条；场景数组去重"],
        ["价格优惠", "priceOriginal, priceCurrent, currency, discounts[]", "仅链接来源或图片可见信息；不保证实时"],
        ["媒体", "mainImageId, imageIds[]", "主图必填；引用资产库 ID"],
        ["来源", "sourceType, sourceUrl, sourceHash", "sourceType: link / image / manual / copy"],
        ["审计", "createdBy, confirmedBy, createdAt, updatedAt, archivedAt", "确认动作单独留痕"],
    ], [1800, 4200, 3360])

    add_heading(doc, "7.2 ProductSnapshot（不可变引用）", 2)
    add_paragraph(doc, "保存创作项目时复制当时有效商品字段并计算 snapshotVersion；快照只追加不修改。下游只读取快照，商品详情可显示其被哪些项目引用。")
    add_heading(doc, "7.3 RecognitionTask（识别任务）", 2)
    add_paragraph(doc, "建议字段：id、teamId、productDraftId、sourceType、inputAssets、normalizedUrl、status、progressStage、provider、model、promptVersion、schemaVersion、rawResponseRef、parsedResult、errorCode、startedAt、finishedAt、latencyMs、token/成本信息。")
    add_heading(doc, "7.4 DigitalAvatar（P1）", 2)
    add_paragraph(doc, "建议字段：id、teamId/null（平台预置）、name、regionTags、languageTags、appearanceTags、rightsScope、sourceAssetIds、moderationStatus、status、createdBy、createdAt。")

    # AI recommendation
    add_heading(doc, "8. AI 模型选型建议", 1)
    add_callout(doc, "推荐方案", "MVP 主模型：Qwen3.7-Plus，使用非思考模式、结构化输出和严格 JSON Schema。Gemini 3.6 Flash 作为海外/英文商品与回归对照备选；Doubao Seed 2.0 Lite（或火山引擎当前可用同级视觉模型）进入 A/B 评测，不在评测前直接绑定生产。", LIGHT_GREEN, GREEN)

    add_heading(doc, "8.1 候选比较", 2)
    add_table(doc, ["模型", "适配理由", "注意点", "建议角色"], [
        ["Qwen3.7-Plus", "官方视觉文档直接给出商品图结构化提取场景；支持图像/视频、长上下文、函数调用与结构化输出；中文电商语义和国内调用链更贴合当前产品。", "结构化输出需使用非思考模式；仍需限制不可见规格的臆测。", "MVP 主模型"],
        ["Gemini 3.6 Flash", "图像理解、分类/VQA、多图、Schema 约束输出和对象定位能力完整，适合作为英文/海外商品与复杂图片的质量标杆。", "需要评估中国大陆链路、成本、数据合规与供应稳定性。", "备选 / 质量对照"],
        ["Doubao Seed 2.0 Lite", "国内服务链路友好，具备图像/视频多模态能力，可与字节生态部署协同。", "公开资料对商品字段 JSON 稳定性的直接证据较少，必须实测有效 JSON 率、幻觉率和成本。", "A/B 候选"],
    ], [1760, 3600, 2460, 1540])

    add_heading(doc, "8.2 为什么不直接绑定单一供应商", 2)
    add_bullets(doc, [
        "模型版本、价格和可用区会变化；业务层应只依赖统一的 ProductRecognitionResult。",
        "不同类目和语言可能表现差异明显；路由层可按语言、类目、失败次数选择 provider。",
        "保留原始响应引用、provider/model/prompt/schema 版本，便于回放、回归和账单核对。",
        "当主模型超时、限流或返回非法 JSON 时，允许一次同模型修复；仍失败再切换备选或人工。",
    ])

    add_heading(doc, "8.3 上线前评测", 2)
    add_paragraph(doc, "建立 150–300 个商品的标注集，至少覆盖服饰、美妆、食品、3C、家居、工具，以及中文、英文和东南亚常见语言。每个样本保留 1–8 张图片、人工金标与允许为空的字段。")
    add_table(doc, ["指标", "定义", "建议门槛"], [
        ["有效 JSON 率", "无需人工修复即可通过 Schema 校验", "≥ 99.5%"],
        ["名称/品类正确率", "与人工金标一致或语义等价", "≥ 95%"],
        ["营销字段通过率", "描述、卖点、受众、场景经审核可直接使用", "≥ 85%"],
        ["硬规格幻觉率", "生成图片中不可见且来源中不存在的确定性规格", "≤ 3%"],
        ["P95 延迟", "1–5 张图，从提交到结构化结果", "≤ 8 秒（需压测确认）"],
        ["降级成功率", "主模型失败后通过修复/备选/人工产出可审核结果", "≥ 98%"],
    ], [2300, 4760, 2300])

    page_break(doc)

    # Contract
    add_heading(doc, "9. 识别输出契约与 Prompt 约束", 1)
    add_heading(doc, "9.1 结构化输出（示例）", 2)
    schema_text = '''{
  "schemaVersion": "1.0",
  "product": {
    "name": {"value": "string|null", "confidence": 0.0, "evidence": ["image:1"]},
    "brand": {"value": "string|null", "confidence": 0.0, "evidence": []},
    "category": {"value": "string|null", "confidence": 0.0, "evidence": []},
    "description": {"value": "string", "confidence": 0.0, "evidence": []},
    "sellingPoints": [{"value": "string", "confidence": 0.0, "evidence": []}],
    "audience": {"value": "string", "confidence": 0.0, "evidence": []},
    "scenes": [{"value": "string", "confidence": 0.0, "evidence": []}],
    "visibleFacts": [{"key": "string", "value": "string", "evidence": []}],
    "price": {"original": null, "current": null, "currency": null},
    "discounts": []
  },
  "warnings": ["string"],
  "needsUserDecision": []
}'''
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(schema_text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(TEXT)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_heading(doc, "9.2 生成约束", 2)
    add_numbered(doc, [
        "只基于提供的图片、页面文本和结构化数据；看不到或无法验证的硬规格返回 null。",
        "把事实与营销推断分开：visibleFacts 必须有证据；受众和场景可推断，但需降低 confidence。",
        "卖点 1–5 条，描述 40–120 个中文字符或等价英文长度；避免医疗、功效和比较级违规表述。",
        "多图先判断是否为同一商品；不确定时写入 needsUserDecision，不得静默合并。",
        "输出必须通过服务端 JSON Schema；模型原始文本不得直接写入正式 Product。",
    ])

    add_heading(doc, "9.3 置信度与证据展示", 2)
    add_table(doc, ["置信度", "前端表现", "默认行为"], [
        ["≥ 0.85", "正常显示，可查看来源", "允许用户直接确认"],
        ["0.60–0.84", "黄色提示“建议核对”", "确认前聚焦该字段"],
        ["< 0.60 / null", "红色提示或空值", "不得以确定语气填充；建议手工补充"],
    ], [1900, 3860, 3600])

    add_heading(doc, "10. 接口与集成建议", 1)
    add_table(doc, ["方法", "接口", "用途", "关键返回"], [
        ["POST", "/api/products/recognition-tasks", "创建链接/图片识别任务", "taskId, draftProductId, status"],
        ["GET", "/api/products/recognition-tasks/{id}", "查询进度与结果", "stage, progress, result, errorCode"],
        ["POST", "/api/products/recognition-tasks/{id}/retry", "重试失败步骤", "newAttemptId, status"],
        ["POST", "/api/products", "确认后创建正式商品", "productId, snapshotVersion"],
        ["GET", "/api/products", "列表、搜索、筛选", "items, pagination, facets"],
        ["GET/PATCH", "/api/products/{id}", "查看或编辑商品", "product, version, updatedAt"],
        ["POST", "/api/products/{id}/archive", "归档商品", "status"],
        ["POST", "/api/products/{id}/snapshots", "为下游生成不可变快照", "snapshotId, snapshotVersion"],
    ], [900, 3100, 3000, 2360])

    add_heading(doc, "10.1 幂等与并发", 2)
    add_bullets(doc, [
        "创建任务接受 Idempotency-Key；同 key 重放返回同一 taskId。",
        "编辑商品使用 version/updatedAt 做乐观锁；冲突返回 409 并要求刷新比较。",
        "任务回调与轮询均按 taskId + attempt 号去重，旧 attempt 不得覆盖新结果。",
        "图片先写资产库并病毒/内容扫描，再进入模型调用；模型服务只接收短期签名 URL。",
    ])

    page_break(doc)

    # Exceptions/security
    add_heading(doc, "11. 异常、降级与可观测性", 1)
    add_table(doc, ["异常", "用户可见反馈", "系统处理"], [
        ["链接不可访问/反爬", "该页面暂不支持自动读取，可上传商品图或手工创建", "记录 domain 与 errorCode；不无限重试"],
        ["图片模糊/遮挡", "标记低置信字段并提示补图", "保留可用字段，允许局部重试"],
        ["多商品冲突", "要求选择主商品或拆分", "暂停正式入库"],
        ["模型非法 JSON", "识别仍在处理或转人工", "Schema 修复一次，再切换 provider"],
        ["超时/限流", "展示可重试状态，可离开页面", "指数退避；达到阈值后降级"],
        ["重复商品", "展示疑似已有商品对比", "允许复用、复制或继续创建"],
        ["内容风险", "提示该素材需调整或审核", "阻止 active，写入 moderationStatus"],
    ], [2200, 3660, 3500])

    add_heading(doc, "11.1 日志与监控", 2)
    add_bullets(doc, [
        "任务级：taskId、provider、model、prompt/schema 版本、阶段耗时、输入图片数、输出字段数、错误码。",
        "模型级：有效 JSON 率、字段缺失率、修复率、切换率、P50/P95 延迟、单任务成本。",
        "业务级：导入来源、确认时长、修改字段、入库、复用、下游入口、归档。",
        "隐私：日志不记录图片原始 URL、完整页面正文、个人信息或模型密钥。",
    ])

    add_heading(doc, "12. 权限、安全与合规", 1)
    add_table(doc, ["领域", "要求"], [
        ["租户隔离", "所有商品、任务、图片和快照必须绑定 teamId；服务端根据会话推导，不接受客户端越权指定。"],
        ["链接安全", "仅允许 HTTP/HTTPS；DNS/IP 解析后阻断本机、内网、云元数据地址与重定向绕过，防止 SSRF。"],
        ["上传安全", "校验 MIME 与文件签名，限制像素/体积，进行恶意文件和内容审核，使用私有存储与短期签名 URL。"],
        ["数据出境", "海外模型调用前确认数据区域、供应商条款与客户授权；支持按租户禁用海外 provider。"],
        ["保留与删除", "原始模型响应设置短期保留策略；归档不等同删除；按账户删除政策清理源文件与派生数据。"],
        ["数字人权利", "自定义数字人上传必须确认肖像权、版权与允许用途；审核失败不得进入 active。"],
    ], [1900, 7460])

    add_heading(doc, "13. 埋点与数据分析", 1)
    add_table(doc, ["事件", "触发时机", "关键属性"], [
        ["product_import_started", "提交链接/图片/手工创建", "sourceType, imageCount, entryPoint"],
        ["product_recognition_completed", "识别任务成功", "provider, model, latency, fieldsFilled"],
        ["product_recognition_failed", "任务失败", "stage, errorCode, provider, retryCount"],
        ["product_review_completed", "用户确认入库", "duration, editedFields, lowConfidenceCount"],
        ["product_reused", "下游选择商品", "productId, snapshotId, entryPoint, projectId"],
        ["product_archived", "商品归档", "ageDays, reuseCount"],
    ], [2700, 2800, 3860])

    page_break(doc)

    # Acceptance/rollout
    add_heading(doc, "14. 验收标准", 1)
    add_heading(doc, "14.1 功能验收", 2)
    add_table(doc, ["AC", "Given / When / Then"], [
        ["AC-01", "给定支持的公开商品链接，当用户提交识别时，系统产出可编辑的名称、描述、卖点、受众、场景与可获取的价格优惠字段。"],
        ["AC-02", "给定 1–8 张清晰商品图片，当识别完成时，系统至少产出描述、卖点、目标人群与使用场景，且每个 AI 字段含置信度和证据。"],
        ["AC-03", "给定不可见硬规格，当模型无法验证时，对应字段为空或进入 warning，不生成确定性参数。"],
        ["AC-04", "给定用户修改过的字段，当重新识别时，系统不覆盖用户值，而是展示差异供选择。"],
        ["AC-05", "给定已确认商品，当从三个创作入口打开商品选择器时，均可搜索、选择并生成不可变快照。"],
        ["AC-06", "给定商品已被历史项目使用，当商品被编辑或归档时，历史项目仍能读取原快照。"],
        ["AC-07", "给定模型超时或 JSON 非法，当自动修复/降级失败时，用户可保留草稿并手工完成入库。"],
        ["AC-08", "给定用户没有目标团队权限，当请求商品或任务时，系统拒绝访问且不泄露资源存在性。"],
    ], [1200, 8160])

    add_heading(doc, "14.2 非功能验收", 2)
    add_bullets(doc, [
        "性能：商品列表首屏 P95 ≤ 2 秒；识别任务异步，1–5 张图片 P95 目标 ≤ 8 秒，以压测结果调整。",
        "可用性：识别失败不丢失已上传图片和已编辑字段；刷新页面可恢复任务。",
        "兼容性：桌面端 Chrome/Edge 最近两个主版本；移动端可查看，MVP 不要求完整编辑。",
        "质量：核心接口具备单元/集成测试；至少覆盖权限、幂等、重复回调、非法 JSON、并发编辑。",
        "可观测性：所有失败均有稳定 errorCode，监控可按 provider/model/domain/类目切片。",
    ])

    add_heading(doc, "15. 实施计划与依赖", 1)
    add_table(doc, ["阶段", "建议交付", "退出条件"], [
        ["W1：基础闭环", "数据模型、图片上传、Qwen3.7-Plus 适配器、Schema、确认页、基础商品列表", "图片识别 → 确认 → 入库可走通；关键测试通过"],
        ["W2：链接与复用", "链接解析、重复检测、详情/编辑、统一选择器、快照接口", "三个下游入口均能复用商品"],
        ["W3：质量与灰度", "Gemini/字节评测、降级、指标看板、安全加固、灰度开关", "评测门槛达标；灰度无 P0 安全问题"],
        ["P1 并行", "数字人预置扩充、自定义上传、授权与审核", "30–50 个资源按内容计划到位"],
    ], [1600, 4900, 2860])

    add_heading(doc, "15.1 外部与内部依赖", 2)
    add_bullets(doc, [
        "产品链接解析：支持站点白名单、页面抓取服务、反爬边界与来源存证。",
        "产品图片解析：模型账号/配额、对象存储、签名 URL、内容审核、Schema 校验与评测集。",
        "资产库：图片上传、缩略图、去重、权限与生命周期。",
        "下游创作：爆款成片、高保真复刻、增强复刻统一接入商品选择器和快照 ID。",
        "内容运营：数字人素材、地区覆盖、肖像/版权授权与标签质量。",
    ])

    add_heading(doc, "16. 评审待确认项", 1)
    add_numbered(doc, [
        "8 月范围是否包含链接识别的实际抓取服务，还是先以白名单站点 / 模拟解析完成闭环？",
        "图片数量、单图大小、像素上限与支持格式是否接受“1–8 张、≤10 MB”的 MVP 假设？",
        "商品正式入库的最小必填字段是否为商品名、主图、描述、至少 1 条卖点？",
        "产品价格是否仅作识别时快照，需不需要在 UI 明示“非实时价格”？",
        "海外模型是否允许处理客户商品图；哪些租户必须强制走国内 provider？",
        "数字人库是否属于同一 W1 工程范围，还是作为 P1 并行项目单独排期？",
        "业务目标 WAU +500、周视频 +5,000 的观察窗口和归因口径由谁负责？",
    ])

    add_heading(doc, "附录 A：当前系统兼容建议", 1)
    add_paragraph(doc, "当前代码已有商品库入口与 ProductBrief 类字段（名称、品类、卖点、受众、场景、价格等）。为减少迁移成本，MVP 可保留这些字段作为 Product 的展示层映射，但需新增 teamId、状态、来源、版本、证据、置信度与识别任务模型。")
    add_table(doc, ["现有概念", "目标概念", "兼容建议"], [
        ["ProductBrief.name/category", "Product.name/category", "直接映射；category 后续改为字典 ID"],
        ["sellingPoints[]", "Product.sellingPoints[]", "保留数组；新增排序与字段级证据"],
        ["audience/scenes[]", "Product.audience/scenes[]", "保留；统一去重与 locale"],
        ["price", "priceOriginal/priceCurrent/currency", "拆为结构化价格，不再用单字符串"],
        ["image/url", "Asset + sourceUrl", "图片走资产库，URL 仅保留来源"],
        ["manual/link_ai_analysis", "sourceType + RecognitionTask", "来源与识别方式解耦"],
    ], [2500, 3060, 3800])

    add_heading(doc, "附录 B：官方模型资料", 1)
    sources = [
        ("阿里云百炼：视觉理解模型（Qwen3.7-Plus 与商品图结构化提取）", "https://help.aliyun.com/zh/model-studio/vision-model/"),
        ("阿里云百炼：Qwen 结构化输出", "https://help.aliyun.com/en/model-studio/qwen-structured-output"),
        ("Google AI for Developers：Gemini 图像理解", "https://ai.google.dev/gemini-api/docs/image-understanding"),
        ("Google AI for Developers：Gemini 模型", "https://ai.google.dev/gemini-api/docs/models"),
        ("火山引擎方舟：模型发布与多模态能力", "https://www.volcengine.com/docs/6492/2165228?lang=en"),
        ("火山引擎：图片理解接口", "https://www.volcengine.com/docs/6349/2198429?lang=en"),
    ]
    for idx, (label, url) in enumerate(sources, start=1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.50)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        hyperlink(p, label, url)

    add_callout(doc, "使用说明", "模型能力、名称、价格与可用区可能变化。工程实现前需重新核对官方文档与账号实际可用模型，并以本 PRD 的评测集和门槛做最终准入。", LIGHT_AMBER, AMBER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
