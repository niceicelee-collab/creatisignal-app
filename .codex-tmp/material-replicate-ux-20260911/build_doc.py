from pathlib import Path
import re
import json
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path(r'D:\Cursor\creatisignal-app-new')
OUT = BASE / 'deliverables/material-replicate-ux-20260911'
NAME = '素材分析与高保真复刻字段优化方案_V1.0'
source = (OUT / (NAME + '.md')).read_text(encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.63)
sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.7)
sec.right_margin = Inches(.7)
sec.footer_distance = Inches(.28)
for grid in sec._sectPr.findall(qn('w:docGrid')):
    sec._sectPr.remove(grid)

def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        style.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:' + attr), None)
    ppr = style.element.get_or_add_pPr()
    for border in ppr.findall(qn('w:pBdr')):
        ppr.remove(border)
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); ppr.append(snap)

for style in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    font(doc.styles[style], 11)
    pf = doc.styles[style].paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = Pt(16)
    pf.widow_control = True
for name, size in [('Title', 23), ('Subtitle', 10), ('Heading 1', 17), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
    font(doc.styles[name], size, name != 'Subtitle')
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(10 if name.startswith('Heading') else 0)
    pf.space_after = Pt(7)
    pf.keep_with_next = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)

def inline(p, text):
    parts = re.split(r'(\*\*.+?\*\*|\[[^]]+\]\(https?://[^)]+\))', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif re.match(r'\[[^]]+\]\(https?://', part):
            title, url = re.match(r'\[([^]]+)\]\(([^)]+)\)', part).groups()
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
            run = OxmlElement('w:r')
            rp = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '245B47'); rp.append(color)
            run.append(rp)
            tx = OxmlElement('w:t'); tx.text = title; run.append(tx)
            hyperlink.append(run); p._p.append(hyperlink)
        else:
            p.add_run(part)

def table(rows):
    data = [x for x in rows if not all(re.fullmatch(r':?-+:?', c.strip()) for c in x)]
    cols = len(data[0])
    t = doc.add_table(rows=0, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if cols == 2:
        fractions = [.24, .76]
    elif data[0][1] == '处理':
        fractions = [.33, .13, .54]
    elif cols == 4:
        fractions = [.19, .27, .27, .27]
    else:
        fractions = [.27, .43, .30]
    for col, frac in zip(t.columns, fractions):
        col.width = Inches(7.1 * frac)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        x = OxmlElement('w:' + side); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), 'D8DFDA'); borders.append(x)
    pr.append(borders)
    for index, texts in enumerate(data):
        row = t.add_row()
        rpr = row._tr.get_or_add_trPr()
        rpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            rpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, text) in enumerate(zip(row.cells, texts)):
            cell.width = Inches(7.1 * fractions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top', 75), ('left', 95), ('bottom', 75), ('right', 95)]:
                e = OxmlElement('w:' + side); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcpr.append(margins)
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EDF2EE'); tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(14)
            if index == 0:
                p.paragraph_format.keep_with_next = True
            if cols == 3 and data[0][1] == '处理' and j == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, text)
            for run in p.runs:
                run.font.size = Pt(10)
                if index == 0:
                    run.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(2)

lines = source.splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line == '<!-- pagebreak -->':
        pass
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        table(rows)
        continue
    elif line.startswith('# '):
        inline(doc.add_paragraph(style='Title'), line[2:])
    elif line.startswith('## '):
        inline(doc.add_paragraph(style='Heading 1'), line[3:])
    elif line.startswith('### '):
        inline(doc.add_paragraph(style='Heading 2'), line[4:])
    elif line.startswith('- '):
        inline(doc.add_paragraph(style='List Bullet'), line[2:])
    elif re.match(r'^\d+\. ', line):
        inline(doc.add_paragraph(), line)
    elif line.startswith('版本 '):
        inline(doc.add_paragraph(style='Subtitle'), line)
    else:
        inline(doc.add_paragraph(), line)
    i += 1

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('CreatiSignal  |  ')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string('777777')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = '素材分析与高保真复刻字段优化方案'
doc.core_properties.subject = '字段删减 呈现方式 当前只读与后续编辑'
doc.core_properties.author = 'CreatiSignal'
doc.core_properties.keywords = '素材分析 高保真复刻 字段优化 产品需求'
target = OUT / (NAME + '.docx')
doc.save(target)
with ZipFile(target) as z:
    assert z.testzip() is None
print(json.dumps({'docx': str(target), 'markdown_characters': len(source), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)}, ensure_ascii=False))
