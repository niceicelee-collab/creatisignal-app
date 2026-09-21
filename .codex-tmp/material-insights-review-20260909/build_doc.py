from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent.parent / 'deliverables' / 'material-insights-review-20260909'
OUT.mkdir(parents=True, exist_ok=True)
source = (ROOT / 'content.md').read_text(encoding='utf-8')
source = re.sub(r'(?m)^\| E(\d{2}) ', r'| ST\1 ', source)
(ROOT / 'content.md').write_text(source, encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.6)
sec.left_margin, sec.right_margin = Inches(.68), Inches(.68)
sec.footer_distance = Inches(.25)

def style_font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    style.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    lang = OxmlElement('w:lang'); lang.set(qn('w:eastAsia'), 'zh-CN')
    style.element.get_or_add_rPr().append(lang)

style_font(doc.styles['Normal'], 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(16)
normal.space_after = Pt(6)
normal.widow_control = True
for name,size in [('Title',23),('Subtitle',12),('Heading 1',18),('Heading 2',12)]:
    style_font(doc.styles[name],size,name != 'Subtitle')
    f=doc.styles[name].paragraph_format
    f.space_before=Pt(7 if name=='Heading 2' else 0)
    f.space_after=Pt(4 if name=='Heading 2' else 10)
    f.line_spacing=Pt(32 if name=='Title' else (26 if name=='Heading 1' else 18))
    f.keep_with_next=True

for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('CreatiSignal  ·  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
footer._p.append(field)

def inline(p, text):
    for n,s in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        if s:
            r=p.add_run(s);r.bold=bool(n%2)

def link(p,url,label):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'24536B');rp.append(c)
    r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def table(lines):
    values = [[v.strip() for v in line.strip('|').split('|')] for line in lines]
    values=[row for row in values if not all(re.fullmatch(r'[:\- ]+',v) for v in row)]
    widths=[1.65,2.65,2.84] if len(values[0])==3 else [2.0,5.14]
    if len(values[0])==2: widths=[2,5.14]
    tab=doc.add_table(rows=0, cols=len(values[0]));tab.autofit=False
    for col,w in zip(tab.columns,widths):col.width=Inches(w)
    pr=tab._tbl.tblPr
    w=pr.find(qn('w:tblW'));w.set(qn('w:w'),str(round(7.14*1440)));w.set(qn('w:type'),'dxa')
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for idx,rowdata in enumerate(values):
        row=tab.add_row()
        trpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if idx==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,(cell,txt) in enumerate(zip(row.cells,rowdata)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EFF2' if idx==0 else ('F7F9FA' if idx%2==0 else 'FFFFFF'));tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top',85),('bottom',85),('left',95),('right',95)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14)
            inline(p,txt)
            for r in p.runs:r.font.size=Pt(10);r.bold=idx==0
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(2)

pages=source.split('\n---PAGE---\n')
for pi,page in enumerate(pages):
    lines=page.strip().splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip());i+=1
            table(rows);continue
        if s.startswith('# '):
            p=doc.add_paragraph(s[2:], 'Title' if pi==0 else 'Heading 1')
            if pi:p.paragraph_format.page_break_before=True
        elif re.fullmatch(r'\*\*.*\*\*',s):
            doc.add_paragraph(s[2:-2], 'Heading 2')
        elif s.startswith('https://'):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7)
            label='打开 TikTok 官方资料'
            if 'attribution' in s:label='TikTok 官方 GMV Max 归因说明'
            elif 'tips-to-measure' in s:label='TikTok 官方 GMV Max 素材评价建议'
            elif 'reporting' in s:label='TikTok 官方 GMV Max 报表说明'
            link(p,s,label)
        else:
            p=doc.add_paragraph()
            if pi==0 and i<6:p.paragraph_format.space_after=Pt(9)
            if re.match(r'^\d+\. ',s):
                p.paragraph_format.space_after=Pt(3 if pi==15 else 5)
            inline(p,s)
        i+=1

doc.core_properties.title='素材洞察与数据源产品优化方案'
doc.core_properties.subject='账户接入 空状态 信息架构 诊断规则 逐页优化'
doc.core_properties.author='CreatiSignal'
doc.core_properties.keywords='素材洞察, 数据源, TikTok, GMV Max, 产品优化'
doc.core_properties.comments=''
path=OUT/'CreatiSignal_素材洞察与数据源产品优化方案_V1.0.docx'
doc.save(path)
(ROOT/'manifest.json').write_text(json.dumps({'path':str(path),'planned_pages':len(pages),'characters':len(source),'tables':len(doc.tables)},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'path':str(path),'planned_pages':len(pages),'characters':len(source),'tables':len(doc.tables)},ensure_ascii=False))
