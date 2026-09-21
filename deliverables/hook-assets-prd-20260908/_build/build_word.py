from pathlib import Path
import re, json
from zipfile import ZipFile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Hook资产_产品需求文档_V1.0.md'
DEST = ROOT / 'Hook资产_产品需求文档_V1.0.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.9)
section.bottom_margin = Cm(1.7)
section.left_margin = section.right_margin = Cm(1.7)
section.header_distance = Cm(.8)
section.footer_distance = Cm(.8)

def font(style, size, bold=False, color='000000'):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for name in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:'+name), 'Microsoft YaHei')

font(doc.styles['Normal'], 10)
norm = doc.styles['Normal'].paragraph_format
norm.line_spacing = 1.45
norm.space_after = Pt(6)
norm.widow_control = True
for name, size in [('Title',26), ('Heading 1',20), ('Heading 2',12)]:
    font(doc.styles[name],size,True)
    pf=doc.styles[name].paragraph_format
    pf.space_before=Pt(8 if name=='Heading 2' else 0)
    pf.space_after=Pt(6 if name=='Heading 2' else 10)
    pf.line_spacing=1.25
    pf.keep_with_next=True
    borders=doc.styles[name].element.get_or_add_pPr().find(qn('w:pBdr'))
    if borders is not None:borders.getparent().remove(borders)
font(doc.styles['Caption'],8,False,'737D84')
doc.styles['Caption'].paragraph_format.space_after=Pt(7)
doc.styles['Caption'].paragraph_format.line_spacing=1.1
for n in ['Header','Footer']:
    font(doc.styles[n],8,False,'000000' if n=='Header' else '7B858D')
    doc.styles[n].paragraph_format.space_after=Pt(0)

header=section.header.paragraphs[0]
header.text='CREATISIGNAL　产品需求　　　　　　　　　　　　　　Hook 资产　V1.0'
footer=section.footer.paragraphs[0]
footer.text='2026年9月8日　产品与工程评审'
footer.paragraph_format.tab_stops.add_tab_stop(Cm(15.6))
footer.add_run('\t')
for instr in ['PAGE','NUMPAGES']:
    if instr=='NUMPAGES':footer.add_run(' / ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),instr)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1' if instr=='PAGE' else '26';r.append(t);fld.append(r);footer._p.append(fld)

def rich(p,text,size=None):
    for token in re.split(r'(\*\*.*?\*\*)',text):
        if not token:continue
        r=p.add_run(token[2:-2] if token.startswith('**') else token)
        if token.startswith('**'):r.bold=True
        if size is not None:r.font.size=Pt(size)
    return p

def paragraph(text,feature=False,style=None):
    p=doc.add_paragraph(style=style)
    rich(p,text,9.5 if feature else None)
    if feature:
        p.paragraph_format.line_spacing=1.4
        p.paragraph_format.space_after=Pt(5)
    return p

def table(rows):
    count=len(rows[0])
    if count==4: widths=[1.3,4.4,3.2,8.7]
    elif rows[0][0]=='编号': widths=[1.3,5.3,11.0]
    elif count==3 and rows[0][0]=='事件':widths=[4.0,5.3,8.3]
    elif count==3:widths=[3.7,5.0,8.9]
    elif count==2:widths=[4.3,13.3]
    else:widths=[17.6/count]*count
    t=doc.add_table(rows=0,cols=count)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(Cm(17.6).twips)))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,value in [('top',65),('bottom',65),('left',90),('right',90)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for idx,values in enumerate(rows):
        row=t.add_row();trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
        for j,(cell,txt) in enumerate(zip(row.cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx==0 or idx%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F2' if idx==0 else 'FAFBFC');cell._tc.get_or_add_tcPr().append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.3
            if rows[0][0]=='编号' and j==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            rich(p,txt,8.8)
            if idx==0:
                p.paragraph_format.keep_with_next=True
                for r in p.runs:r.bold=True
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(0);after.paragraph_format.line_spacing=1
    after.paragraph_format.space_before=Pt(0);after.add_run().font.size=Pt(3)

text=SOURCE.read_text(encoding='utf-8')
sections=text.split('<!-- page -->')
expected_text=[]
image_count=0
for idx,part in enumerate(sections):
    lines=part.strip().splitlines();i=0;feature='![' in part
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('# '):
            style='Title' if idx==0 else 'Heading 1'
            p=paragraph(line[2:],style=style)
            if idx>0:p.paragraph_format.page_break_before=True
            else:p.paragraph_format.space_after=Pt(16)
            expected_text.append(line[2:])
        elif line.startswith('## '):
            paragraph(line[3:],style='Heading 2');expected_text.append(line[3:])
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1
            picture=p.add_run().add_picture(str(ROOT/m[2]),width=Cm(16.3))
            picture._inline.docPr.set('descr',m[1]+' 静态原型示例')
            caption=paragraph(m[1]+'　静态原型示例',style='Caption')
            caption.paragraph_format.keep_with_next=True
            image_count+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                r=lines[i].strip()
                if not re.match(r'^\|[\s:|-]+\|$',r):rows.append([x.strip() for x in r.strip('|').split('|')])
                i+=1
            table(rows);expected_text.extend(c for r in rows for c in r);continue
        elif line.startswith('- '):
            p=paragraph('• '+line[2:],feature)
            p.paragraph_format.left_indent=Cm(.35);p.paragraph_format.first_line_indent=Cm(-.35)
            expected_text.append(line[2:])
        else:
            paragraph(line,feature);expected_text.append(line)
        i+=1

doc.core_properties.title='Hook资产产品需求文档'
doc.core_properties.subject='Hook资产提取 保存与替换生成产品需求'
doc.core_properties.author='CreatiSignal'
doc.core_properties.keywords='Hook,Qwen3.7plus,Seedance2.0,Seedance2.5,产品需求'
doc.save(DEST)

check=Document(DEST)
actual='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
missing=[s for s in expected_text if s.replace('**','') not in actual]
with ZipFile(DEST) as z:
    invalid=z.testzip();media=[n for n in z.namelist() if n.startswith('word/media/')]
    xml=z.read('word/document.xml').decode('utf-8')
report={'file':str(DEST),'bytes':DEST.stat().st_size,'paragraphs':len(check.paragraphs),'tables':len(check.tables),'images':len(check.inline_shapes),'embedded_media':len(media),'source_sections':len(sections),'explicit_page_starts':xml.count('w:pageBreakBefore'),'missing_source_blocks':missing,'zip_error':invalid,'has_full_system_prompt':all(s in actual for s in ['你是一名电商短视频创意分析师','第五部分为待确认项和结论','Qwen3.7plus','Seedance2.0','Seedance2.5']),'all_acceptance_items':all(f'A{i:02d}' in actual for i in range(1,33))}
(ROOT/'_build'/'word_structural_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
assert not missing and invalid is None and len(check.inline_shapes)==12 and report['has_full_system_prompt'] and report['all_acceptance_items']
