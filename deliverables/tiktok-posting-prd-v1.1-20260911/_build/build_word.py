from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from content import ROOT, PAGES, SOURCES

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(11.6929)
sec.page_height = Inches(8.2677)
sec.left_margin = sec.right_margin = Inches(0.53)
sec.top_margin = Inches(0.43)
sec.bottom_margin = Inches(0.48)
sec.header_distance = Inches(0.15)
sec.footer_distance = Inches(0.19)

def font_style(style, size, bold=False, color='000000'):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + key), 'Microsoft YaHei')
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:' + key), None)
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)

for name, size, bold in [('Normal',10.2,False),('Title',24,True),('Subtitle',9.5,False),('Heading 1',19,True),('Heading 2',12.3,True),('Caption',10.2,False),('Footer',8,False)]:
    font_style(doc.styles[name],size,bold)
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.23
normal.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:
    fmt=doc.styles[name].paragraph_format
    fmt.space_before=Pt(0 if name!='Heading 2' else 7)
    fmt.space_after=Pt(9 if name!='Heading 2' else 5)
    fmt.line_spacing=1.1
    fmt.keep_with_next=True
doc.styles['Subtitle'].paragraph_format.space_after=Pt(11)
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.styles['Caption'].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.italic=False

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('CreatiSignal  ·  TikTok 发帖产品需求文档  V1.1    ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE')
footer._p.append(fld)

def table(headers, rows):
    widths=[1.85,8.76] if len(headers)==2 else [1.79,5.44,3.38]
    if len(headers)==3 and headers[1]=='默认':widths=[2.57,.85,7.19]
    if headers[0]=='编号':widths=[.81,9.80]
    if headers[0]=='界面指标':widths=[1.68,3.35,5.58]
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF6DF' if i==0 else ('FAFBF8' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge,amount in [('top',85),('bottom',85),('left',130),('right',130)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(amount));e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.18
            if i==0:p.paragraph_format.keep_with_next=True
            if (headers[0]=='编号' and j==0) or (len(headers)==3 and headers[1]=='默认' and j==1):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(value);r.font.size=Pt(9.6 if i==0 else 9.4);r.bold=i==0
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(0);p.add_run().font.size=Pt(3)

def hyperlink(p, label, url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'255C45');rp.append(color)
    r.append(rp);txt=OxmlElement('w:t');txt.text=label;r.append(txt);h.append(r);p._p.append(h)

for idx,page in enumerate(PAGES):
    heading=doc.add_paragraph((page.get('proto','')+'  '+page['title']).strip(),style='Title' if idx==0 else 'Heading 1')
    if idx:heading.paragraph_format.page_break_before=True
    if page.get('sub'):doc.add_paragraph(page['sub'],style='Subtitle')
    if 'proto' in page:
        doc.add_paragraph(page['note'],style='Caption')
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
        shape=p.add_run().add_picture(str(ROOT/'prototypes'/(page['proto']+'.png')),width=Inches(9.70))
        shape._inline.docPr.set('descr',page['proto']+' '+page['title']+'。'+page['note'])
        continue
    for block in page['blocks']:
        if block[0]=='table':table(block[1],block[2])
        elif block[0]=='sources':
            for start in range(0,len(SOURCES),3):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
                for i,(label,url) in enumerate(SOURCES[start:start+3]):
                    if i:p.add_run('     ')
                    hyperlink(p,label,url)
        else:
            p=doc.add_paragraph(block[1],style='Heading 2' if block[0]=='h' else 'Normal')
            if block[0]=='note':
                for r in p.runs:r.font.size=Pt(9.2);r.font.color.rgb=RGBColor.from_string('526154')

doc.core_properties.title='TikTok 发帖产品需求文档'
doc.core_properties.subject='新建任务 任务列表 发布状态与 GMV Max 衔接'
doc.core_properties.author='CreatiSignal'
doc.core_properties.keywords='TikTok,产品需求,GMV Max'
doc.core_properties.comments=''
for border in doc.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
out=ROOT/'TikTok发帖_产品需求文档_V1.1.docx'
doc.save(out)
print(out)
print('Editable tables:',len(doc.tables),'Embedded prototypes:',len(doc.inline_shapes))
