from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'UGC爆款_产品需求文档_V1.0.md'
OUTPUT=ROOT/'UGC爆款_产品需求文档_V1.0.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.6)
sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.7)
sec.gutter=Cm(0)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))):sec._sectPr.remove(grid)

def fontset(obj,size,bold=False,color='171A20'):
    obj.font.name='Microsoft YaHei'
    obj.font.size=Pt(size);obj.font.bold=bold;obj.font.color.rgb=RGBColor.from_string(color)
    obj.font.italic=False
    rp=obj._element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),'Microsoft YaHei')
    for k in list(rf.attrib):
        if 'Theme' in k:del rf.attrib[k]

for st,size,bold in [('Normal',10.2,False),('Title',27,True),('Subtitle',10.2,False),('Heading 1',21,True),('Heading 2',12,True),('List Bullet',10.2,False),('List Number',10.2,False),('Caption',8,False),('Header',8,False),('Footer',8,False)]:
    s=doc.styles[st];fontset(s,size,bold,'000000' if st in ['Title','Subtitle','Heading 1','Heading 2','Header'] else '171A20')
    f=s.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(5)
    f.line_spacing=Pt(15.5);f.widow_control=True
    sp=s._element.get_or_add_pPr()
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');sp.append(snap)
    for border in list(sp.findall(qn('w:pBdr'))):sp.remove(border)
    if st in ['Title','Heading 1','Heading 2']:
        f.keep_with_next=True;f.keep_together=True
        f.space_after=Pt(8 if st!='Heading 2' else 4)
        f.space_before=Pt(9 if st=='Heading 2' else 0)
        f.line_spacing=Pt(size*1.28)
    if st in ['List Bullet','List Number']:
        f.left_indent=Cm(.4);f.first_line_indent=Cm(-.32)
    if st in ['Caption','Header','Footer']:
        f.space_after=Pt(0);f.line_spacing=1.1
    if st=='Caption':f.keep_with_next=False
    if st in ['Title','Subtitle']:s.base_style=doc.styles['Normal']
    if st in ['Header','Footer']:
        f.tab_stops.clear_all()

# Explicitly clear built-in Title border residue in Office renderers.
title_borders=OxmlElement('w:pBdr')
for edge in ['top','bottom','left','right','between']:
    border=OxmlElement('w:'+edge);border.set(qn('w:val'),'nil');title_borders.append(border)
doc.styles['Title']._element.get_or_add_pPr().append(title_borders)

# Preserve native editable paragraph and table text; use true list styles.
header=sec.header.paragraphs[0]
header.paragraph_format.tab_stops.add_tab_stop(Cm(17.4),WD_ALIGN_PARAGRAPH.RIGHT)
fontset(header.add_run('CREATISIGNAL　/　产品需求\tUGC爆款　V1.0'),8,color='000000')
footer=sec.footer.paragraphs[0]
footer.paragraph_format.tab_stops.add_tab_stop(Cm(17.4),WD_ALIGN_PARAGRAPH.RIGHT)
fontset(footer.add_run('2026年9月8日　·　产品与研发评审\t'),8,color='878E97')
def field(p,code):
    r=p.add_run();fontset(r,8,color='878E97')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),code)
    t=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';t.append(tt);f.append(t);r._r.addnext(f)
field(footer,'PAGE');fontset(footer.add_run(' / '),8,color='878E97');field(footer,'NUMPAGES')

expected=[]
def paragraph(txt,style='Normal',feature=False,dense=False):
    txt=re.sub(r'\*\*(.*?)\*\*',r'\1',txt)
    p=doc.add_paragraph(txt,style=style);expected.append(txt)
    if style in ['Normal','List Bullet','List Number']:
        size=9.6 if feature else 9.5 if dense else 10.2
        for r in p.runs:fontset(r,size)
        f=p.paragraph_format
        f.line_spacing=Pt(14.1 if feature else 14.3 if dense else 15.5)
        f.space_after=Pt(4 if feature or dense else 5)
        f.keep_together=False;f.keep_with_next=False
    return p

def xmlval(parent,tag,attrs):
    x=OxmlElement(tag)
    for k,v in attrs.items():x.set(qn(k),str(v))
    parent.append(x);return x

def widths_for(headers):
    n=len(headers)
    if n==2:return [3.3,14.1]
    if headers[0]=='编号' and n==3:return [1.4,5.5,10.5]
    if n==3:return [3.3,7.5,6.6]
    if n==4:return [3.4,5.7,3.0,5.3]
    if n==5:return [1.4,5.1,2.1,8.8] if len(headers)==4 else [1.3,4.0,1.8,3.8,6.5]
    return [17.4/n]*n

table_count=0
def addtable(rows,feature=False):
    global table_count
    t=doc.add_table(rows=len(rows),cols=len(rows[0]));table_count+=1
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths_for(rows[0])
    if rows[0]==['编号','功能','优先级','关键交付']:widths=[1.5,5.1,2.0,8.8]
    if rows[0][0]=='后续优先级':widths=[2.0,7.0,8.4]
    props=t._tbl.tblPr
    tblw=props.find(qn('w:tblW'));tblw.set(qn('w:type'),'dxa');tblw.set(qn('w:w'),str(round(17.4/2.54*1440)))
    borders=xmlval(props,'w:tblBorders',{})
    for edge in ['top','left','bottom','right','insideH','insideV']:
        xmlval(borders,'w:'+edge,{'w:val':'single','w:sz':4,'w:color':'D9D9D9'})
    for i,col in enumerate(t.columns):col.width=Cm(widths[i])
    for ri,(row,values) in enumerate(zip(t.rows,rows)):
        trpr=row._tr.get_or_add_trPr()
        xmlval(trpr,'w:cantSplit',{})
        if ri==0:xmlval(trpr,'w:tblHeader',{'w:val':'true'})
        for ci,(cell,txt) in enumerate(zip(row.cells,values)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            xmlval(tcpr,'w:shd',{'w:fill':'E7EDF0' if ri==0 else 'FAFBFD' if ri%2==0 else 'FFFFFF'})
            mar=xmlval(tcpr,'w:tcMar',{})
            for edge in ['top','bottom','left','right']:
                xmlval(mar,'w:'+edge,{'w:w':65 if edge in ['top','bottom'] else 105,'w:type':'dxa'})
            p=cell.paragraphs[0];txt=re.sub(r'\*\*(.*?)\*\*',r'\1',txt)
            p.text=txt;expected.append(txt)
            f=p.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=Pt(12.4 if feature else 13)
            f.keep_with_next=False;f.keep_together=False;f.widow_control=True
            if widths[ci]<=2.1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:fontset(r,8.6 if feature else 9,ri==0)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=Pt(4);p.add_run().font.size=Pt(1)

blocks=SOURCE.read_text(encoding='utf-8').split('<!-- page -->')
image_count=0
for bi,block in enumerate(blocks):
    feature='![' in block;dense=bi in [4,20,21,22,23,24,26,27,28,29]
    lines=block.strip().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('# '):
            p=paragraph(line[2:],'Title' if bi==0 else 'Heading 1')
            if bi>0:p.paragraph_format.page_break_before=True
        elif line.startswith('## '):paragraph(line[3:],'Heading 2',feature,dense)
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(ROOT/m[2]),width=Cm(11.5 if '|' in block else 13.0))
            shape._inline.docPr.set('descr',m[1]);shape._inline.docPr.set('title',m[1]);image_count+=1
            cap=doc.add_paragraph(m[1]+'　静态示例','Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after=Pt(5)
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                value=lines[i].strip()
                if not re.match(r'^\|[\s:|-]+\|$',value):rows.append([c.strip() for c in value.strip('|').split('|')])
                i+=1
            addtable(rows,feature);continue
        elif line.startswith('- '):paragraph(line[2:],'List Bullet',feature,dense)
        elif re.match(r'^\d+\. ',line):paragraph(re.sub(r'^\d+\. ','',line),'List Number',feature,dense)
        else:paragraph(line,'Subtitle' if bi==0 and line.startswith('CreatiSignal') else 'Normal',feature,dense)
        i+=1

doc.core_properties.title='UGC爆款产品需求文档'
doc.core_properties.subject='产品功能 运营流程 数据埋点与验收'
doc.core_properties.author='CreatiSignal'
doc.core_properties.keywords='UGC爆款,产品需求,运营审核,我的资产,收藏'
doc.core_properties.comments=''
doc.save(OUTPUT)

check=Document(OUTPUT)
actual=[''.join(p.itertext()) for p in []]
actual=[''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in check._element.body.iter(qn('w:p'))]
missing=Counter(expected)-Counter(actual)
assert not missing,missing
assert len(check.inline_shapes)==15
assert len(blocks)==30
assert all(('F'+str(i).zfill(2)) in '\n'.join(actual) for i in range(1,9))
assert all(('B'+str(i).zfill(2)) in '\n'.join(actual) for i in range(1,8))
assert all(('A'+str(i).zfill(2)) in '\n'.join(actual) for i in range(1,33))
with ZipFile(OUTPUT) as z:assert z.testzip() is None
report={'file':str(OUTPUT),'bytes':OUTPUT.stat().st_size,'logical_sections':len(blocks),'native_paragraphs':len(check.paragraphs),'native_tables':len(check.tables),'embedded_images':len(check.inline_shapes),'source_text_blocks':len(expected),'missing_text_blocks':len(missing),'zip_integrity':'passed'}
(ROOT/'_build/word_structure_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
